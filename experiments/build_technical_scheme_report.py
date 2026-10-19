"""Reconstruct the technical scheme + chronological experiment report, read-only to runs.

No simulation is executed and no existing report or raw experiment CSV is replaced.
"""
from __future__ import annotations
import argparse
from collections import Counter,defaultdict
from datetime import datetime,timezone,timedelta
import hashlib
import json
import math
from pathlib import Path
import shutil
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
import yaml
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from experiments.run_mixed_strategies import read,number,input_signature
from experiments.run_mixed_3600 import write_csv
from experiments.build_all_strategy_report import configure_document,table
from experiments.technical_scheme_content import build_technology

NAME='智能路由算法技术方案及实验'
OUT=ROOT/'results'/'technical_scheme_experiments_20261016'
SINGLE=ROOT/'results'/'all_strategy_single_biz_20261016'
DENSE=ROOT/'results'/'mixed_3600_latency_20261016'
SPARSE=ROOT/'results'/'mixed_3600_latency_sparse_peak_20261016'
OTHER=ROOT/'results'/'mixed_3600_other_strategies_20261016'
BOUNDED=ROOT/'results'/'mixed_3600_bounded_all_strategies_20261016'
FRESH=ROOT/'results'/'mixed_3600_random_batch_20261016_111126_73e445bc6f54'
STRATEGIES=('latency','balance','energy');BIZ=('LL','HR','HC','MC')
LABEL={'latency':'时延','balance':'均衡','energy':'能耗'}
COLOR={'latency':'#2864A0','balance':'#D88722','energy':'#27856B'}
CAP={1:200,2:250,3:500}
SOURCES=set()


def track(path):
    path=Path(path);SOURCES.add(path.resolve());return path


def rows(path):return [number(r) for r in read(track(path))]
def js(path):return json.loads(track(path).read_text(encoding='utf-8'))
def val(row,key,default=None):
    x=row.get(key,default)
    return default if x in ('',None) else x
def fmt(x,d=3):return '—' if x is None or x=='' else f'{float(x):.{d}f}'
def integer(x):return '—' if x is None or x=='' else str(round(float(x)))
def pct(x):return '—' if x is None or x=='' else f'{float(x)*100:.2f}%'
def trip(row,*keys,d=3):return ' / '.join(fmt(val(row,k),d) for k in keys)
def relative(a,b):return 100*(a/b-1) if b else 0
def marker(s):return f"场景{int(s['scenario'])}/{LABEL[s['strategy']]}"+(f"/{s['biz']}" if s['biz']!='ALL' else '')


def mixed_normalize(raw,phase):
    aliases=dict(generated='generated',accepted='initially_accepted',completed='completed',end_active='active_at_end',
        dropped='dropped',rcap='rejected_cap',rnone='rejected_no_feasible',accept_rate='initial_accept_rate',success_rate='final_survival_rate',
        avg_active='time_avg_active',peak_active='peak_concurrency',service_s='observed_service_seconds',
        delay='lifetime_delay_mean_ms',p95='lifetime_delay_p95_ms',dmax='lifetime_delay_max_ms',weighted_delay='service_weighted_delay_ms',
        tc='compute_delay_service_weighted_ms',loss='lifetime_loss_max',energy='energy_total_J',ec='energy_compute_J',et='energy_tx_J',
        etask='energy_per_initially_accepted_J',egbit='energy_J_per_input_Gbit',gbits='observed_input_Gbit',
        segments='allocation_segments',reroute='reroutes',migrations='migrations',source_pct='source_compute_segment_pct',ground_pct='ground_compute_segment_pct',
        link_peak='peak_link_util_pct',link_cv='mean_link_cv',link_jain='mean_link_jain',node_peak='peak_node_util_pct',node_cv='mean_node_cv',node_jain='mean_node_jain',
        pmean='mean_total_power_W',ppeak='peak_total_power_W',pcmean='mean_compute_power_W',ptmean='mean_tx_power_W',
        pcpeak='peak_compute_power_W',ptpeak='peak_tx_power_W',solve99='solve_p99_ms',fallback='fallback_attempts',recovered='fallback_accepted',
        wall='wall_seconds',cap_s='time_at_cap_s',cap_pct='time_at_cap_pct',near_s='time_above_90pct_cap_s',first_cap='first_cap_t',
        survival_admitted='initially_accepted_survival_rate',sha='input_sha256')
    s=dict(phase=phase,scenario=int(raw['scenario']),strategy=raw.get('strategy','latency'),biz=raw['biz'],T=3600.,mode='event')
    s.update({key:val(raw,original) for key,original in aliases.items()})
    s['pmean']=s['energy']/3600;s['pcmean']=s['ec']/3600;s['ptmean']=s['et']/3600
    s['complete_rate']=s['completed']/s['generated']
    s['survival_admitted']=(s['completed']+s['end_active'])/max(1,s['accepted'])
    if s['biz']!='ALL':s['margin']={'LL':20,'HR':200,'HC':200,'MC':200}[s['biz']]-s['dmax']
    return s


def stable_node_moments(path,scenario):
    """Reconstruct compute loads from live allocations, avoiding running-moment drift.

    This is reporting-only postprocessing, never a change to route selection or raw data.
    Every simultaneous start/end is grouped; zero-length states do not affect integrals.
    """
    cfg=js(path/'effective_config.json');capacity=cfg['compute_capacity_tflops']
    n={1:500,2:528,3:1006}[scenario]
    events=defaultdict(list)
    for a in rows(path/'routing_allocations.csv'):
        if a['service_duration_s']<=1e-12:continue
        entry=(a['compute_node'],a['allocation_id'],a['compute_rate_used']/(capacity[a['compute_node_type']]*1e12))
        events[a['t_alloc']].append((1,entry));events[a['service_t_end']].append((-1,entry))
    live=defaultdict(dict);cv=jain=0.;previous=0.;cv_integral=jain_integral=0.;peak=0.
    for when,changes in sorted(events.items()):
        dt=when-previous;cv_integral+=cv*dt;jain_integral+=jain*dt;previous=when
        for sign,(node,aid,inc) in changes:
            if sign==1:live[node][aid]=inc
            else:
                del live[node][aid]
                if not live[node]:del live[node]
        loads=[math.fsum(tasks.values()) for tasks in live.values()]
        total=math.fsum(loads);sq=math.fsum(u*u for u in loads);mean=total/n
        if mean>1e-15:
            cv=math.sqrt(max(0.,sq/n-mean*mean))/mean;jain=total*total/(n*sq)
        else:cv=jain=0.
        peak=max(peak,max(loads,default=0.))
        assert cv<=math.sqrt(n-1)+1e-7
        assert max(loads,default=0.)<=1+1e-7
    assert previous<=3600 and not live
    return dict(node_cv=cv_integral/3600,node_jain=jain_integral/3600,node_peak=peak*100,
                theoretical_cv_upper_bound=math.sqrt(n-1))


def review_node_statistics(data):
    audit=[]
    for phase in ('E4','E5','E6','E7','E8'):
        for s in data[phase]:
            if s['biz']!='ALL':continue
            stable=stable_node_moments(ROOT/s['raw_path'],s['scenario'])
            old_cv=s['node_cv'];old_jain=s['node_jain']
            assert abs(stable['node_peak']-s['node_peak'])<1e-6
            s.update(node_cv_raw=old_cv,node_jain_raw=old_jain,node_cv=stable['node_cv'],node_jain=stable['node_jain'],node_cv_audit='reconstructed_from_live_allocations')
            audit.append(dict(phase=phase,scenario=s['scenario'],strategy=s['strategy'],
                raw_node_cv=old_cv,reviewed_node_cv=s['node_cv'],raw_node_jain=old_jain,reviewed_node_jain=s['node_jain'],
                theoretical_cv_upper_bound=stable['theoretical_cv_upper_bound'],
                raw_exceeds_theoretical_bound=old_cv>stable['theoretical_cv_upper_bound'],
                node_peak_matches_original='PASS',method='direct live allocation sums, including idle nodes'))
    return audit


def enrich_single(s,path):
    gen=rows(path/'generated_tasks.csv');routes=rows(path/'routing_allocations.csv');kpi=js(path/'kpi.json');slots=rows(path/'slot_stats.csv')
    assert len(gen)==s['generated'] and len(routes)==s['accepted']
    assert all(r['event_type']=='INITIAL' for r in routes)
    assert all(abs(r['duration_s']-30)<1e-8 for r in gen)
    s.update(T=len(slots)*30.,ec=kpi['compute_energy_J'],et=kpi['tx_energy_J'],energy=kpi['total_energy_J'],
        segments=len(routes),reroute=kpi['reroutes'],migrations=kpi['migrations'],preemptions=kpi['preemptions'],
        tc=kpi['avg_compute_latency_ms'],loss=max(r['max_loss'] for r in routes),
        source_pct=100*sum(r['src']==r['compute_node'] for r in routes)/len(routes),
        ground_pct=100*sum(r['compute_node_type'] in ('EDGE','DC') for r in routes)/len(routes),
        pre_hops=kpi['avg_pre_hops'],post_hops=kpi['avg_post_hops'])
    index={r['task_id']:r for r in gen}
    s['gbits']=sum(index[r['task_id']]['data_volume_bits'] for r in routes)/1e9
    s['egbit']=s['energy']/s['gbits'];s['service_s']=s['accepted']*30
    s['etask']=s['energy']/s['accepted']
    s['pmean']=s['energy']/s['T'];s['pcmean']=s['ec']/s['T'];s['ptmean']=s['et']/s['T']
    power=defaultdict(lambda:[0.,0.])
    for a in routes:power[a['t_alloc']][0]+=a['power_used_W'];power[a['t_alloc']][1]+=a['energy_tx_J']/30
    s.update(pcpeak=max(p[0] for p in power.values()),ptpeak=max(p[1] for p in power.values()),ppeak=max(sum(p) for p in power.values()))
    # Old low-load snapshots did not save a config; preserve that limitation.
    s['config_saved']=(path/'effective_config.json').exists()
    s['raw_path']=str(path.relative_to(ROOT));return s


def datasets():
    phases={}
    aliases1=dict(generated='arrived',accepted='accepted',accept_rate='success_rate',success_rate='success_rate',dropped='dropped',
        delay='avg_delay_ms',p95='p95_delay_ms',dmax='max_delay_ms',tc='avg_compute_delay_ms',peak_active='max_concurrency',
        link_peak='peak_link_util_pct',node_peak='peak_node_util_pct',etask='avg_energy_per_task_J',solve99='solve_p99_ms')
    e1=[]
    for a in rows(ROOT/'results'/'single_biz_redefined_summary_20261016.csv'):
        s=dict(phase='E1',scenario=int(a['scenario']),strategy='latency',biz=a['biz'],mode='slot',completed=None,end_active=None)
        s.update({k:val(a,v) for k,v in aliases1.items()})
        sc=s['scenario'];b=s['biz'];path=ROOT/'results'/f'single_biz_redefined_sc{sc}_20261016'/b
        slots=rows(path/'slot_stats.csv');s.update(avg_active=float(np.mean([v['active'] for v in slots])),min_active=min(v['active'] for v in slots))
        s.update(rcap=0,rnone=s['generated']-s['accepted'])
        e1.append(enrich_single(s,path))
    phases['E1']=e1
    e2=[]
    aliases2=dict(generated='arrived',accepted='accepted',accept_rate='request_accept_rate',success_rate='request_accept_rate',dropped='dropped',
        delay='avg_delay_ms',p95='p95_delay_ms',dmax='max_delay_ms',peak_active='max_active',avg_active='avg_active',min_active='min_active',
        link_peak='peak_link_util_pct',node_peak='peak_node_util_pct',etask='avg_energy_per_accepted_task_J',solve99='solve_p99_ms',
        rcap='rejected_cap',rnone='rejected_no_feasible')
    for a in rows(ROOT/'results'/'concurrency_stress_summary_20261016.csv'):
        s=dict(phase='E2',scenario=int(a['scenario']),strategy='latency',biz=a['biz'],mode='slot',completed=None,end_active=None)
        s.update({k:val(a,v) for k,v in aliases2.items()})
        path=ROOT/'results'/f"single_biz_concurrency_stress_sc{s['scenario']}_20261016"/s['biz'];e2.append(enrich_single(s,path))
    phases['E2']=e2
    e3=[]
    aliases3=dict(generated='arrived',accepted='accepted',accept_rate='request_success_rate',success_rate='request_success_rate',
        dropped='dropped',rcap='rejected_cap',rnone='rejected_no_feasible',delay='delay_mean_ms',p95='delay_p95_ms',dmax='delay_max_ms',
        link_peak='peak_link_util_pct',link_cv='mean_link_cv',link_jain='mean_link_jain',node_peak='peak_node_util_pct',node_cv='mean_node_cv',
        min_active='min_active',avg_active='avg_active',peak_active='max_active',solve99='solve_p99_ms',fallback='fallback_attempts',recovered='fallback_accepted',
        common_n='common_accepted_tasks',common_e='common_energy_mean_J',common_d='common_delay_mean_ms',sha='input_sha256')
    for a in rows(SINGLE/'all_metrics.csv'):
        s=dict(phase='E3',scenario=int(a['scenario']),strategy=a['strategy'],biz=a['biz'],mode='slot',completed=None,end_active=None)
        s.update({k:val(a,v) for k,v in aliases3.items()})
        path=SINGLE/s['strategy']/f"sc{s['scenario']}"/s['biz'];e3.append(enrich_single(s,path))
    phases['E3']=e3
    dense=rows(DENSE/'mixed_all_metrics.csv')
    for phase,variant in [('E4','fixed30'),('E5','variable_random')]:
        phases[phase]=[mixed_normalize(a,phase) for a in dense if a['variant']==variant]
        for sc in (1,2,3):
            path=DENSE/variant/f'sc{sc}';routes=rows(path/'routing_allocations.csv')
            powers=defaultdict(lambda:defaultdict(lambda:[0.,0.]))
            for a in routes:
                dt=a['service_duration_s']
                if dt<=1e-12:continue
                for b in ('ALL',a['biz']):
                    for when,sign in ((a['t_alloc'],1),(a['service_t_end'],-1)):
                        powers[b][when][0]+=sign*a['power_used_W'];powers[b][when][1]+=sign*a['energy_tx_J']/dt
            for b,events in powers.items():
                c=t=pc=pt=total=0.
                for when,(dc,dt) in sorted(events.items()):
                    c+=dc;t+=dt;pc=max(pc,c);pt=max(pt,t);total=max(total,c+t)
                s=next(s for s in phases[phase] if s['scenario']==sc and s['biz']==b)
                s.update(pcpeak=pc,ptpeak=pt,ppeak=total,raw_path=str(path.relative_to(ROOT)))
    for phase,path in [('E6',OTHER),('E7',BOUNDED),('E8',FRESH)]:
        phases[phase]=[mixed_normalize(a,phase) for a in rows(path/'mixed_all_metrics.csv')]
        for s in phases[phase]:
            raw_path=(SPARSE/'variable_random' if phase=='E6' and s['strategy']=='latency' else path/s['strategy'])/f"sc{s['scenario']}"
            s['raw_path']=str(raw_path.relative_to(ROOT))
    return phases


class Report:
    def __init__(self,out):
        self.out=out;self.doc=Document();configure_document(self.doc)
        self.doc.sections[0].header.paragraphs[0].text=NAME+' | 2026-10-16'
        self.doc.sections[0].header.paragraphs[0].runs[0].font.size=Pt(8)
        self.doc.styles['Heading 1'].paragraph_format.page_break_before=True
        self.doc.styles['Normal'].paragraph_format.line_spacing=1.18
        self.figures=[];self.formulas=[];self.figdir=out/'figures';self.figdir.mkdir(exist_ok=True)
        self.architecture=self.figdir/'framework.png';self.timeline=self.figdir/'event_timeline.png'
    def h1(self,t):self.doc.add_heading(t,1)
    def h2(self,t):self.doc.add_heading(t,2)
    def h3(self,t):self.doc.add_heading(t,3)
    def p(self,t):return self.doc.add_paragraph(t)
    def t(self,headers,data):return table(self.doc,headers,data,size=8.4)
    def code(self,t):
        p=self.p(t)
        for run in p.runs:run.font.name='Consolas';run.font.size=Pt(8.5)
    def eq(self,label,text,source):
        n=len(self.formulas)+1
        pieces=[]
        for part in text.split('；'):
            while len(part)>86:
                cut=max(part.rfind(symbol,36,82) for symbol in ('＋','+','，','；','；',' '))
                if cut<36:cut=78
                pieces.append(part[:cut]);part=part[cut:]
            pieces.append(part)
        for index,piece in enumerate(pieces):
            p=self.doc.add_paragraph();p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.keep_with_next=True
            om=OxmlElement('m:oMathPara');obj=OxmlElement('m:oMath');mr=OxmlElement('m:r')
            wr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:eastAsia'),'宋体');wr.append(fonts)
            sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');wr.append(sz);mr.append(wr)
            tx=OxmlElement('m:t');tx.text=piece+(f'    ({n})' if index==len(pieces)-1 else '')
            mr.append(tx);obj.append(mr);om.append(obj);p._p.append(om)
        cap=self.p(f'式({n}) {label}。对应：{source}')
        cap.paragraph_format.space_after=Pt(4)
        for run in cap.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('666666')
        self.formulas.append(dict(number=n,label=label,formula=text,source=source))
    def fig(self,path,title,phase,analyses,source='当前源码或相应阶段CSV；见数据索引'):
        path=Path(path);n=len(self.figures)+1
        self.doc.add_picture(str(path),width=Cm(17))
        self.doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p=self.p(f'图{n}  {title}（{phase}）')
        for run in p.runs:run.font.size=Pt(9);run.bold=True
        for index,a in enumerate(analyses):self.p(('读图与口径：' if index==0 else '结果与分析：' if index==1 else '边界与解释：')+a)
        p=self.p('数据来源：'+source)
        for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('666666')
        self.figures.append(dict(number=n,phase=phase,title=title,path=str(path.relative_to(self.out)) if path.is_relative_to(self.out) else str(path.relative_to(ROOT)),source=source,analysis='\n'.join(analyses)))


def diagrams(r):
    plt.rcParams.update({'font.sans-serif':['Microsoft YaHei','SimHei','DejaVu Sans'],'axes.unicode_minus':False,'font.size':10})
    fig,ax=plt.subplots(figsize=(9,5));ax.axis('off');ax.set_xlim(0,10);ax.set_ylim(0,6)
    items=[(1.1,5,'CLI / config.yaml'),(4.8,5,'Simulator\n传统片驱动 / 事件引擎'),(8.6,5,'topology / network\n星座、可见性、信道'),
           (1.1,3,'tasks / TaskGenerator\n端点、业务、时长、种子'),(4.8,3,'solvers S0–S5\n策略评分＋公共校验'),(8.6,3,'ResourceLedger\n带宽 / 算力 / 功率'),
           (1.1,1,'experiments\n配对/独立输入与审计'),(4.8,1,'Metrics\n任务、路由、资源、积分'),(8.6,1,'CSV / JSON\n图表 / Word / PDF')]
    for x,y,t in items:
        patch=FancyBboxPatch((x-1.05,y-.45),2.1,.9,boxstyle='round,pad=.08',facecolor='#e6eef6',edgecolor='#2864A0');ax.add_patch(patch);ax.text(x,y,t,ha='center',va='center',fontsize=9)
    for a,b in [(0,1),(2,1),(3,4),(1,4),(4,5),(5,7),(6,7),(7,8)]:
        x1,y1,_=items[a];x2,y2,_=items[b];ax.annotate('',(x2,y2-.55 if y2>y1 else y2+.55 if y2<y1 else y2),(x1,y1),arrowprops=dict(arrowstyle='->',color='#52616d',lw=1.2,shrinkA=38,shrinkB=10))
    fig.tight_layout();fig.savefig(r.architecture,dpi=210);plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,3.1));ax.set_xlim(0,65);ax.set_ylim(0,4);ax.set_yticks([1,2,3],['业务B','业务A','拓扑']);ax.set_xlabel('仿真时间 s')
    ax.broken_barh([(0,30),(30,30)],(2.75,.45),facecolors=['#abc8e3','#82a7c7']);ax.text(15,3,'快照0：片内冻结',ha='center',fontsize=9);ax.text(45,3,'快照1：重新构建',ha='center',fontsize=9)
    ax.broken_barh([(8,22),(30,18)],(1.75,.45),facecolors=['#d88722','#d88722']);ax.text(9,2.3,'真实到达',fontsize=9);ax.text(32,2.3,'重算后继续服务',fontsize=9);ax.text(49,2,'完成释放',fontsize=9)
    ax.broken_barh([(37,19)],(.75,.45),facecolors='#27856B');ax.text(37,1.3,'片内新业务',fontsize=9)
    ax.axvline(30,color='#555',ls='--');ax.text(30,3.6,'边界：对存续业务重算',ha='center',fontsize=9)
    ax.set_xticks([0,8,30,37,48,56,60]);ax.grid(axis='x',alpha=.2);ax.spines[['top','right']].set_visible(False)
    fig.tight_layout();fig.savefig(r.timeline,dpi=210);plt.close(fig)


def grouped(r,name,items,keys,labels,title,by='scenario',log=False):
    fig,axes=plt.subplots(1,len(keys),figsize=(8.6,3.5),squeeze=False);axes=axes[0];x=np.arange(len(BIZ));width=.25
    for ax,key,ylab in zip(axes,keys,labels):
        for j,sc in enumerate((1,2,3)):
            selected={s['biz']:s for s in items if s['scenario']==sc}
            ax.bar(x+(j-1)*width,[selected[b][key] for b in BIZ],width,label=f'场景{sc}')
        ax.set_xticks(x,BIZ);ax.set_ylabel(ylab);ax.grid(axis='y',alpha=.2)
        if log:ax.set_yscale('log')
    fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',ncol=3,fontsize=9);fig.suptitle(title)
    fig.tight_layout(rect=(0,.1,1,.94));path=r.figdir/(name+'.png');fig.savefig(path,dpi=210);plt.close(fig);return path


def metric_overview(r,items,name,title):
    stats=[s for s in items if s['biz']=='ALL'];fig,axes=plt.subplots(1,3,figsize=(8.6,3.5));xx=np.arange(3)
    for ax,key,lab in zip(axes,('delay','egbit','link_peak'),('生命周期时延均值 ms','单位输入能耗 J/Gbit','链路峰值 %')):
        for j,st in enumerate(STRATEGIES):
            selected={s['scenario']:s for s in stats if s['strategy']==st}
            if not selected:continue
            ax.bar(xx+(j-1)*.25,[selected[i][key] for i in (1,2,3)],.25,label=LABEL[st],color=COLOR[st])
        ax.set_xticks(xx,['场景1','场景2','场景3']);ax.set_ylabel(lab);ax.grid(axis='y',alpha=.2)
    fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',ncol=3);fig.suptitle(title);fig.tight_layout(rect=(0,.1,1,.94))
    path=r.figdir/(name+'.png');fig.savefig(path,dpi=210);plt.close(fig);return path


def concurrency_plot(r,groups,name,title):
    fig,axes=plt.subplots(1,3,figsize=(8.6,3.5))
    for sc,ax in enumerate(axes,1):
        for label,path,color in groups(sc):
            data=rows(path/'event_timeline.csv');ax.step([a['t']/60 for a in data],[a['active'] for a in data],where='post',label=label,color=color,lw=1)
        ax.axhline(CAP[sc],color='#888',ls=':',lw=.7);ax.set_title(f'场景{sc}');ax.set_xlabel('min');ax.set_ylabel('实际并发');ax.grid(alpha=.2)
    fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',ncol=3);fig.suptitle(title);fig.tight_layout(rect=(0,.1,1,.94))
    path=r.figdir/(name+'.png');fig.savefig(path,dpi=210);plt.close(fig);return path


def outcome_plot(r,items,name,title):
    stats=[s for s in items if s['biz']=='ALL'];fig,ax=plt.subplots(figsize=(8.6,3.3));x=np.arange(len(stats));bottom=np.zeros(len(stats))
    for key,label,color in [('completed','完整完成','#27856B'),('end_active','期末在服','#8cbcc5'),('dropped','掉线','#b34c40'),('rcap','上限拒绝','#d88722'),('rnone','不可行拒绝','#777')]:
        vals=[100*float(s[key])/s['generated'] for s in stats];ax.bar(x,vals,bottom=bottom,label=label,color=color);bottom+=vals
    ax.set_xticks(x,[f"{s['scenario']}-{LABEL[s['strategy']]}" for s in stats],rotation=30);ax.set_ylim(0,102);ax.set_ylabel('占生成请求 %');ax.legend(ncol=3,fontsize=8);ax.set_title(title)
    fig.tight_layout();path=r.figdir/(name+'.png');fig.savefig(path,dpi=210);plt.close(fig);return path


def source_chart(r,root,name,prefix):
    src=track(root/'figures'/name);target=r.figdir/f'{prefix}_{name}';shutil.copy2(src,target);return target


def single_tables(r,items,label):
    r.t(['场景/业务','接纳/生成','请求成功/已接纳QoS','平均/峰值并发','时延均值/P95/最大ms','链路/算力峰值%'],[
        [f"{s['scenario']}/{s['biz']}",f"{integer(s['accepted'])}/{integer(s['generated'])}",pct(s['accept_rate'])+' / '+pct(s['admitted_qos_rate']),trip(s,'avg_active','peak_active',d=2),trip(s,'delay','p95','dmax'),trip(s,'link_peak','node_peak',d=4)] for s in items])
    r.t(['场景/业务','计算/传输能耗J','J/Gbit','平均/峰值动态W','处理ms/最大丢包','求解P99ms'],[
        [f"{s['scenario']}/{s['biz']}",trip(s,'ec','et'),fmt(s['egbit']),trip(s,'pmean','ppeak'),f"{fmt(s['tc'],5)}/{s['loss']:.2e}",fmt(s['solve99'])] for s in items])
    for sc in (1,2,3):
        sub=[s for s in items if s['scenario']==sc]
        r.h3(f'{label} 场景{sc}逐业务指标分析')
        for s in sub:
            r.p(f"{s['biz']}：接纳{integer(s['accepted'])}/{integer(s['generated'])}（{pct(s['accept_rate'])}），容量/不可行拒绝{integer(s['rcap'])}/{integer(s['rnone'])}，掉线{integer(s['dropped'])}；"
                f"最低片首并发{integer(s.get('min_active'))}，平均/峰值并发{fmt(s['avg_active'],2)}/{integer(s['peak_active'])}，占上限{100*s['avg_active']/CAP[sc]:.2f}%/{100*s['peak_active']/CAP[sc]:.2f}%。"
                f"时延均值/P95/最大{trip(s,'delay','p95','dmax')} ms，相对预算的最小余量{({'LL':20,'HR':200,'HC':200,'MC':200}[s['biz']]-s['dmax']):.3f} ms；"
                f"处理项{s['tc']:.5f} ms、最大路径丢包{s['loss']:.2e}。均值反映典型样本，P95与最大值反映尾部风险，不能只看均值。")
            r.p(f"资源与能量：链路/单节点峰值{trip(s,'link_peak','node_peak',d=4)}%，计算/传输能量{trip(s,'ec','et')} J，传输占{100*s['et']/s['energy']:.2f}%；"
                f"每接纳任务{s['etask']:.5f} J、单位输入{s['egbit']:.4f} J/Gbit，实际模型输入{s['gbits']:.3f} Gbit。平均/峰值动态功率{trip(s,'pmean','ppeak')} W，"
                f"计算/传输平均W为{trip(s,'pcmean','ptmean')}、各自峰值W为{trip(s,'pcpeak','ptpeak')}；源端/地面处理段比例{trip(s,'source_pct','ground_pct',d=2)}%，前段/后段平均跳数{trip(s,'pre_hops','post_hops',d=2)}，求解P99 {s['solve99']:.3f} ms。"
                '不同业务速率差异大，绝对J不能独立代表策略效率；求解时间为当时机器墙钟观测。')


def single_strategy_tables(r,items):
    for sc in (1,2,3):
        for biz in BIZ:
            selected=[s for s in items if s['scenario']==sc and s['biz']==biz];l=next(s for s in selected if s['strategy']=='latency')
            r.h3(f'场景{sc}／{biz}：三策略配对结果')
            r.t(['策略','接纳/生成','上限/不可行拒绝','时延均值/P95/最大ms','处理ms/损失','均值/峰值并发'],[
                [LABEL[s['strategy']],f"{integer(s['accepted'])}/{integer(s['generated'])}",f"{integer(s['rcap'])}/{integer(s['rnone'])}",trip(s,'delay','p95','dmax'),f"{fmt(s['tc'],5)}/{s['loss']:.2e}",trip(s,'avg_active','peak_active',d=2)] for s in selected])
            r.t(['策略','总/计算/传输能耗J','J/Gbit','配对任务J/时延ms','源端/地面%','前/后段跳数'],[
                [LABEL[s['strategy']],trip(s,'energy','ec','et'),fmt(s['egbit']),trip(s,'common_e','common_d'),trip(s,'source_pct','ground_pct',d=2),trip(s,'pre_hops','post_hops',d=2)] for s in selected])
            r.t(['策略','链路峰值%/CV/Jain','节点峰值%/CV','平均/峰值W','求解P99ms','补充搜索/成功'],[
                [LABEL[s['strategy']],trip(s,'link_peak','link_cv','link_jain',d=4),trip(s,'node_peak','node_cv',d=4),trip(s,'pmean','ppeak'),fmt(s['solve99']),f"{integer(s.get('fallback'))}/{integer(s.get('recovered'))}"] for s in selected])
            for s in selected:
                r.p(f"{LABEL[s['strategy']]}：请求成功率{pct(s['accept_rate'])}、已接纳QoS率{pct(s['admitted_qos_rate'])}，相对{20 if biz=='LL' else 200} ms预算最大时延余量"
                    f"{(20 if biz=='LL' else 200)-s['dmax']:.3f} ms；共同成功任务{integer(s['common_n'])}个，配对平均能耗相对时延策略变化{relative(s['common_e'],l['common_e']):+.2f}%，"
                    f"配对时延变化{relative(s['common_d'],l['common_d']):+.2f}%。链路峰值/CV与节点峰值/CV分别为{trip(s,'link_peak','link_cv',d=4)}和{trip(s,'node_peak','node_cv',d=4)}，"
                    f"链路Jain {fmt(s['link_jain'],4)}。源端/地面比例{trip(s,'source_pct','ground_pct',d=2)}%，"
                    f"计算能量占总量{100*s['ec']/s['energy']:.3f}%，平均/峰值总W为{trip(s,'pmean','ppeak')}，计算/传输平均W为{trip(s,'pcmean','ptmean')}、各自峰值W为{trip(s,'pcpeak','ptpeak')}；求解P99 {fmt(s['solve99'])} ms，"
                    f"补充搜索{integer(s.get('fallback'))}次、恢复{integer(s.get('recovered'))}次。资源、QoS与求解代价应同时评估，不以单指标宣称全优。")
            r.p('本组为单片30 s任务，路由段数等于接纳任务数，重算、迁移、抢占和跨片掉线均为0；这不是算法跳过新片计算，而是旧任务在下一片前已经结束。'
                '表中计算/传输能量分别对应处理与数据搬运；单任务J、J/Gbit和配对J的分母不同，配对差值只在共同成功任务集合上解释。节点CV未与链路Jain强行合并为一个指标。')


def mixed_tables(r,items,phase,detail=True):
    overall=[s for s in items if s['biz']=='ALL']
    r.t(['场景/策略','接纳/生成','完成/在服/掉线','接纳/完整完成率','容量/不可行拒绝','均值/峰值并发'],[
        [marker(s),f"{integer(s['accepted'])}/{integer(s['generated'])}",f"{integer(s['completed'])}/{integer(s['end_active'])}/{integer(s['dropped'])}",f"{pct(s['accept_rate'])}/{pct(s['complete_rate'])}",f"{integer(s['rcap'])}/{integer(s['rnone'])}",trip(s,'avg_active','peak_active',d=2)] for s in overall])
    r.t(['场景/策略','时延均值/P95/最大ms','总/计算/传输kJ','J/Gbit','平均/峰值W','链路峰值%/CV'],[
        [marker(s),trip(s,'delay','p95','dmax'),'/'.join(fmt(s[k]/1000) for k in ('energy','ec','et')),fmt(s['egbit']),trip(s,'pmean','ppeak'),trip(s,'link_peak','link_cv')] for s in overall])
    r.t(['场景/策略','节点峰值%/CV','链路/节点Jain','计算/传输峰值W','路由段/重算/迁移','满并发秒/首达秒','求解P99ms/搜索恢复'],[
        [marker(s),trip(s,'node_peak','node_cv',d=4),trip(s,'link_jain','node_jain',d=4),trip(s,'pcpeak','ptpeak'),
         f"{integer(s['segments'])}/{integer(s['reroute'])}/{integer(s['migrations'])}",trip(s,'cap_s','first_cap'),f"{fmt(s['solve99'])}/{integer(s.get('recovered'))}"] for s in overall])
    r.p('“—”表示该阶段未保存或未独立审计该量，不补0、不伪造历史数据。全网链路/节点分布指标只在ALL行报告，不拆成没有可定义分母的“某业务全网CV”。本报告可由保存的服务段补算动态功率，但不会给没有原始事件证据的早期表捏造满并发持续秒数。')
    for s in overall:
        r.h3(marker(s)+'：总体指标说明')
        r.p(f"任务与服务：生成{s['generated']:.0f}、接纳{s['accepted']:.0f}，完整完成{s['completed']:.0f}、期末在服{s['end_active']:.0f}、掉线{s['dropped']:.0f}，"
            f"容量/不可行拒绝{integer(s['rcap'])}/{integer(s['rnone'])}；接纳率{pct(s['accept_rate'])}、完整完成率{pct(s['complete_rate'])}、存续率{pct(s['success_rate'])}，"
            f"已接纳存续率{pct(s['survival_admitted'])}、已服务QoS通过率{pct(s['admitted_qos_rate'])}。累计服务{s['service_s']:.3f}秒，对应平均并发{s['avg_active']:.3f}，峰值{s['peak_active']:.0f}；"
            f"满并发{fmt(s.get('cap_s'))}秒、占比{fmt(s.get('cap_pct'))}%，≥90%上限{fmt(s.get('near_s'))}秒、首达{fmt(s.get('first_cap'))}秒。完整完成与“完成＋期末在服”分母相同但含义不同。")
        r.p(f"时延与能量：生命周期最坏时延的任务均值/P95/最大{trip(s,'delay','p95','dmax')} ms，服务加权时延{fmt(s['weighted_delay'])} ms，处理均值{fmt(s['tc'],5)} ms。"
            f"原始实服输入{s['gbits']:.3f} Gbit，计算/传输能量{fmt(s['ec']/1000)}/{fmt(s['et']/1000)} kJ，传输占{100*s['et']/s['energy']:.2f}%；"
            f"每初始接纳任务{s['etask']:.3f} J、归一化{s['egbit']:.3f} J/Gbit。平均计算/传输/总W为{trip(s,'pcmean','ptmean','pmean')}，"
            f"各项峰值计算/传输/同刻总W为{trip(s,'pcpeak','ptpeak','ppeak')}，总峰值不按各自峰值相加。")
        r.p(f"资源与代价：链路峰值{fmt(s['link_peak'])}%、CV {fmt(s['link_cv'])}、Jain {fmt(s['link_jain'],4)}；节点峰值{fmt(s['node_peak'],4)}%、CV {fmt(s['node_cv'])}、Jain {fmt(s['node_jain'],4)}。"
            f"源端/地面计算段比例{trip(s,'source_pct','ground_pct',d=2)}%，路由段/重算/迁移{integer(s['segments'])}/{integer(s['reroute'])}/{integer(s['migrations'])}。"
            f"求解P99 {fmt(s['solve99'])} ms，整组耗时{fmt(s.get('wall'),1)} s，补充搜索/成功{integer(s.get('fallback'))}/{integer(s.get('recovered'))}；"
            '这些是当前残余资源下启发式的实测，不是全局最优证明或无系统负载影响的性能基准。')
    if not detail:return
    r.h2('分业务完整结果及逐项解释')
    for s in overall:
        selected=[a for a in items if a['scenario']==s['scenario'] and a['strategy']==s['strategy'] and a['biz']!='ALL']
        r.h3(marker(s)+'／四业务')
        r.t(['业务','完成/在服/生成','时延均值/P95/最大ms','处理ms/最大丢包','J/Gbit','均值/峰值活跃'],[
            [a['biz'],f"{integer(a['completed'])}/{integer(a['end_active'])}/{integer(a['generated'])}",trip(a,'delay','p95','dmax'),f"{fmt(a['tc'],5)}/{a['loss']:.2e}",fmt(a['egbit']),trip(a,'avg_active','peak_active',d=2)] for a in selected])
        r.t(['业务','计算/传输kJ','输入Gbit/服务秒','源端/地面%','路由段/重算/迁移','容量/不可行/掉线'],[
            [a['biz'],f"{fmt(a['ec']/1000,4)}/{fmt(a['et']/1000,4)}",trip(a,'gbits','service_s'),trip(a,'source_pct','ground_pct',d=2),
             f"{integer(a['segments'])}/{integer(a['reroute'])}/{integer(a['migrations'])}",f"{integer(a['rcap'])}/{integer(a['rnone'])}/{integer(a['dropped'])}"] for a in selected])
        for a in selected:
            budget=20 if a['biz']=='LL' else 200;loss_budget=.0002 if a['biz']=='HR' else .01
            r.p(f"{a['biz']}：接纳率{pct(a['accept_rate'])}、完整完成率{pct(a['complete_rate'])}，最大时延距{budget} ms预算还有{budget-a['dmax']:.3f} ms；"
                f"P95 {a['p95']:.3f} ms、服务加权时延{a['weighted_delay']:.3f} ms，最大路径丢包{a['loss']:.2e}相对预算{loss_budget:.2e}。"
                f"能量占该组总体的{100*a['energy']/s['energy']:.3f}%，其中传输占自身能量{100*a['et']/a['energy']:.3f}%；"
                f"平均/峰值动态W为{trip(a,'pmean','ppeak')}，每初始接纳任务{a['etask']:.4f} J。"
                f"源端/地面处理{trip(a,'source_pct','ground_pct',d=2)}%，迁移{integer(a['migrations'])}次；重算并不要求更换计算点。"
                '时延样本只含已服务任务，失败率与时延应分开读；输入bit与服务秒数是能量和并发的统计分母。')


def figure_single(r,items):
    definitions=[('01_delay.png','配对任务时延','common_d','ms'),('02_energy_ratio.png','配对任务相对能耗','common_e','J'),
        ('03_link_cv.png','链路利用率CV','link_cv','CV'),('04_success.png','请求成功率','accept_rate','比率'),('05_solver.png','求解P99','solve99','ms')]
    for file,title,key,unit in definitions:
        lines=[]
        for sc in (1,2,3):
            sub=[s for s in items if s['scenario']==sc];lo=min(sub,key=lambda s:s[key]);hi=max(sub,key=lambda s:s[key])
            lines.append(f"场景{sc}范围{lo[key]:.4g}–{hi[key]:.4g} {unit}（最低{lo['biz']}/{LABEL[lo['strategy']]}，最高{hi['biz']}/{LABEL[hi['strategy']]}）")
        note={'common_d':'按共同成功任务对比；不是生命周期跨片最大时延。MC按最坏目的统计，仍可能与S1即时后段之和排序不同。',
              'common_e':'图中按每场景/业务时延策略归一化为1；表中给实际J。共同成功交集控制拒绝样本偏差，但单种子不证明显著性。',
              'link_cv':'含空闲链路；CV与最大链路峰值是不同目标，平方增量也不承诺每个CV均最低。',
              'accept_rate':'纵轴截断在接近100%的区间，用来辨认小差异；不能忽略主动超额请求和传播不可行拒绝。',
              'solve99':'为墙钟P99，含并行负载影响，不计入网络时延，不能用该单次观测代替严格复杂度分析。'}[key]
        r.fig(source_chart(r,SINGLE,file,'E3'),title,'E3',['横向为四业务，分面为三场景，三色为同输入三策略。','；'.join(lines)+'。',note],str((SINGLE/'all_metrics.csv').relative_to(ROOT)))
    for file,title,analysis in [
        ('06_hc_link_timeseries.png','HC最忙链路随时间',f"场景2的三策略峰值依次为"+'/'.join(fmt(next(s for s in items if s['scenario']==2 and s['biz']=='HC' and s['strategy']==st)['link_peak']) for st in STRATEGIES)+'%。峰值未必恰好出现在150 s并发最高片，因为端点、速率与快照也变化。'),
        ('07_placement_compute.png','计算放置与单节点算力峰值','热力图逐格给出源端计算比例及节点利用率。源端早压缩往往降低传输，算力峰值仍远小于100%；对比图须注意两幅色标的单位与数量级不同。'),
        ('08_concurrency.png','MC并发压力轨迹','常规片约60%上限，150 s为超额峰值片；场景3存在少量拒绝时曲线可不完全重合，但三个场景均实际触及各自硬上限。')]:
        r.fig(source_chart(r,SINGLE,file,'E3'),title,'E3',['图中位置/曲线对应三场景三策略；这是单片30 s业务，不是长会话轨迹。',analysis,'源端比例按本轮初始路由计数；节点峰值是单节点占用，不是全网算力已经饱和。'],str((SINGLE/'all_slot_metrics.csv').relative_to(ROOT)))


def mixed_figures(r,items,root,phase):
    overall=[s for s in items if s['biz']=='ALL']
    defs=[('06_concurrency.png','实际并发轨迹',('avg_active','peak_active'),'平均/峰值任务'),
          ('01_delay.png','分业务生命周期时延',('delay','p95'),'ms'),('02_energy.png','分业务单位输入能耗',('egbit',),'J/Gbit'),
          ('03_load.png','链路热点与CV',('link_peak','link_cv'),'百分比/CV'),('04_power.png','全网动态平均与峰值功率',('pmean','ppeak'),'W'),
          ('05_placement.png','源端处理与迁移',('source_pct','migrations'),'百分比/次数'),('07_solver.png','单次求解P99',('solve99',),'ms')]
    for file,title,keys,unit in defs:
        summaries=[]
        for sc in (1,2,3):
            group=[s for s in overall if s['scenario']==sc]
            summaries.append(f"场景{sc}按时延/均衡/能耗："+'；'.join('/'.join(fmt(s[k],2) for k in keys) for s in group))
        notes={'06_concurrency.png':'同输入的E7三策略曲线重合；E8每策略输入不同，曲线差异含业务效应。3600 s归零来自生成时截止规则。',
               '01_delay.png':'图按业务分组，段内最坏目的再跨片取最大；上文表和本图不可误读为平均会话持续时间。',
               '02_energy.png':'对数纵轴比较量级；HC原始输入大，主导总能量，低码率业务J/Gbit也可能较高。',
               '03_load.png':'最忙链路峰值低不保证CV低，CV低不代表有更大容量；E8场景2实际出现不同排序。',
               '04_power.png':'平均W＝J/3600，峰值按同一时刻求和；不是平台静态功耗或射频铭牌功率。',
               '05_placement.png':'按路由段计源端比例，迁移只指计算节点变化；低迁移不等于不逐片重算。',
               '07_solver.png':'双层千星的候选搜索与可行性补充可能增加开销；并行机器负载也影响墙钟数字。'}[file]
        if file in ('01_delay.png','02_energy.png'):
            extra=[]
            for biz in BIZ:
                sub=[s for s in items if s['biz']==biz]
                extra.append(f"{biz}范围{min(s[keys[0]] for s in sub):.3f}–{max(s[keys[0]] for s in sub):.3f}")
            summaries.append('分业务'+'；'.join(extra))
        r.fig(source_chart(r,root,file,phase),title,phase,[f'指标为{unit}；三色依次为时延、均衡、能耗。','。'.join(summaries)+'。',notes+(' 本阶段为非配对输入，差值不能直接归因于算法。' if phase=='E8' else ' 本阶段已核验逐任务输入相同，可作本批策略对照。')],str((root/'mixed_all_metrics.csv').relative_to(ROOT)))


def build_experiments(r,data):
    r.h1('8. 实验体系、时间顺序与统一指标字典')
    catalog=[('E1','单业务低负载功能验证',12,'3600 s，30 s业务','三场景四业务；仅S1时延；旧配置快照缺失'),
        ('E2','单业务并发压力验证',12,'300 s，10片','均值≥半上限且峰值触顶；仅S1时延'),
        ('E3','单业务三策略配对实验',36,'300 s，10片','每场景/业务三策略同任务；含算法改进前后'),
        ('E4','固定30 s混合业务参考',3,'3600 s，120片','3:2:3:1；仅S1时延；每10片一个压力峰'),
        ('E5','原时长高到达混合实验',3,'3600 s，120片','沿用高到达导致持续过载；未限制尾部结束'),
        ('E6','原时长低背景＋短峰修订',9,'3600 s，120片','3时延＋6其他策略；同场景配对；允许期末在服'),
        ('E7','全部业务窗口内截止',9,'3600 s，120片','9组重跑，尾部生成时收短；同场景配对，记A'),
        ('E8','变化业务独立随机复测',9,'3600 s，120片','场景和策略都独立生成新业务；非配对，记B')]
    r.t(['阶段','名称','结果组数','时间设计','比较口径'],catalog)
    r.p('顺序依据原Word及用户逐轮修订要求，不依文件最后修改时间猜测。E2的时延对照与E3中同条件时延结果属于复测/控制，不当作独立随机样本叠加；E6复用其已保存的三组时延参考，仅新增两策略六组。E7截止规则使完成标签变化，不自动改变窗口内承载与能量；E8澄清前配对参考排除在正式统计之外。')
    r.t(['指标族','定义与分析问题'],[
        ['任务结果','生成、接纳、完整完成、期末在服、掉线、两种拒绝；是过载、路径问题，还是观察截断？'],
        ['并发/服务','平均＝服务秒/T；峰值＝真实事件最大active；是否真的达到硬上限、持续多久？'],
        ['QoS','任务时延均值/P95/最大、服务加权时延、处理项、损失与预算余量；均值是否掩盖尾风险？'],
        ['能量/功率','计算/传输/总J、J/任务、J/Gbit、平均/峰值W；分母是否一致、是否同刻求和？'],
        ['负载分布','链路/节点峰值、CV与Jain；是否改善热点、是否只是资源池改变？'],
        ['路由结构','源端/地面计算比例、前/后跳数、分配段/重算/迁移；早压缩是否减少搬运？'],
        ['求解代价','单次P99、整组墙钟、补充搜索/恢复；效率收益是否伴随更高计算开销？'],
        ['审计元数据','种子、输入SHA-256、物理配置、出处；不作为性能指标参与排名']])
    r.p('下文按阶段给完整汇总、逐业务解释和每幅图的读法/观察/边界。底层字段及别名见metric_dictionary.csv，规范化实际数值见report_data.csv；历史未采集量标“—”。没有将当前Word中删除过的旧阶段默默抹去，但也不将旧场景版本的档案数字并入当前主结论。')
    r.p('统计复核说明：本次对33组事件实验的节点CV/Jain进行了独立重建，未发现超出非负n维负载理论上限√(n−1)的记录。'
        '路由服务段逐时刻重建活动申请、直接求和后积分的CV与原值最大差约2.9×10⁻¹¹，Jain差约6.3×10⁻¹⁵，属于浮点误差。'
        '原值保存在report_data的node_cv_raw/node_jain_raw与node_cv_recalculation.csv，正文使用等价复核值；没有修改任何原始CSV、路由、成功率或能量。')
    r.h1('9. E1：单业务低负载多场景功能验证')
    r.p('每次只启用LL/HR/HC/MC之一，30 s任务片首进入，运行3600 s。历史到达基准为8/6/2/1每片，沿用高峰倍率，并非当前混合权重3/2/3/1。证据来自12行摘要、各组原始任务/路由和120片统计；这轮缺少effective_config，不应声称使用今天的批量入口可以逐任务精确复现。')
    single_tables(r,data['E1'],'E1')
    for name,keys,labels,title in [('E1_delay',('delay','p95'),('平均时延 ms','P95 ms'),'低负载四业务时延'),('E1_resource',('link_peak','node_peak'),('链路峰值 %','节点峰值 %'),'低负载资源占用'),('E1_energy',('egbit','solve99'),('J/Gbit','求解P99 ms'),'低负载能量与求解开销')]:
        observations=[]
        for key,unit in zip(keys,labels):
            observations.append(unit+'范围'+f"{min(s[key] for s in data['E1']):.5g}–{max(s[key] for s in data['E1']):.5g}")
        r.fig(grouped(r,name,data['E1'],keys,labels,title),title,'E1',['每个业务下三根柱对应三个场景，各面板单位独立。',
            '；'.join(observations)+'。12组均接纳成功，但最大实际并发仅27，不能说明200/250/500上限下的容量；HC的高码率能制造局部热点。',
            '场景3的较低典型时延与更多LEO连接一致，但端点池与随机序列也不同；不能把跨场景差值全部归因于网络分层。'], 'results/single_biz_redefined_summary_20261016.csv及对应原始CSV')
    r.h1('10. E2：提高任务量后的单业务并发压测')
    r.p('每组300 s、10片，任务30 s；常规片请求为并发上限60%，第5号片为峰值：场景1为200、场景2为250、场景3注入525但硬上限仍500。运行12组S1时延实验。与E1相比同时改变了到达量和观测时长，不能把差值当成只增并发的严格单因素试验。')
    single_tables(r,data['E2'],'E2')
    for name,keys,labels,title in [('E2_concurrency',('avg_active','peak_active'),('平均并发','峰值并发'),'达到最大并发的证据'),('E2_margin',('dmax','link_peak'),('最大时延 ms','链路峰值 %'),'满载验证不等于充足QoS裕量')]:
        observations='；'.join(unit+'范围'+f"{min(s[key] for s in data['E2']):.5g}–{max(s[key] for s in data['E2']):.5g}" for key,unit in zip(keys,labels))
        r.fig(grouped(r,name,data['E2'],keys,labels,title),title,'E2',['四业务分别比较三场景，平均与峰值不能互换。',
            observations+'。各组平均实际并发约64%上限并均触顶；场景3 MC最大时延199.266 ms，仅余0.734 ms。场景2 HC链路99.8481%，而算力峰值远低于1%。',
            '25左右的峰值超额拒绝是设计负载的一部分；GEO相关不可行另列。当前未建显式排队，不以QoS达标宣称真实满链路也有同样时延。'],'results/concurrency_stress_summary_20261016.csv')
    r.h1('11. E3：单业务多场景三策略配对实验与算法修订')
    r.p('沿用E2压力形状，对三场景×四业务×三策略共36组实际运行。每个场景/业务的三策略逐任务输入与物理配置已通过哈希/字段校验。全部正式求解器为S1，不把它误写成S0–S5性能矩阵。比较先看接纳与QoS，再看共同成功任务能量/时延，最后看负载、功率、放置和求解开销。')
    before=rows(SINGLE/'algorithm_before_after.csv')
    r.h2('11.1 仅算法层面的修订及真实前后结果')
    r.p('负载目标改为本次分配后的平方利用率增量，候选最终评分聚合前后段重叠链路；均衡/能耗首选失败时增加低传播候选，仍按原目标排序。旧版对空闲低速任务的固定ε可能主导路径选择；单一能耗/负载最短路也可能违反时延预算，即使另有可行传播路径。没有通过扩容或放宽QoS解决。')
    r.t(['场景/业务/策略','旧/新接纳','旧/新不可行','旧/新总J','旧/新链路峰值%'],[
        [f"{int(a['scenario'])}/{a['biz']}/{LABEL[a['strategy']]}",f"{integer(a['old_accepted'])}/{integer(a['new_accepted'])}",f"{integer(a['old_no_feasible'])}/{integer(a['new_no_feasible'])}",f"{fmt(a['old_energy_J'])}/{fmt(a['new_energy_J'])}",f"{fmt(a['old_peak_link_pct'])}/{fmt(a['new_peak_link_pct'])}"] for a in before])
    for st in ('balance','energy'):
        a=[x for x in before if x['strategy']==st]
        r.p(f"{LABEL[st]}修订前后：12组不可行拒绝总数{sum(x['old_no_feasible'] for x in a):.0f}→{sum(x['new_no_feasible'] for x in a):.0f}，接纳{sum(x['old_accepted'] for x in a):.0f}→{sum(x['new_accepted'] for x in a):.0f}。"
            '总J可能同时受接纳集合变化影响，不能只用旧/新总J差宣称效率；对齐服务任务与归一化能耗更重要。旧源代码未独立封存时，当前摘要可作为结果证据，但不是可执行的旧版本回放包。')
    r.h2('11.2 三场景四业务完整指标与逐组比较')
    single_strategy_tables(r,data['E3'])
    r.h2('11.3 八幅原有对比图的逐图说明')
    figure_single(r,data['E3'])
    r.h1('12. E4：3600 s固定30 s混合业务参考')
    r.p('LL∶HR∶HC∶MC＝3∶2∶3∶1，每片都按9的倍数生成，全部持续30 s、片首进入。120片中每10片有一个压力峰；仅测试S1时延。这是混合业务基准，不是四组单业务结果相加，因为资源账本共享、同片按优先级共同分配。')
    mixed_tables(r,data['E4'],'E4')
    path=concurrency_plot(r,lambda sc:[('固定30s',DENSE/'fixed30'/f'sc{sc}',COLOR['latency'])],'E4_concurrency','固定30秒混合业务的周期峰值')
    r.fig(path,'固定30 s混合业务并发','E4',['横轴60分钟，每10片一个峰值，虚线为硬上限。','平均并发125.30/162.70/317.275，实际峰值200/250/500；所有已接纳任务在一个片内结束，因此无跨片重算。','场景3存在3个传播不可行请求，容量拒绝另列；拒绝并非都来自算力不足。'],str(DENSE.relative_to(ROOT))+'/fixed30/scN/event_timeline.csv')
    r.fig(metric_overview(r,data['E4'],'E4_overview','固定30秒混合业务的时延、能量与热点'),'固定混合业务关键指标','E4',[
        '三个面板分别是总体时延、J/Gbit、最忙链路百分比，不能将不同单位合成一个评分。',
        '三场景单位能耗31.097/51.645/27.880 J/Gbit，场景2链路峰值接近100%，体现高码率任务星地搬运的成本。',
        '总体指标混合四类业务，HC主导能量，LL时延仍应单列；跨场景不是同一端点样本。'],str(DENSE.relative_to(ROOT))+'/mixed_all_metrics.csv（fixed30）')
    r.h1('13. E5：原时长但高到达强度的混合实验')
    r.p('保持与固定30 s参考相同的高到达任务数量/属性配对基础，恢复业务原始时长并在片内随机到达；长会话跨越多个快照。此阶段没有生成时尾部截止，存在右删失。高到达不适合原时长，持续满并发导致大量scenario_concurrency_cap拒绝，不能简单归咎于路由算法。')
    r.eq('到达、时长与请求并发负载近似','N_offered≈λ·E[D]；E[D]=(3×180+2×360+3×1200+1×1200)/9=673.333 s','仅负载设计近似；实际请求峰值由事件时间精确计算')
    mixed_tables(r,data['E5'],'E5')
    r.fig(outcome_plot(r,data['E5'],'E5_outcomes','高到达原时长的请求结果组成'),'高到达导致的低接纳率','E5',[
        '每个场景的柱按生成请求分母分解为完成、在服、掉线、容量拒绝及不可行。',
        '容量拒绝13885/18027/35287，请求接纳仅约8%；实际平均并发198.50/248.19/496.44，几乎全程顶格。',
        '窗口结束仍有200/250/500在服；场景1/3还各有2/3个LL跨片掉线，需与容量拒绝分开诊断。'],str(DENSE.relative_to(ROOT))+'/mixed_all_metrics.csv（variable_random）')
    r.fig(concurrency_plot(r,lambda sc:[('30s参考',DENSE/'fixed30'/f'sc{sc}',COLOR['latency']),('原时长高到达',DENSE/'variable_random'/f'sc{sc}',COLOR['balance'])],'E5_compare','同到达强度：30s与原时长占用不同'),'长会话造成持续满载','E5',[
        '比较的是到达强度相同但持续时间和片内时刻不同的两种配置，不是算法更换。',
        '30 s组能在片末释放；673.333 s加权原始时长令每批任务占多个片，早期积累后持续触顶。',
        '5个掉线LL在新快照中的乐观时延下界约20.35–22.73 ms，已经超过20 ms；片内静止不能消除跨片几何变化。'],str(DENSE.relative_to(ROOT))+'/mixed_failure_diagnostics.csv及事件轨迹')
    failures=rows(DENSE/'mixed_failure_diagnostics.csv')
    r.t(['场景/业务','类型/时刻s','传播＋固定项下界ms','预算ms','诊断'],[[f"{int(a['scenario'])}/{a['biz']}",f"{a['kind']}/{fmt(a['t_sec'],1)}",fmt(a['delay_lower_bound_ms']),fmt(a['delay_budget_ms'],1),a['cause']] for a in failures])
    r.p('诊断表同时保留固定30 s参考中的3个GEO相关初始不可行：乐观下界约260–261 ms已高于200 ms。其余5行为长任务掉线。下界超过预算能证明相应快照的传播约束不可行；下界不超过预算则不能倒推出必有满足带宽、损失和放置约束的路由。')
    r.h1('14. E6：降低背景、保留短峰后的三策略混合实验')
    r.p('为测试单场景多业务分配而不是持续过载，背景任务按20%标称占用估计，随机散到120片；1800–1830 s安排一次脉冲，只按输入到达/结束预测将请求峰值推到105%硬上限的最小9倍数，再验证实际承载峰值。总生成比例3∶2∶3∶1，不要求每片比例相同。此阶段仍未限制尾部业务在3600 s内结束。')
    r.eq('稀疏背景与短峰规模','N_bg=9·round[0.2N_capT/(9E[D])]；N_pulse=9·min{u:offered_peak(u)≥ceil(1.05N_cap)}','experiments/run_mixed_3600.py:sparse_peak_profile')
    r.p('先运行三组S1时延，后以保存的相同逐任务输入补测均衡/能耗六组，合并9组；同一场景三策略严格配对，不能把复用参考说成9次全新生成。背景216/270/531、脉冲180/225/423，总生成396/495/954；旧固定种子42/1042/2042分别用于端点速率、时间、背景片号。')
    mixed_tables(r,data['E6'],'E6')
    r.fig(concurrency_plot(r,lambda sc:[('高到达',DENSE/'variable_random'/f'sc{sc}',COLOR['balance']),('低背景短峰',SPARSE/'variable_random'/f'sc{sc}',COLOR['energy'])],'E6_concurrency','到达机制修订：持续过载转为单次短峰'),'稀疏业务输入修订','E6',[
        '两曲线使用不同生成数量，不是同输入算法对照；均沿用时延策略。',
        '平均并发由约198.5/248.2/496.4降为62.89/78.27/156.48，仍各触及200/250/500；接纳率约96%–97%，不再要求平均至少半上限。',
        '接纳率上升主要源于移除长期超额到达，不是路由算法突然获得十倍能力；新样本没有复现5个掉线不等于跨片问题已数学消失。'],str(SPARSE.relative_to(ROOT))+'/arrival_profile_summary.csv')
    r.fig(metric_overview(r,data['E6'],'E6_strategies','原时长短峰的三策略对照'),'原时长三策略性能','E6',[
        '同场景同任务，先比较QoS与承载，再看延迟、归一化能量和热点。',
        '场景2时延/均衡/能耗的J/Gbit为56.690/11.521/10.780，最忙链路约99.778%/3.527%/11.027%。均衡与能耗都减少原始大流量搬运，但目标不是同一个量。',
        '三策略的完整完成/期末在服均为330/51、411/65、815/110；不能把期末在服计为完整完成。'],str(OTHER.relative_to(ROOT))+'/mixed_all_metrics.csv')
    r.fig(outcome_plot(r,data['E6'],'E6_censoring','原时长短峰：完整完成与期末在服'),'右删失状态的含义','E6',[
        '每柱分母是生成总数，绿色完整完成与蓝色期末在服必须分开。',
        '51/65/110个已接纳会话配置结束晚于3600 s，窗口内正常服务，能量只计观察区间，不能标为失败或补计未来能量。',
        '此问题随后用生成时截止规则解决；仅在报告里改状态名称不会令业务真的在窗口内完成。'],str(OTHER.relative_to(ROOT))+'/strategy_summary.csv')
    r.h1('15. E7：窗口内全部截止的三场景三策略实验A')
    r.p('沿用E6的端点、到达、速率和原采样时长，生成时设D_actual=min(D_sample,3600−a)，重算数据量与总操作量；9组全部实际重跑，三策略同场景严格配对。收短会话数51/65/110，未掉线的接纳业务全部实际完成。该修订改善状态解释，不是算法升级。')
    mixed_tables(r,data['E7'],'E7')
    old={(s['scenario'],s['strategy']):s for s in data['E6'] if s['biz']=='ALL'}
    r.t(['场景/策略','旧完成/在服','新完成/在服','窗口服务秒差','窗口能耗J差'],[
        [marker(s),f"{integer(old[s['scenario'],s['strategy']]['completed'])}/{integer(old[s['scenario'],s['strategy']]['end_active'])}",f"{integer(s['completed'])}/{integer(s['end_active'])}",fmt(s['service_s']-old[s['scenario'],s['strategy']]['service_s'],6),fmt(s['energy']-old[s['scenario'],s['strategy']]['energy'],6)] for s in data['E7'] if s['biz']=='ALL'])
    r.p('表中窗口服务与能量差约0，说明原事件引擎本来就只观测到3600 s。新任务在生成时把未来尾段移除，所以接纳/并发/窗口能量可以不变、完整完成数却增加；这不能当作可靠性或能效改善。')
    paired=rows(BOUNDED/'common_service_comparison.csv')
    r.t(['场景/策略','共同服务任务','共同服务秒','同区间相对时延策略节能%','两者服务加权时延ms'],[
        [f"{int(a['scenario'])}/{LABEL[a['strategy']]}",integer(a['common_tasks']),fmt(a['common_service_seconds']),fmt(a['common_energy_reduction_pct'],2),f"{fmt(a['latency_common_delay_ms'])}/{fmt(a['strategy_common_delay_ms'])}"] for a in paired if a['biz']=='ALL'])
    r.p('共同任务与重叠服务区间使比较分母一致，避免通过更早掉线/拒绝高成本任务获得表面节能。该控制仅适用于E7配对输入，不适用于E8同名任务。当前三场景能耗策略同区间相对时延节能约68.13%、80.98%、67.64%；同时需要看其时延与链路热点，不把单目标节能称为所有指标全优。')
    mixed_figures(r,data['E7'],BOUNDED,'E7')
    r.h1('16. E8：每场景、每策略均变化业务的独立随机复测B')
    manifest=js(FRESH/'batch_manifest.json')
    r.p(f"按用户确认，九组测试均独立生成业务，而不是只跨场景变换、三策略共享任务。主种子{manifest['master_seed']}，每组分别派生业务、时间和背景片号种子，共27个互异随机流；批次{manifest['batch_id']}。同一配置可显式重放，但新--prepare默认用新系统熵源。没有为了追求某种算法排序反复挑选种子。")
    r.t(['场景/策略','业务种子','时间种子','背景种子','背景/脉冲','预测峰值'],[
        [f"{a['scenario']}/{LABEL[a['strategy']]}",a['seeds']['task'],a['seeds']['time'],a['seeds']['schedule'],f"{a['background_tasks']}/{a['pulse_tasks']}",a['offered_peak']] for a in manifest['cases']])
    workload=rows(FRESH/'workload_summary.csv')
    r.t(['场景/策略','生成数','请求平均并发','请求输入Gbit','平均实际时长s','HC平均速率Mbps'],[
        [f"{int(a['scenario'])}/{LABEL[a['strategy']]}",integer(a['generated']),fmt(a['offered_time_avg_active']),fmt(a['offered_input_Gbit']),fmt(a['mean_actual_duration_s']),fmt(next(z['mean_rate_mbps'] for z in workload if z['scenario']==a['scenario'] and z['strategy']==a['strategy'] and z['biz']=='HC'))] for a in workload if a['biz']=='ALL'])
    r.p('九组任务输入哈希均不同且均区别于A，算法、资源、业务处理和QoS配置不变。每组请求输入量、实际平均并发也不同，因此下列差值是非配对表现，不是可归因于算法的同任务收益。Task ID会从头编号，不进行跨策略或跨批次同名任务配对。澄清前配对参考目录preclarification_paired_reference未进入正式统计。')
    mixed_tables(r,data['E8'],'E8')
    r.p('关键非一致排序：场景2均衡/能耗链路峰值3.502%/13.926%，但CV为3.163/2.229，能耗策略的CV反而较低；单位能耗10.532/10.514 J/Gbit仅差约0.17%。这既反映不同指标目标不一致，也包含随机业务输入影响，不能据此证明能耗算法在负载均衡上因果优于均衡算法。')
    mixed_figures(r,data['E8'],FRESH,'E8')
    r.h2('16.1 A/B跨批次变化与排序稳定性')
    base={(s['scenario'],s['strategy']):s for s in data['E7'] if s['biz']=='ALL'}
    r.t(['场景/策略','A/B平均并发','A/B时延ms','A/B J/Gbit','归一化能耗变化','A/B链路CV'],[
        [marker(s),f"{fmt(a['avg_active'],2)}/{fmt(s['avg_active'],2)}",f"{fmt(a['delay'])}/{fmt(s['delay'])}",f"{fmt(a['egbit'])}/{fmt(s['egbit'])}",f"{relative(s['egbit'],a['egbit']):+.2f}%",f"{fmt(a['link_cv'])}/{fmt(s['link_cv'])}"] for s in data['E8'] if s['biz']=='ALL' for a in [base[s['scenario'],s['strategy']]]])
    r.fig(source_chart(r,FRESH,'08_cross_batch.png','E8'),'不同随机输入的A/B单位能耗','E8',[
        '淡色A、深色B，在同一场景下比较策略归一化J/Gbit，不按同名task_id求差。',
        'A与B的单位能耗排序均为能耗＜均衡＜时延，但变化幅度依场景/策略不同；场景3均衡比A上升约6.55%，场景2均衡比A下降约8.58%。',
        '归一化只去除部分数据量尺度，端点距离、业务速率/时长与接纳集合仍有影响。只有两批，不能宣称统计显著或稳定置信区间。'],str(FRESH.relative_to(ROOT))+'/cross_batch_comparison.csv')
    r.h1('17. 综合结论、适用范围与待完善项')
    r.t(['问题','本次证据','应如何解读'],[
        ['算力是不是主要瓶颈','HC单任务最高0.5T，对400T节点很小；节点峰值远低于1%','当前是带宽/几何/QoS与并发硬上限更敏感，不是已经验证算力饱和'],
        ['地面越强越值得卸载吗','场景2时延策略HC常把SGL推近满载且传输能量占主导','需权衡先搬原始数据与先在源端压缩；不以TFLOPS单独决策'],
        ['均衡为何也节能','凸负载代价倾向少搬高码率原始流量','两目标可相关，但最忙链路、CV、J/Gbit并不保证同排名'],
        ['能耗为何有时低时延','源端放置、候选差异以及MC目标/统计量不同','不说明时延目标必然代码错误；E8还含独立输入差异'],
        ['成功率从约8%到96%','从持续过载改为低背景＋短峰','主要是实验负载修订，不能当作算法十倍提升'],
        ['完成数增加但J没变','E6→E7生成时收短尾段，观察窗口未变','是状态与业务配置变化，不是凭空增加服务或节能'],
        ['S0–S5能否直接公平排名','当前多策略目标与S1不完全一致，S5另有实现限制','先修正目标/模型与复现约束，再开展统一参数算法矩阵；本文不制造新结果']])
    r.p('优先待完善：统一S2/S3/S4/S5的balance目标；修正S5源端流守恒与消环/路径恢复；固定S4跨进程随机性；在需要工程时延保证时加入链路串行发送/排队、迁移和真实协议开销。若研究能量管理，应增加平台静态负载、射频功率及电池/补给模型。以上为建议，不是本次已完成修改或测得的额外收益。')
    r.p('实验扩展应明确两条路线：同输入配对适合隔离算法差值；独立业务多批次适合观察样本分布与稳健性。保持每批输入记录、随机种子和真实拒绝原因，不为“看起来有区别”改容量、放宽QoS或选择性删掉高成本任务。')


def appendices(r,data,cfg,network_counts):
    r.h1('附录A. 参数、来源、运行与审计索引')
    r.t(['地面站','纬度°','经度°','类型','当前场景2参与'],[[s['name'],s['lat'],s['lon'],s['type'],'是' if s['type'] in ('EDGE','DC') else '否'] for s in cfg['ground']['stations']])
    r.p('轨道、链路、业务与算法参数已在技术章节列出。地面站坐标是实验建模坐标，不代表真实站点设备已部署。附带的source_manifest.json记录当前源码和引用数据的SHA-256；source_file_index.csv提供文件路径、大小与摘要。报告所有图都有编号、来源及说明，不依赖读者自行猜测图义。')
    r.h2('A.1 历史矩阵与未纳入主比较的记录')
    historic=rows(ROOT/'results'/'compare.csv')
    r.t(['历史标签','场景编号','策略','求解器','到达/接受','最终QoS率','总能耗J','求解P99ms'],[
        [a['tag'],integer(a['scenario']),a['strategy'],a['solver'],f"{integer(a['arrived'])}/{integer(a['accepted'])}",pct(a['qos_overall']),fmt(a['energy_J'],1),fmt(a['solve_p99_ms'])] for a in historic])
    r.p('此39行历史档案仅原样列示证据，不产生当前三场景下的算法名次：没有对应runs目录或effective_config，且旧分析中的场景组成、端点回源、功率与容量和当前版本不同。tag含ablation的3行也不能作为当前模型的压缩/功率/抢占消融证据。旧分析关于S4速度或S5 gap＝0的文字，在缺少本轮原始验证材料时只作历史声称，不升级为本报告已验证事实。')
    r.h2('A.2 数据表与字段字典')
    r.p('report_data.csv是各阶段规范化长表；phase、scenario、strategy、biz为结果定位键，raw_path可回到原始目录。metric_dictionary.csv解释字段、单位、分母与限制；formula_index.csv列全部公式及代码位置；figure_index.csv列每幅图、阶段、来源及完整分析。缺失值是“未记录/不适用”，与实测零不同。')
    r.h2('A.2.1 节点离散度的独立数值复核')
    audited=[s for values in data.values() for s in values if s.get('node_cv_audit')]
    r.t(['阶段/场景/策略','原始节点CV','复核节点CV','理论上限','原始/复核Jain'],[
        [f"{s['phase']}/{s['scenario']}/{LABEL[s['strategy']]}",fmt(s['node_cv_raw'],4),fmt(s['node_cv'],4),
         fmt(math.sqrt({1:500,2:528,3:1006}[s['scenario']]-1),4),trip(s,'node_jain_raw','node_jain',d=6)] for s in audited])
    r.p('复核从每个Allocation的开始/结束时刻重建计算节点上仍活动的申请集合，用math.fsum对当前申请直接求和，空闲节点计0；'
        '同刻释放/申请合并后再对有长度区间积分，用于独立检查增量矩统计的数值稳定性。'
        '非负n维负载的CV≤√(n−1)，全部33组原值与复核值均满足此界；CV/Jain及节点峰值与原值在浮点误差内一致，未发现本批节点分布指标异常。'
        '此次复核没有改动UtilizationMonitor或其他运行时代码，也不改变任务承载、路径、能量和其他原始结果。')
    r.t(['阶段','主要输入文件','原始路径/限制'],[
        ['E1','single_biz_redefined_summary_20261016.csv','single_biz_redefined_scN_20261016/BIZ；无有效配置快照'],
        ['E2','concurrency_stress_summary_20261016.csv','single_biz_concurrency_stress_scN_20261016/BIZ；无有效配置快照'],
        ['E3','all_strategy_single_biz_20261016/all_metrics.csv','36组配置、输入配对、资源及QoS审计'],
        ['E4/E5','mixed_3600_latency_20261016/mixed_all_metrics.csv','分别fixed30/variable_random，不混用variant'],
        ['E6','mixed_3600_other_strategies_20261016/mixed_all_metrics.csv','latency参考位于sparse_peak/variable_random'],
        ['E7','mixed_3600_bounded_all_strategies_20261016/mixed_all_metrics.csv','9组重跑，全部实际结束≤3600'],
        ['E8','mixed_3600_random_batch_…/mixed_all_metrics.csv','独立随机9组；preclarification参考不参与']])
    r.h2('A.3 本次重构验收')
    r.p('重构不运行长时间仿真；检查全部引用原始文件存在，任务数/服务与能量分区一致，E3/E6/E7三策略输入配对，E8九组输入互异，当前三场景t=0节点和链路数符合配置。保留所有旧Word与CSV。新文档独立命名，并通过Word导出PDF检查页数、空白页、嵌入图及公式对象；详细结果见report_validation.json。')
    r.code('python experiments/build_technical_scheme_report.py\npython experiments/validate_technical_scheme_report.py --docx results/technical_scheme_experiments_20261016/智能路由算法技术方案及实验.docx\npython -m unittest discover -s tests -v')


def verify_data(data):
    checks=[]
    for phase,items in data.items():
        for s in items:
            assert abs(s['energy']-s['ec']-s['et'])<max(.01,s['energy']*1e-10)
            assert s['peak_active']<=CAP[s['scenario']]
            if s['biz']!='ALL':
                assert s['dmax']<=(20 if s['biz']=='LL' else 200)+1e-6
                assert s['loss']<=(.0002 if s['biz']=='HR' else .01)+1e-12
                s['admitted_qos_rate']=1.
            if s['mode']=='event':
                assert s['generated']==sum(s[k] for k in ('completed','end_active','dropped','rcap','rnone'))
                assert abs(s['service_s']/3600-s['avg_active'])<1e-6
            if phase in ('E7','E8'):assert s['end_active']==0
        checks.append(dict(phase=phase,metric_rows=len(items),energy_and_status='PASS'))
        for s in items:
            if s['biz']=='ALL':s['admitted_qos_rate']=1.
    for phase in ('E3','E6','E7','E8'):
        keys={};configs={}
        for s in data[phase]:
            if phase!='E3' and s['biz']!='ALL':continue
            path=ROOT/s['raw_path'];generated=read(track(path/'generated_tasks.csv'));signature=input_signature(generated)
            keys[s['scenario'],s['biz'],s['strategy']]=signature
            cfg=js(path/'effective_config.json')
            for field in ('strategy','event_output_dir'):cfg['simulation'].pop(field,None)
            if phase=='E8':
                cfg['simulation'].pop('random_seed',None)
                for field in ('time_seed','fixed_arrival_overrides','arrival_profile'):cfg['traffic'].pop(field,None)
            configs[s['scenario'],s['biz'],s['strategy']]=cfg
            if phase in ('E7','E8'):
                assert all(float(g['t_end'])<=3600 and g['final_status']!='ACTIVE_AT_END' for g in generated)
        if phase=='E8':assert len(set(keys.values()))==9
        else:
            for sc,biz,st in keys:assert len({keys[sc,biz,p] for p in STRATEGIES})==1
        for sc,biz,st in configs:assert all(configs[sc,biz,st]==configs[sc,biz,p] for p in STRATEGIES)
    return checks


DICTIONARY={
    'generated':('个','全部生成请求数，包括拒绝'), 'accepted':('个','初始成功获得分配的任务数'),
    'completed':('个','事件引擎完整完成；传统片引擎未独立导出此项，留空'), 'end_active':('个','窗口停止时仍正常服务，预定结束在窗口后'),
    'dropped':('个','已接纳后跨片失去可行分配'), 'rcap':('个','并发硬上限拒绝，不等于带宽/算力不足'),
    'rnone':('个','该求解器未找到满足硬约束候选，不自动证明全局无解'),
    'accept_rate':('0–1','accepted/generated'), 'complete_rate':('0–1','completed/generated；仅事件结果'),
    'success_rate':('0–1','单片：接纳且QoS通过/生成；事件：(完成+期末在服)/生成'),
    'survival_admitted':('0–1','(完成+在服)/初始接纳'), 'avg_active':('个','事件：总实服秒/T；传统等长片：片统计平均'),
    'peak_active':('个','事件峰值/传统片首最大实际接纳'), 'min_active':('个','传统片最小实际接纳；不与事件末尾归零比较'),
    'service_s':('任务·s','事件真实服务区间；单片30s采用接纳数×30的模型服务量'),
    'delay':('ms','单片任务时延均值；事件任务生命周期跨片最坏时延均值'), 'p95':('ms','对应delay任务样本的95分位'),
    'dmax':('ms','任务样本中的最大端到端时延'), 'weighted_delay':('ms','按服务段生效秒数加权的最坏目的时延'),
    'tc':('ms','单片处理项均值；事件服务时间加权处理项'), 'loss':('0–1','已服务任务最大累计路径丢包率'),
    'margin':('ms','时延预算−最大时延；越小越接近约束'),
    'energy':('J','当前观察窗口业务增量总能量'), 'ec':('J','η×实服操作量'), 'et':('J','β×各段实际传输bit累积'),
    'etask':('J/任务','总能量/初始接纳任务数；不等于每完整完成任务J'), 'egbit':('J/Gbit','总能量/实际服务原始输入Gbit'),
    'gbits':('10^9 bit','实际服务原始输入量；MC分母只计输入一次'),
    'link_peak':('%','最忙单链路相对自身容量的全程峰值'), 'link_cv':('无量纲','含空闲链路的CV，再按时间平均'),
    'link_jain':('0–1','瞬时Jain再按时间平均，空闲贡献约定0'),
    'node_peak':('%','最忙计算节点利用率；现标定同时等于计算功率利用率'), 'node_cv':('无量纲','计算节点利用率CV的时间平均'),
    'node_jain':('0–1','计算节点Jain的时间平均，非CV平均值的简单函数'),
    'pmean':('W','总动态J/T'), 'ppeak':('W','同刻全网计算＋传输W总和峰值'), 'pcmean':('W','计算J/T'), 'ptmean':('W','传输J/T'),
    'pcpeak':('W','同刻计算W全网和峰值'), 'ptpeak':('W','同刻传输W全网和峰值'),
    'source_pct':('%','源端计算的路由段数/全部路由段数'), 'ground_pct':('%','EDGE/DC计算路由段占比'),
    'pre_hops':('跳','单片前段平均跳数'), 'post_hops':('跳','单片每任务各目的后段跳数的均值再平均'),
    'segments':('段','每次初始分配或跨片重算分配1段'), 'reroute':('次','跨片重算成功次数，不必换计算点'),
    'migrations':('次','计算节点发生变化的跨片重算'), 'preemptions':('次','旧分支抢占次数；正式事件模式禁用'),
    'solve99':('ms','全部单次求解墙钟耗时P99，含失败尝试'), 'fallback':('次','S1原策略候选失败后补充搜索尝试'),
    'recovered':('次','补充候选恢复成功次数'), 'wall':('s','整组仿真墙钟；受并行负载影响'),
    'cap_s':('s','实际active等于硬上限的累计事件区间'), 'cap_pct':('%','cap_s/T×100'), 'near_s':('s','实际并发≥90%上限累计时间'),
    'first_cap':('s','第一次实际达到硬上限的时刻'), 'common_n':('个','三策略共同成功任务集合大小'),
    'common_e':('J/任务','共同成功任务的平均能量'), 'common_d':('ms','共同成功任务的平均时延'),
    'sha':('SHA-256','保存的输入指纹，非性能指标'), 'config_saved':('布尔','该早期运行是否保存有效配置快照'),
    'node_cv_raw':('无量纲','旧事件监视器原始节点CV；本批独立复核与原值浮点一致'),
    'node_jain_raw':('0–1','旧事件监视器原始节点Jain；与独立复核值并存'),
    'node_cv_audit':('标记','当前节点CV/Jain改用实际活动分配直接重建的报告复核值'),
    'admitted_qos_rate':('0–1','已服务任务QoS通过率，按分业务最大时延/损失核验；不包含拒绝任务'),
    'phase':('标识','E1–E8实验阶段'), 'scenario':('编号','当前三场景1/2/3'), 'strategy':('标签','latency/balance/energy'),
    'biz':('标签','ALL总体或LL/HR/HC/MC'), 'T':('s','观察窗口长度'), 'mode':('标签','slot或event统计口径'), 'raw_path':('相对路径','原始运行目录')}


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--out',type=Path,default=OUT);args=ap.parse_args();out=args.out.resolve();out.mkdir(parents=True,exist_ok=True)
    cfg=yaml.safe_load(track(ROOT/'config.yaml').read_text(encoding='utf-8'))
    from simulator import Simulator
    from network import build_snapshot
    counts={}
    for sc in (1,2,3):
        sim=Simulator(cfg,scenario=sc);snap=build_snapshot(0,sim.all_nodes,cfg,sc)
        counts[sc]=dict(nodes=len(snap.nodes),links=len(snap.links),types=dict(Counter(n.type.value for n in snap.nodes.values())))
    assert [counts[sc]['nodes'] for sc in (1,2,3)]==[500,528,1006]
    data=datasets();checks=verify_data(data);node_audit=review_node_statistics(data)
    r=Report(out);diagrams(r)
    cover=r.doc.add_paragraph();cover.paragraph_format.space_before=Pt(85)
    p=r.doc.add_paragraph(NAME,style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=r.p('三场景 · 四业务 · S0–S5求解器 · 三种优化策略\n技术模型、代码实现与实验全过程');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=r.p('资料截止：2026年10月16日（Asia/Shanghai）\n依据当前代码、已保存原始实验数据及最新Word重构\n不新增仿真，不覆盖原报告');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=r.doc.add_paragraph('目录',style='Title');p.paragraph_format.page_break_before=True
    for run in p.runs:run.font.size=Pt(17)
    p=r.doc.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');p._p.append(field)
    r.p('正文先介绍模型和实现，再按E1–E8的历史顺序组织实验。各章图下提供读图、实际观察与解释边界；附录和伴随CSV提供字段、公式、数据来源索引。')
    build_technology(r,cfg,counts);build_experiments(r,data);appendices(r,data,cfg,counts)
    target=out/(NAME+'.docx');r.doc.save(target)
    # Original reports are read as evidence, never overwritten.
    for p in (ROOT/'全场景全策略单业务测试.docx',DENSE/'全场景全策略单业务测试.docx'):track(p)
    source_docs=[]
    for p in (ROOT/'全场景全策略单业务测试.docx',DENSE/'全场景全策略单业务测试.docx'):
        doc=Document(p);source_docs.append(dict(path=str(p.relative_to(ROOT)),paragraphs=len(doc.paragraphs),tables=len(doc.tables),figures=len(doc.inline_shapes),sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
    normalized=[s for values in data.values() for s in values];fields=list(dict.fromkeys(k for s in normalized for k in s))
    assert set(fields)<=set(DICTIONARY),set(fields)-set(DICTIONARY)
    write_csv(out/'report_data.csv',[{k:s.get(k,'') for k in fields} for s in normalized])
    write_csv(out/'metric_dictionary.csv',[dict(field=k,unit=DICTIONARY[k][0],definition=DICTIONARY[k][1]) for k in fields])
    write_csv(out/'figure_index.csv',r.figures);write_csv(out/'formula_index.csv',r.formulas)
    write_csv(out/'node_cv_recalculation.csv',node_audit)
    write_csv(out/'experiment_catalog.csv',[dict(**check,comparison='independent/unpaired' if check['phase']=='E8' else 'paired within scene and business' if check['phase'] in ('E3','E6','E7') else 'single policy / descriptive') for check in checks])
    for p in list(ROOT.glob('*.py'))+list((ROOT/'solvers').glob('*.py'))+list((ROOT/'experiments').glob('*.py'))+list((ROOT/'tests').glob('*.py'))+[ROOT/'使用说明.md']:track(p)
    source_index=[dict(path=str(p.relative_to(ROOT)),bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in sorted(SOURCES)]
    write_csv(out/'source_file_index.csv',source_index)
    manifest=dict(report=str(target),created_at=datetime.now(timezone(timedelta(hours=8))).isoformat(),
        no_new_simulation=True,existing_reports_unchanged=True,source_documents=source_docs,network_t0=counts,
        experimental_phases=checks,metric_rows=len(normalized),figures=len(r.figures),equations=len(r.formulas),
        tables=len(r.doc.tables),paragraphs=len(r.doc.paragraphs),source_files=source_index,
        node_cv_reviewed_runs=len(node_audit),node_cv_raw_out_of_bound=sum(a['raw_exceeds_theoretical_bound'] for a in node_audit),
        explicit_limitations=['S0-S5 performance not unified under current parameters','S5 source flow constraint discrepancy',
            'S2-S4 balance objective discrepancy','E1/E2 frozen config missing','E8 unpaired inputs'])
    (out/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    mirror=ROOT/(NAME+'.docx')
    try:shutil.copy2(target,mirror);manifest['project_root_copy']=str(mirror)
    except PermissionError:manifest['project_root_copy_locked']=str(mirror)
    (out/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in manifest.items() if k!='source_files'},ensure_ascii=True),flush=True)


if __name__=='__main__':main()
