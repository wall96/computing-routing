"""Replace the old 3600-second supplement with nine closed-window strategy runs."""
from __future__ import annotations
import hashlib
import json
from pathlib import Path
import shutil
import sys
from collections import Counter

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Cm,Pt
from experiments.build_all_strategy_report import table,paragraph
from experiments.run_mixed_strategies import read,number,common_comparison
from experiments.run_bounded_mixed import OUT,STRATEGIES

BASE_REPORT=ROOT/'results'/'all_strategy_single_biz_20261016'/'全场景全策略单业务测试_追加混合实验前备份.docx'
NAME='全场景全策略单业务测试.docx'
LABEL={'latency':'最小时延','balance':'负载均衡','energy':'能耗优化'}
COLORS={'latency':'#2864A0','balance':'#D88722','energy':'#27856B'}
BIZ=('LL','HR','HC','MC')
FIG=OUT/'figures'


def load_data():
    manifest=json.loads((OUT/'comparison_manifest.json').read_text(encoding='utf-8'))
    assert manifest['complete'] and manifest['run_groups']==9
    all_metrics=read(OUT/'mixed_all_metrics.csv')
    runs={}
    for sc in (1,2,3):
        for strategy in STRATEGIES:
            path=OUT/strategy/f'sc{sc}'
            stats={r['biz']:number(r) for r in all_metrics
                   if int(float(r['scenario']))==sc and r['strategy']==strategy}
            life=read(path/'lifecycle_tasks.csv');adjust=read(path/'duration_adjustments.csv')
            cfg=json.loads((path/'effective_config.json').read_text(encoding='utf-8'))
            audit=json.loads((path/'closed_window_audit.json').read_text(encoding='utf-8'))
            assert stats['ALL']['active_at_end']==0
            assert all(float(r['t_end'])<=3600 for r in adjust)
            assert all(r['status']!='ACTIVE_AT_END' for r in life)
            assert stats['ALL']['completed']+stats['ALL']['dropped']==stats['ALL']['initially_accepted']
            runs[sc,strategy]=dict(stats=stats,path=path,cfg=cfg,audit=audit,adjust=adjust,life=life,
                slots=[number(r) for r in read(path/'slot_stats.csv')])
    return runs


def plots(runs,independent=False):
    FIG.mkdir(exist_ok=True)
    plt.rcParams.update({'font.sans-serif':['Microsoft YaHei','SimHei','DejaVu Sans'],
        'axes.unicode_minus':False,'font.size':10,'axes.spines.top':False,'axes.spines.right':False})
    def finish(fig,axes,name,title):
        fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',ncol=3,fontsize=9)
        fig.suptitle(title,fontsize=12);fig.tight_layout(rect=(0,.09,1,.95))
        fig.savefig(FIG/name,dpi=220);plt.close(fig)
    for filename,key,title,ylabel,log in (
        ('01_delay.png','lifetime_delay_mean_ms','已接纳任务生命周期最坏时延的均值','时延（ms）',False),
        ('02_energy.png','energy_J_per_input_Gbit','按完整实际业务输入量归一化的能耗','能耗（J/Gbit，对数坐标）',True)):
        fig,axes=plt.subplots(1,3,figsize=(8.5,3.6));x=np.arange(4)
        for sc,ax in enumerate(axes,1):
            for j,st in enumerate(STRATEGIES):
                ax.bar(x+(j-1)*.25,[runs[sc,st]['stats'][b][key] for b in BIZ],.25,
                       color=COLORS[st],label=LABEL[st])
            ax.set_xticks(x,BIZ);ax.set_title(f'场景{sc}');ax.set_ylabel(ylabel);ax.grid(axis='y',alpha=.2)
            if log:ax.set_yscale('log')
        finish(fig,axes,filename,title)
    for filename,keys,titles,labels,title in (
        ('03_load.png',('peak_link_util_pct','mean_link_cv'),('最忙链路峰值','链路利用率CV'),
         ('利用率（%）','CV（越低越均匀）'),'三策略的链路热点与负载分布'),
        ('04_power.png',('mean_total_power_W','peak_total_power_W'),('全网平均动态功率','全网峰值动态功率'),
         ('功率（W）','功率（W）'),'动态计算＋传输功率：不含平台静态功耗'),
        ('05_placement.png',('source_compute_segment_pct','migrations'),('源卫星处理的路由段比例','计算节点迁移次数'),
         ('路由段比例（%）','迁移次数'),'计算位置与跨片稳定性')):
        fig,axes=plt.subplots(1,2,figsize=(8.5,3.6));x=np.arange(3)
        for ax,key,subtitle,ylab in zip(axes,keys,titles,labels):
            for j,st in enumerate(STRATEGIES):
                ax.bar(x+(j-1)*.25,[runs[sc,st]['stats']['ALL'][key] for sc in (1,2,3)],.25,
                       label=LABEL[st],color=COLORS[st])
            ax.set_xticks(x,['场景1','场景2','场景3']);ax.set_title(subtitle);ax.set_ylabel(ylab);ax.grid(axis='y',alpha=.2)
        finish(fig,axes,filename,title)
    fig,axes=plt.subplots(1,3,figsize=(8.5,3.6))
    for sc,ax in enumerate(axes,1):
        for j,st in enumerate(STRATEGIES):
            events=read(runs[sc,st]['path']/'event_timeline.csv')
            ax.step([float(r['t'])/60 for r in events],[float(r['active']) for r in events],
                where='post',color=COLORS[st],label=LABEL[st],ls=('-','--',':')[j],lw=1.1)
        ax.axhline({1:200,2:250,3:500}[sc],color='#888888',lw=.8,ls=':')
        ax.set_title(f'场景{sc}');ax.set_xlabel('时间（min）');ax.set_ylabel('实际并发任务数');ax.grid(alpha=.2)
    finish(fig,axes,'06_concurrency.png','独立随机输入：短峰达到上限、3600秒全部截止' if independent else '短峰达到上限、3600秒全部截止（相同结果的曲线重合）')
    fig,ax=plt.subplots(figsize=(8.5,3.6));x=np.arange(3)
    for j,st in enumerate(STRATEGIES):
        ax.bar(x+(j-1)*.25,[runs[sc,st]['stats']['ALL']['solve_p99_ms'] for sc in (1,2,3)],.25,
               color=COLORS[st],label=LABEL[st])
    ax.set_xticks(x,['场景1','场景2','场景3']);ax.set_ylabel('单次求解P99（ms）');ax.grid(axis='y',alpha=.2)
    finish(fig,[ax],'07_solver.png','实测求解开销：墙钟耗时，不计入网络时延')


def picture(doc,name,caption):
    doc.add_picture(str(FIG/name),width=Cm(17))
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    p=doc.add_paragraph(caption)
    for r in p.runs:r.italic=True;r.font.size=Pt(9)


def build_document(runs):
    doc=Document(BASE_REPORT);before=len(doc.paragraphs)
    doc.add_page_break()
    doc.add_heading('11. 追加：3600秒混合业务三策略实验（全部业务窗口内截止版）',level=1)
    paragraph(doc,'本章整体替换此前3600 s混合业务仿真记录，不在旧补充章后重复追加。前文三场景、三策略、四类单业务的36组实验正文和图表保留。'
        '本轮重新实际运行三个场景×最小时延/负载均衡/能耗优化共9组，每组3600 s、120个30 s时间片；不复用旧结果冒充重跑。'
        '四类业务生成总数比例仍为LL∶HR∶HC∶MC＝3∶2∶3∶1，全部业务在生成时就满足结束时间不晚于3600 s。')
    doc.add_heading('11.1 业务截止规则与场景',level=2)
    paragraph(doc,'设原采样会话时长为Dsample，到达时刻为a，实际配置时长D＝min(Dsample,3600−a)，结束时刻e＝a＋D≤3600。'
        '例如3000 s到达、原采样1200 s的业务，实际配置时长为600 s，并于3600 s完成。实际输入数据量和操作次数分别改为rate×D和rate×D×每bit操作数，'
        '速率、每秒算力需求及QoS不变；这不是观测窗口结束时把未完成任务强行改为完成。')
    paragraph(doc,'原采样时长仍为LL 60–300 s、HR 120–600 s、HC/MC 600–1800 s；尾部会话被截止规则收短后，实际时长可低于原采样下限。'
        '采用“原时刻到达＋窗口内截止”，而不是重新抽取到达时刻、删掉尾部任务或停止生成某类业务；因此原有生成总数与3∶2∶3∶1比例保持。'
        'sampled_duration_s、实际duration_s、duration_clipped以及两种结束时刻记录在duration_adjustments.csv。')
    table(doc,['场景','网络组成','并发硬上限','本轮生成任务数'],[
        ['1','500 LEO；无EDGE/DC/GEO',200,int(runs[1,'latency']['stats']['ALL']['generated'])],
        ['2','500 LEO＋18 EDGE＋10 DC',250,int(runs[2,'latency']['stats']['ALL']['generated'])],
        ['3','双层LEO各500＋6 GEO；无地面',500,int(runs[3,'latency']['stats']['ALL']['generated'])]])
    paragraph(doc,'业务源和目的均为卫星；地面节点仅提供中转和计算，LL禁止GEO端点。LEO/GEO/EDGE/DC算力分别为400/1350/6000/32000 TFLOPS，'
        '计算功率上限分别为150/500/1400/10200 W，η按Pmax/Cmax标定。以上为当前实验设计值，不是本次测得的真实平台功耗。'
        '三策略的拓扑、链路、资源容量、功率预算、速率、压缩、处理逻辑和QoS均相同，仅路由选择目标不同。')
    table(doc,['业务','生成权重','时延预算ms','丢包预算','每bit等效操作数','压缩后/前比率'],[
        ['LL',3,20,'1%',50,.10],['HR',2,200,'0.02%',200,.33],
        ['HC',3,200,'1%',1000,.20],['MC',1,200,'1%','40＋15×3目的＝85',.05]])
    table(doc,['场景','LL/HR/HC/MC生成','收短会话数','实际时长均值s','最大配置结束时刻s'],[
        [sc,'/'.join(str(int(runs[sc,'latency']['stats'][b]['generated'])) for b in BIZ),
         runs[sc,'latency']['audit']['duration_clipped_tasks'],
         f"{np.mean([float(r['duration_s']) for r in runs[sc,'latency']['adjust']]):.3f}",
         f"{runs[sc,'latency']['audit']['max_requested_end_s']:.1f}"] for sc in (1,2,3)])
    doc.add_heading('11.2 到达负载、三策略与运行方法',level=2)
    paragraph(doc,'保持低强度随机背景＋1800–1830 s一次短脉冲的到达计划。背景任务数按原时长加权均值673.333 s、20%并发上限估计并取9的倍数；'
        '背景随机分配到120片，脉冲用请求重叠预测达到约105%上限的最小9倍数，实际承载峰值另行验证。只要求短时达到上限，不要求全程平均至少半满。'
        '全局生成比例严格保持，单片数量和活跃业务构成不强求该比例。')
    table(doc,['场景','背景任务','脉冲任务','脉冲窗口s','实际峰值（三策略）'],[
        [sc,p['background_tasks'],p['pulse_tasks'],'1800–1830',
         '/'.join(str(int(runs[sc,st]['stats']['ALL']['peak_concurrency'])) for st in STRATEGIES)]
        for sc in (1,2,3) for p in [runs[sc,'latency']['cfg']['traffic']['arrival_profile']]])
    table(doc,['策略','S1候选选择目标','可行性保证与局限'],[
        ['最小时延','压缩前路径传播＋各目的压缩后路径传播之和＋处理时延','满足带宽/算力/功率/时延/丢包硬约束；不是全局最坏目的时延优化器'],
        ['负载均衡','预测分配后链路与计算节点利用率平方和的增量；重叠链路先合并带宽','首选不满足QoS则补充最小时延候选，并用原负载目标排序；不保证每个CV指标单独最优'],
        ['能耗优化','计算能耗＋前段原始数据与后段压缩数据的传输能耗','首选不满足QoS则补充最小时延候选，按能耗排序；不单独优化瞬时功率峰值']])
    paragraph(doc,'使用相同逐任务输入和种子：源目的/速率42，时长/片内偏移1042，背景片号2042。到达在片内均匀随机分布，拓扑片内静态、每30 s更新。'
        '实际到达时才分配，实际业务结束时释放，每个新片重算全部存续会话；不启用片内抢占，不为未来请求预留资源。'
        '截止规则在生成数据量和操作量之前生效；3600 s的完成事件先处理，再结束仿真，最后活跃数为0。')
    doc.add_heading('11.3 关键指标定义与模型边界',level=2)
    table(doc,['指标','口径'],[
        ['完成率/初始接纳率','完成数/生成数；初始接纳数/生成数。存在掉线时两者不同，本轮无期末删失项'],
        ['期末仍在服','本轮应恒为0；对全部生成任务检查a＋D≤3600，不能只修改显示标签'],
        ['拒绝/掉线','并发上限拒绝、初始no_feasible和接纳后跨片掉线分别统计'],
        ['生命周期时延均值/P95/最大','每个已接纳任务先对已服务路由段取最坏目的端到端时延的最大值，再在任务间统计'],
        ['服务加权时延/处理时延','以实际服务秒数加权；不同于唯一任务之间的生命周期最大值均值'],
        ['端到端时延','前段传播＋处理＋后段传播；处理含固定项、参考块bit数×每bit操作数/节点吞吐和简化负载放大'],
        ['平均/峰值并发','全部任务实际服务秒数/3600；峰值取真实事件时刻，不用片末数量的均值替代'],
        ['满并发持续秒数','对实际active＝并发上限的事件区间长度求和，另列其占3600 s的比例'],
        ['计算/传输能耗','按各路由段真实服务秒数积分，跨片不重复计算整会话能耗'],
        ['J/Gbit','总能耗/实际服务原始输入Gbit；不以压缩后的数据量或请求数替代'],
        ['动态平均/峰值功率','总能耗/3600；峰值按同一时刻各服务段计算＋传输功率求和，不能把两项各自峰值直接相加'],
        ['链路/节点峰值利用率','最忙单资源相对自身容量的峰值；二者不是同一种容量尺度'],
        ['CV/Jain','包含空闲资源，按真实事件区间计算再按时长平均；CV小、Jain大表示更均匀'],
        ['源端/地面计算比例','按路由服务段数量统计，不等于唯一任务占比；跨片同一任务可计多段'],
        ['重算/迁移','跨片重路由可不换计算节点；迁移专指计算节点改变，尚未模拟状态传输开销'],
        ['求解P99/运行耗时','单次路由求解与整组墙钟时间；本轮并行运行，受系统负载影响，不加入模拟网络时延']])
    paragraph(doc,'会话持续秒数不直接加到毫秒级分组/数据块QoS中。当前不计链路串行化、显式排队、抖动、重传、求解等待和迁移状态传输。'
        '计算处理使用1/(1−利用率)的简化负载放大，利用率截断0.95，不是显式排队模型。功率只含业务增量计算和传输，不含平台静态功耗，也不是射频发射机铭牌功率。')
    doc.add_heading('11.4 九组总体结果与窗口内完成验证',level=2)
    cases=[(sc,st,runs[sc,st]['stats']['ALL']) for sc in (1,2,3) for st in STRATEGIES]
    table(doc,['场景/策略','接纳/生成','完成/在服/掉线','完成率','上限拒绝/不可行','平均/峰值并发'],[
        [f'{sc}/{LABEL[st]}',f"{int(s['initially_accepted'])}/{int(s['generated'])}",
         f"{int(s['completed'])}/{int(s['active_at_end'])}/{int(s['dropped'])}",f"{s['completed']/s['generated']:.2%}",
         f"{int(s['rejected_cap'])}/{int(s['rejected_no_feasible'])}",f"{s['time_avg_active']:.2f}/{int(s['peak_concurrency'])}"] for sc,st,s in cases])
    table(doc,['场景/策略','时延均值/P95 ms','总能耗kJ','J/Gbit','全网平均/峰值W','链路峰值%'],[
        [f'{sc}/{LABEL[st]}',f"{s['lifetime_delay_mean_ms']:.3f}/{s['lifetime_delay_p95_ms']:.3f}",
         f"{s['energy_total_J']/1000:.3f}",f"{s['energy_J_per_input_Gbit']:.3f}",
         f"{s['mean_total_power_W']:.2f}/{s['peak_total_power_W']:.2f}",f"{s['peak_link_util_pct']:.3f}"] for sc,st,s in cases])
    table(doc,['场景/策略','首次满并发s','满并发s/占比','实际服务秒','最大业务结束s','片末活跃'],[
        [f'{sc}/{LABEL[st]}',f"{s['first_cap_t']:.3f}",f"{s['time_at_cap_s']:.3f}/{s['time_at_cap_pct']:.3f}%",
         f"{s['observed_service_seconds']:.3f}",runs[sc,st]['audit']['max_requested_end_s'],0] for sc,st,s in cases])
    picture(doc,'06_concurrency.png','图9：实际事件并发随时间变化。结果相同的三策略曲线重合；3600 s降至0来自业务配置截止，不是掉线或故障。')
    doc.add_heading('11.5 时延与分业务完整结果',level=2)
    for sc in (1,2,3):
        doc.add_heading(f'场景{sc}',level=3)
        for st in STRATEGIES:
            paragraph(doc,LABEL[st])
            table(doc,['业务','完成/生成','完成率','时延均值/P95/最大ms','J/Gbit','平均/峰值活跃'],[
                [b,f"{int(s['completed'])}/{int(s['generated'])}",f"{s['completed']/s['generated']:.2%}",
                 f"{s['lifetime_delay_mean_ms']:.3f}/{s['lifetime_delay_p95_ms']:.3f}/{s['lifetime_delay_max_ms']:.3f}",
                 f"{s['energy_J_per_input_Gbit']:.3f}",f"{s['time_avg_active']:.2f}/{int(s['peak_concurrency'])}"]
                for b in BIZ for s in [runs[sc,st]['stats'][b]]])
            table(doc,['业务','计算/传输能耗kJ','最大丢包率','处理ms','路由段/重算/迁移','源端/地面计算%'],[
                [b,f"{s['energy_compute_J']/1000:.4f}/{s['energy_tx_J']/1000:.4f}",f"{s['lifetime_loss_max']:.2e}",
                 f"{s['compute_delay_service_weighted_ms']:.4f}",
                 f"{int(s['allocation_segments'])}/{int(s['reroutes'])}/{int(s['migrations'])}",
                 f"{s['source_compute_segment_pct']:.2f}/{s['ground_compute_segment_pct']:.2f}"]
                for b in BIZ for s in [runs[sc,st]['stats'][b]]])
    picture(doc,'01_delay.png','图10：分业务生命周期最坏时延的均值。时延指标取最坏目的并跨片取最大，与S1即时多目的路径代价之和不是同一个优化量。')
    picture(doc,'02_energy.png','图11：单位实际输入数据能耗（对数坐标）。相同任务、相同服务区间的对照另见common_service_comparison.csv。')
    doc.add_heading('11.6 负载、功率与计算放置',level=2)
    table(doc,['场景/策略','链路CV/Jain','节点CV/Jain','节点峰值%','计算/传输平均W','计算/传输各自峰值W'],[
        [f'{sc}/{LABEL[st]}',f"{s['mean_link_cv']:.3f}/{s['mean_link_jain']:.4f}",
         f"{s['mean_node_cv']:.3f}/{s['mean_node_jain']:.4f}",f"{s['peak_node_util_pct']:.4f}",
         f"{s['mean_compute_power_W']:.3f}/{s['mean_tx_power_W']:.3f}",
         f"{s['peak_compute_power_W']:.3f}/{s['peak_tx_power_W']:.3f}"] for sc,st,s in cases])
    picture(doc,'03_load.png','图12：负载均衡策略的链路热点和分布比较。平方利用率增量目标不承诺链路与计算节点每一个峰值或CV都同时最小。')
    picture(doc,'04_power.png','图13：全网动态平均与瞬时峰值功率。峰值按同一时刻计算＋传输项相加；不含平台静态功耗。')
    picture(doc,'05_placement.png','图14：源卫星处理比例与迁移次数。先在源端压缩通常减少原始大流量的跨链路搬运；比例按服务段计数。')
    doc.add_heading('11.7 三策略实际效果与原因分析',level=2)
    pair={ (int(float(r['scenario'])),r['strategy'],r['biz']):number(r)
           for r in read(OUT/'common_service_comparison.csv') }
    for sc in (1,2,3):
        l,b,e=[runs[sc,st]['stats']['ALL'] for st in STRATEGIES]
        paragraph(doc,f"场景{sc}：三策略完成率均为{l['completed']/l['generated']:.2%}，峰值{int(l['peak_concurrency'])}，"
            f"期末在服和跨片掉线均为0。最小时延/负载均衡/能耗优化的生命周期时延均值分别为"
            f"{l['lifetime_delay_mean_ms']:.3f}/{b['lifetime_delay_mean_ms']:.3f}/{e['lifetime_delay_mean_ms']:.3f} ms，"
            f"链路峰值分别为{l['peak_link_util_pct']:.3f}%/{b['peak_link_util_pct']:.3f}%/{e['peak_link_util_pct']:.3f}%。"
            f"同任务、同服务区间下，负载均衡与能耗优化相对最小时延分别节能"
            f"{pair[sc,'balance','ALL']['common_energy_reduction_pct']:.2f}%与{pair[sc,'energy','ALL']['common_energy_reduction_pct']:.2f}%。")
        ll=[runs[sc,st]['stats']['LL'] for st in STRATEGIES]
        paragraph(doc,f"场景{sc}的LL时延均值按三策略顺序为"+
            '/'.join(f"{x['lifetime_delay_mean_ms']:.3f}" for x in ll)+' ms；最大值分别为'+
            '/'.join(f"{x['lifetime_delay_max_ms']:.3f}" for x in ll)+' ms。'
            '已服务分段均须满足20 ms预算，跨片存续任务也逐片重新校验。')
    paragraph(doc,'负载均衡在本组数据中用更低的链路峰值和CV换取更高的部分业务时延与求解开销。'
        '它同时倾向源端早压缩，因此也会显著省传输能量；省电并非能耗策略独有的结果。节点算力利用率很低，'
        '链路利用率项对当前目标影响更大，不能由链路CV改善直接推断计算节点CV必然改善。')
    paragraph(doc,'能耗优化在本次三个场景均取得最低单位输入能耗，主要通过源端处理和压缩后传输减少β×bit的路径累积。'
        '计算能耗占比相对小，不能据此推断真实平台总功耗很低。能耗目标最小化总J，固定3600 s窗口的平均W随总J同比例变化，'
        '但不承诺全网峰值W在任意输入下都最小。')
    paragraph(doc,'三策略都逐片重算路由。能耗优化的计算节点迁移很少，表示计算位置更稳定，不代表不重算或路径不变化。'
        '最小时延目标没有迁移惩罚，同等或接近代价的节点也可能随拓扑和负载变化而切换；迁移次数不是物理上必需迁移次数。'
        '本轮节能统计尚未包含迁移状态传输，不能将迁移减少额外折算为已经测得的节能。')
    s2=[runs[2,st]['stats']['ALL'] for st in STRATEGIES]
    paragraph(doc,f"场景2最小时延的链路峰值仍达{s2[0]['peak_link_util_pct']:.3f}%，地面计算服务段占{s2[0]['ground_compute_segment_pct']:.2f}%；"
        f"负载均衡和能耗优化将链路峰值降到{s2[1]['peak_link_util_pct']:.3f}%、{s2[2]['peak_link_util_pct']:.3f}%。"
        '地面算力更强不代表卸载一定更优：对高码率、强压缩业务，先搬原始数据到地面可能增加前段带宽和传输能量。'
        '当前接近满带宽链路没有显式排队，不能把模型中的QoS达标等同于真实饱和网络中的时延保证。')
    paragraph(doc,'最小时延是策略名称，不保证本表的生命周期最坏时延均值必然低于其他策略。'
        '对于MC，S1按多个目的路径代价求和排序，而结果取最坏目的；跨片统计又取生命周期最大值。'
        '即时目标、约束下的启发式候选集合以及最终统计量不同，其他策略在个别场景取得更低统计时延并不自动说明程序错误。')
    doc.add_heading('11.8 求解开销与可行性补充搜索',level=2)
    table(doc,['场景/策略','求解P99 ms','整组耗时s','补充搜索/恢复成功','迁移/重算'],[
        [f'{sc}/{LABEL[st]}',f"{s['solve_p99_ms']:.3f}",f"{s['wall_seconds']:.1f}",
         f"{int(s['fallback_attempts'])}/{int(s['fallback_accepted'])}",f"{int(s['migrations'])}/{int(s['reroutes'])}"] for sc,st,s in cases])
    picture(doc,'07_solver.png','图15：三策略单次求解耗时P99。并行运行会影响墙钟数值，因此它是本机实测开销，不是网络传播时延或严格复杂度证明。')
    paragraph(doc,'能耗/拥塞最短路径可能超出时延预算，即使传播最短路径仍可行。已有补充搜索只增加候选路径，不改变资源或放宽QoS；'
        '负载均衡在双层千星场景需要评估更多候选并触发更多补充搜索，因而运行更慢。本轮不为强行拉开指标而更改三策略目标或业务参数。')
    doc.add_heading('11.9 与上次3600秒记录的区别',level=2)
    old_root=ROOT/'results'/'mixed_3600_other_strategies_20261016'
    previous={(int(float(r['scenario'])),r['strategy']):number(r) for r in read(old_root/'strategy_summary.csv')}
    table(doc,['场景/策略','旧完成/在服','新完成/在服','窗口内服务秒差','窗口内能耗差J'],[
        [f'{sc}/{LABEL[st]}',f"{int(o['completed'])}/{int(o['active_at_end'])}",f"{int(s['completed'])}/{int(s['active_at_end'])}",
         f"{s['observed_service_seconds']-o['observed_service_seconds']:.6f}",f"{s['energy_total_J']-o['energy_total_J']:.6f}"]
        for sc,st,s in cases for o in [previous[sc,st]]])
    paragraph(doc,'这次提高的是“窗口内完整完成”的可解释性：任务本身被配置为不晚于3600 s结束，不再有右删失会话。'
        '旧版已经只积截至3600 s实际发生的服务与能耗，因此新旧窗口内服务量、资源占用和能量可以保持一致；'
        '完成数上升不能宣称为路由算法可靠性提升。表中列出真实重跑差值，未将原时长未观测的未来数据量计入新业务。')
    doc.add_heading('11.10 验证、限制与复现',level=2)
    paragraph(doc,'11项自动测试通过，包含新增的“时长、数据量及操作量一致收短”和“恰好在结束时刻完成并释放资源”。'
        '九组独立审计全部任务a＋D≤3600、没有ACTIVE_AT_END、全部未掉线接纳任务实际服务秒等于配置时长、三策略逐任务输入相同、'
        '物理配置相同、QoS、资源上下限、路由连续性、服务时间与计算能耗积分、功率峰值一致性。审计与源码SHA-256写入结果目录。')
    paragraph(doc,'仍有少量并发上限拒绝，这是脉冲为确保达到上限而安排的超额请求，不应归因于算力不足。'
        '在当前轻背景负载下三策略接纳数相同，区分主要体现在路径、计算位置、时延、能耗、链路热点、迁移和求解时间。'
        '本轮单种子、单次脉冲，不能代替多种子统计结论；未来可另测更高持续负载、显式排队和迁移成本，但不属于本次结果。')
    paragraph(doc,'输出：results/mixed_3600_bounded_all_strategies_20261016/{latency,balance,energy}/sc{1,2,3}。'
        '每组含全部生成任务、最终任务、路由路径与计算节点、带宽/算力/功率资源事件、slot_stats、event_timeline、生命周期和duration_adjustments等CSV，'
        '另有kpi.json、effective_config.json及窗口审计。总体strategy_summary.csv 9行，mixed_all_metrics.csv 45行，'
        'mixed_all_slot_metrics.csv 1080行，common_service_comparison.csv 30行。')
    paragraph(doc,'复现：python experiments/run_bounded_mixed.py --strategy latency --scenario 1；依次替换三种策略与三个场景。'
        '九组完成后运行python experiments/run_bounded_mixed.py --summarize，再运行python experiments/build_bounded_mixed_report.py。'
        '一般入口config.yaml中traffic.truncate_at_simulation_end=true，使用实际simulation.t_end_sec作为上限；显式关闭仅用于复现历史未截止实验。')
    return doc,before


def main():
    runs=load_data();plots(runs);doc,before=build_document(runs)
    output=OUT/NAME;doc.save(output)
    targets=[ROOT/'results'/'mixed_3600_latency_20261016'/NAME,
             ROOT/'results'/'mixed_3600_latency_sparse_peak_20261016'/NAME]
    replacement=[]
    for target in targets:
        backup=target.with_name(target.stem+'_业务截止规则修改前备份.docx')
        record=dict(path=str(target),replaced=False)
        try:
            if target.exists() and not backup.exists():shutil.copy2(target,backup)
            doc.save(target);record.update(replaced=True,backup=str(backup))
        except PermissionError:
            record['reason']='Word文件被占用；未关闭用户窗口或覆盖未保存内容'
        replacement.append(record)
    old=Document(BASE_REPORT);new=Document(output)
    assert [p.text for p in old.paragraphs]==[p.text for p in new.paragraphs[:before]]
    for a,b in zip(old.tables,new.tables):
        assert [[c.text for c in r.cells] for r in a.rows]==[[c.text for c in r.cells] for r in b.rows]
    assert len(new.inline_shapes)==len(old.inline_shapes)+7
    assert sum('11. 追加：3600秒混合业务' in p.text for p in new.paragraphs)==1
    manifest=dict(report=str(output),old_3600_chapter_replaced=True,chapter_count=1,original_single_business_content='PASS',
        original_paragraphs=before,new_paragraphs=len(new.paragraphs),total_figures=len(new.inline_shapes),
        total_tables=len(new.tables),replacement=replacement,baseline_sha256=hashlib.sha256(BASE_REPORT.read_bytes()).hexdigest(),
        builder_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    (OUT/'append_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    comparison=json.loads((OUT/'comparison_manifest.json').read_text(encoding='utf-8'))
    comparison.update(word_report_modified=True,report=str(output))
    (OUT/'comparison_manifest.json').write_text(json.dumps(comparison,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=True))


if __name__=='__main__':main()
