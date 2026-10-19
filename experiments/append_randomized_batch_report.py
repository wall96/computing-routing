"""Append fresh-batch results without replacing the earlier mixed/single-biz chapters."""
from __future__ import annotations
import argparse
import hashlib
import json
import re
from pathlib import Path
import shutil
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from docx import Document
from docx.shared import Cm,Pt
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from experiments import build_bounded_mixed_report as charts
from experiments.build_all_strategy_report import paragraph,table
from experiments.run_randomized_batch import PREVIOUS,dump
from experiments.run_mixed_strategies import read,number,STRATEGIES

NAME='全场景全策略单业务测试.docx'
LABEL=charts.LABEL
BIZ=('LL','HR','HC','MC')
MAIN=ROOT/'results'/'mixed_3600_latency_20261016'/NAME
TITLE_SUFFIX='新随机业务批次：三场景三策略复测与跨批次分析'


def next_chapter(doc):
    numbers=[int(m.group(1)) for p in doc.paragraphs if p.style.name=='Heading 1'
             for m in [re.match(r'^(\d+)\.',p.text)] if m]
    return max(numbers,default=11)+1


def cross_plot(batch,runs,old):
    fig,axes=plt.subplots(1,3,figsize=(8.5,3.8));x=np.arange(3)
    for sc,ax in enumerate(axes,1):
        for j,st in enumerate(STRATEGIES):
            ax.bar(j-.18,old[sc,st]['energy_J_per_input_Gbit'],.35,color=charts.COLORS[st],alpha=.42,
                label='A '+LABEL[st])
            ax.bar(j+.18,runs[sc,st]['stats']['ALL']['energy_J_per_input_Gbit'],.35,color=charts.COLORS[st],
                label='B '+LABEL[st])
        ax.set_xticks(x,['时延','均衡','能耗']);ax.set_title(f'场景{sc}');ax.set_ylabel('J/Gbit');ax.grid(axis='y',alpha=.2)
    fig.suptitle('两批不同随机输入：归一化能耗与策略排序');fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',ncol=3,fontsize=8)
    fig.tight_layout(rect=(0,.13,1,.94));fig.savefig(batch/'figures'/'08_cross_batch.png',dpi=220);plt.close(fig)


def image(doc,batch,name,caption):
    doc.add_picture(str(batch/'figures'/name),width=Cm(17))
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    p=doc.add_paragraph(caption)
    for run in p.runs:run.font.size=Pt(9);run.italic=True


def append_document(batch,runs,old,base):
    manifest=json.loads((batch/'batch_manifest.json').read_text(encoding='utf-8'))
    doc=Document(base);before=len(doc.paragraphs)
    assert manifest['batch_id'] not in '\n'.join(p.text for p in doc.paragraphs),'Backup already contains this batch'
    chapter=next_chapter(doc)
    figures=len(doc.inline_shapes)
    cases=[(sc,st,runs[sc,st]['stats']['ALL']) for sc in (1,2,3) for st in STRATEGIES]
    contrasts={(int(float(r['scenario'])),r['strategy']):number(r) for r in read(batch/'unpaired_strategy_comparison.csv')}
    doc.add_page_break();doc.add_heading(f'{chapter}. {TITLE_SUFFIX}',level=1)
    paragraph(doc,'本章在上一批全部业务于3600 s内截止的三策略实验之后追加，不替换前面的单业务试验或第11章。'
        '基准批次记为A，本轮全新随机输入记为B，并由批次标识区分不同轮次。重新实际运行3场景×3策略共9组，每组3600 s、120个30 s拓扑时间片。'
        '本次只更换业务抽样与到达计划的随机种子，不调整算法、物理资源、功率预算、处理逻辑、压缩率或QoS。')
    doc.add_heading(f'{chapter}.1 测试目的、随机化与非配对比较',level=2)
    paragraph(doc,'按照用户确认，每个批次自动生成新的主种子，再为每个场景、每种策略派生互不相同的业务、时间及背景片号种子，共27个派生种子。'
        '因此场景之间不再沿用42/1042/2042，也不直接复用上一批任务；源卫星、目的卫星、业务速率、时长、片内到达时刻与背景片号重新抽取。'
        '同一场景内三种策略也分别使用不同输入，这是非配对比较。所有差异同时包含任务数据与算法影响，不能将单次差值直接归因于策略，'
        '也不能将单位能耗差值称为同任务节能率。澄清前已生成的配对参考数据独立封存，不进入本章正式汇总。')
    paragraph(doc,f"本轮主种子：{manifest['master_seed']}；批次标识：{manifest['batch_id']}。默认由系统熵源生成；显式传入--seed可重放。"
        '种子在看到路由结果之前一次性确定，不为改善指标挑选种子或反复重抽。预估并发只用于按既有规则确定脉冲规模，不读取算法输出。')
    table(doc,['场景/策略','业务种子','时间种子','背景片号种子','背景/脉冲数','请求预测峰值'],[
        [f"{c['scenario']}/{LABEL[c['strategy']]}",c['seeds']['task'],c['seeds']['time'],c['seeds']['schedule'],
         f"{c['background_tasks']}/{c['pulse_tasks']}",c['offered_peak']] for c in manifest['cases']])
    audits=read(batch/'batch_input_audit.csv')
    workload={(int(float(r['scenario'])),r['strategy'],r['biz']):number(r) for r in read(batch/'workload_summary.csv')}
    table(doc,['场景/策略','A输入SHA-256前16位','B输入SHA-256前16位','输入与配置检查'],[
        [f"{r['scenario']}/{LABEL[r['strategy']]}",r['batch_A_sha256'][:16],r['batch_B_sha256'][:16],'A≠B；B九组不同；物理/算法不变'] for r in audits])
    paragraph(doc,'哈希只覆盖任务输入，不包括结果状态和拒绝原因；完整SHA-256见batch_input_audit.csv。'
        'Task ID在各次运行中从头编号，同名ID不表示不同运行中的同一业务。接口使用(batch_id,scenario,strategy,task_id)作为业务联合键。'
        '本轮跨策略、跨场景和跨批次均不进行“同名任务”配对。')
    doc.add_heading(f'{chapter}.2 场景、业务与实验方法',level=2)
    table(doc,['场景','网络组成','最大承载并发','A生成数','B时延/均衡/能耗生成数'],[
        [sc,{1:'500 LEO；无地面/GEO',2:'500 LEO＋18 EDGE＋10 DC',3:'双层LEO各500＋6 GEO；无地面'}[sc],
         {1:200,2:250,3:500}[sc],int(old[sc,'latency']['generated']),
         '/'.join(str(int(runs[sc,st]['stats']['ALL']['generated'])) for st in STRATEGIES)] for sc in (1,2,3)])
    paragraph(doc,'所有业务均在卫星上产生、源目的均为卫星；LL只允许LEO端点，地面节点仅提供转发和算力。'
        'LEO/GEO/EDGE/DC分别为400/1350/6000/32000 TFLOPS，计算功率上限150/500/1400/10200 W。'
        '这是本项目当前模型配置，不作为真实在轨硬件测量结论。')
    table(doc,['业务','生成权重','原抽样时长s','速率','时延/丢包预算','操作/bit；压缩率'],[
        ['LL',3,'60–300','2–10 Mbps','20 ms / 1%','50；0.10'],
        ['HR',2,'120–600','10–100 kbps','200 ms / 0.02%','200；0.33'],
        ['HC',3,'600–1800','100–500 Mbps','200 ms / 1%','1000；0.20'],
        ['MC',1,'600–1800','10–100 kbps','200 ms / 1%','40＋15×3＝85；0.05']])
    paragraph(doc,'3∶2∶3∶1约束是3600 s内各业务的生成总数比例，不是每片活跃比例。背景规模仍按20%标称占用估计，'
        '随机散布到120片；1800–1830 s有一次短脉冲，请求重叠预测达到105%并发上限后停止增加脉冲。'
        '不同随机时长/到达组合可能需要不同脉冲数，因此三场景任务总数不强行等于上一批。实际承载峰值另由事件日志验证。')
    paragraph(doc,'任务实际持续时间D=min(原采样时长,3600−到达时刻)，实际数据量=速率×D，总计算量=数据量×每bit操作数。'
        '尾部会话可能短于原时长下限；duration_adjustments.csv同时保存原抽样与实际配置。'
        '每30 s拓扑更新时重新计算所有存续业务，片内拓扑冻结；按真实到达分配、真实结束释放，3600 s完成事件也处理，不抢占。')
    table(doc,['场景/策略','B中LL/HR/HC/MC数','截止收短数','实际时长均值s','最大结束时刻s'],[
        [f'{sc}/{LABEL[st]}','/'.join(str(int(runs[sc,st]['stats'][b]['generated'])) for b in BIZ),
         int(runs[sc,st]['audit']['duration_clipped_tasks']),
         f"{np.mean([float(r['duration_s']) for r in runs[sc,st]['adjust']]):.3f}",
         f"{runs[sc,st]['audit']['max_requested_end_s']:.1f}"] for sc,st,_ in cases])
    table(doc,['场景/策略','请求预测平均并发','实际平均并发','请求/实服Gbit','HC平均速率Mbps'],[
        [f"{c['scenario']}/{LABEL[c['strategy']]}",f"{c['offered_time_avg_active']:.3f}",
         f"{runs[c['scenario'],c['strategy']]['stats']['ALL']['time_avg_active']:.3f}",
         f"{workload[c['scenario'],c['strategy'],'ALL']['offered_input_Gbit']:.3f}/{runs[c['scenario'],c['strategy']]['stats']['ALL']['observed_input_Gbit']:.3f}",
         f"{workload[c['scenario'],c['strategy'],'HC']['mean_rate_mbps']:.3f}"] for c in manifest['cases']])
    paragraph(doc,'三策略沿用S1：最小时延按即时传播＋处理代价选候选；负载均衡最小化预测链路/算力利用率平方和增量；'
        '能耗优化最小化计算＋传输增量能量。原有QoS可行性补充搜索保留，不放宽业务预算、不增加容量。')
    doc.add_heading(f'{chapter}.3 指标定义与统计口径',level=2)
    table(doc,['指标','定义/解释'],[
        ['完成率/接纳率','完整完成数/生成数；初始接纳数/生成数。接纳后掉线单列，不混为容量拒绝。'],
        ['平均/峰值并发','实际服务秒总和/3600；事件级实际活跃数最大值。满并发秒数按事件区间积分。'],
        ['时延均值/P95/最大','每个已接纳任务先跨片取最坏目的端时延最大值，再对任务求均值/95分位/最大。'],
        ['处理时延','固定开销＋Lref×每bit操作数/可用算力，并使用原有利用率修正；表中为服务时间加权均值。'],
        ['能耗与J/Gbit','服务段计算J＋传输J之和；除以实际服务原始输入Gbit归一化，掉线前消耗也计入。'],
        ['链路/节点CV及Jain','按事件区间积分至3600 s；CV越小通常越均匀，Jain越大通常越公平，含未使用资源。'],
        ['利用率峰值','任意事件时刻任意单链路/节点的最高占用比例，不是全网平均。'],
        ['平均/峰值动态功率','总业务增量J/3600；同一时刻全网计算W＋传输W之和的最大值，不含平台静态功耗。'],
        ['重算/迁移','服务段数−已接纳任务数；计算节点切换次数。重算并不必然迁移。'],
        ['求解P99','单次路由求解墙钟时间的99分位，不计入业务网络时延；并行运行会影响实测耗时。']])
    paragraph(doc,'当前端到端时延只含计算前传播＋业务处理＋计算后传播，MC取最坏目的；不含链路串行发送、显式排队、重传或迁移状态开销。'
        '计算能量采用η×实际操作数，传输能量按β×每段传输bit数累计；这不是整星/地面站的真实总电力消耗。'
        '即时最小时延目标与“跨片、最坏目的”的生命周期指标并不完全相同，不能预设后者在所有场景都由latency取得最低。')
    doc.add_heading(f'{chapter}.4 新批次B：总体运行结果',level=2)
    table(doc,['场景/策略','完成/生成','接纳率','容量/不可行拒绝','掉线/期末在服','平均/峰值并发','满并发秒'],[
        [f'{sc}/{LABEL[st]}',f"{int(s['completed'])}/{int(s['generated'])}",f"{s['initial_accept_rate']:.2%}",
         f"{int(s['rejected_cap'])}/{int(s['rejected_no_feasible'])}",f"{int(s['dropped'])}/{int(s['active_at_end'])}",
         f"{s['time_avg_active']:.2f}/{int(s['peak_concurrency'])}",f"{s['time_at_cap_s']:.3f}"] for sc,st,s in cases])
    table(doc,['场景/策略','时延均值/P95 ms','J/Gbit','总能耗kJ','链路峰值% / CV','平均/峰值动态W'],[
        [f'{sc}/{LABEL[st]}',f"{s['lifetime_delay_mean_ms']:.3f}/{s['lifetime_delay_p95_ms']:.3f}",
         f"{s['energy_J_per_input_Gbit']:.3f}",f"{s['energy_total_J']/1000:.3f}",
         f"{s['peak_link_util_pct']:.3f}/{s['mean_link_cv']:.3f}",f"{s['mean_total_power_W']:.2f}/{s['peak_total_power_W']:.2f}"] for sc,st,s in cases])
    image(doc,batch,'06_concurrency.png',f'图{figures+1}：九组使用独立任务输入，实际并发轨迹不同，均须触及场景上限；请求峰值与承载峰值不同。')
    doc.add_heading(f'{chapter}.5 分业务指标：三场景×三策略×四业务',level=2)
    for sc in (1,2,3):
        doc.add_heading(f'场景{sc}分业务结果',level=3)
        for st in STRATEGIES:
            paragraph(doc,LABEL[st])
            table(doc,['业务','完成/生成','时延均值/P95/最大ms','J/Gbit','平均/峰值活跃','处理ms'],[
                [b,f"{int(s['completed'])}/{int(s['generated'])}",
                 f"{s['lifetime_delay_mean_ms']:.3f}/{s['lifetime_delay_p95_ms']:.3f}/{s['lifetime_delay_max_ms']:.3f}",
                 f"{s['energy_J_per_input_Gbit']:.3f}",f"{s['time_avg_active']:.2f}/{int(s['peak_concurrency'])}",
                 f"{s['compute_delay_service_weighted_ms']:.4f}"] for b in BIZ for s in [runs[sc,st]['stats'][b]]])
            table(doc,['业务','计算/传输kJ','实际输入Gbit','最大丢包率','路由段/重算/迁移','源端/地面处理%'],[
                [b,f"{s['energy_compute_J']/1000:.4f}/{s['energy_tx_J']/1000:.4f}",f"{s['observed_input_Gbit']:.3f}",
                 f"{s['lifetime_loss_max']:.2e}",f"{int(s['allocation_segments'])}/{int(s['reroutes'])}/{int(s['migrations'])}",
                 f"{s['source_compute_segment_pct']:.2f}/{s['ground_compute_segment_pct']:.2f}"] for b in BIZ for s in [runs[sc,st]['stats'][b]]])
    image(doc,batch,'01_delay.png',f'图{figures+2}：新输入下四类业务的生命周期最坏时延均值，按场景及策略对照。')
    image(doc,batch,'02_energy.png',f'图{figures+3}：各类业务单位输入能耗，对数轴。策略输入不同，图中差值不是同任务节能率。')
    doc.add_heading(f'{chapter}.6 策略差异、负载与功率分析',level=2)
    table(doc,['场景/策略','节点峰值% / CV','链路/节点Jain','源端/地面处理%','迁移','求解P99ms','补充搜索/成功'],[
        [f'{sc}/{LABEL[st]}',f"{s['peak_node_util_pct']:.4f}/{s['mean_node_cv']:.3f}",
         f"{s['mean_link_jain']:.4f}/{s['mean_node_jain']:.4f}",f"{s['source_compute_segment_pct']:.2f}/{s['ground_compute_segment_pct']:.2f}",
         int(s['migrations']),f"{s['solve_p99_ms']:.2f}",f"{int(s['fallback_attempts'])}/{int(s['fallback_accepted'])}"] for sc,st,s in cases])
    table(doc,['场景','对照策略','时延/对照 输入Gbit','时延/对照 J/Gbit','非配对单位能耗差%'],[
        [sc,LABEL[st],f"{p['latency_input_Gbit']:.3f}/{p['strategy_input_Gbit']:.3f}",
         f"{p['latency_energy_J_per_Gbit']:.3f}/{p['strategy_energy_J_per_Gbit']:.3f}",f"{p['normalized_energy_difference_pct']:+.2f}"]
        for sc in (1,2,3) for st in ('balance','energy') for p in [contrasts[sc,st]]])
    for sc in (1,2,3):
        l,b,e=[runs[sc,st]['stats']['ALL'] for st in STRATEGIES]
        paragraph(doc,f"场景{sc}：按时延/均衡/能耗顺序，完成数为{int(l['completed'])}/{int(b['completed'])}/{int(e['completed'])}，"
            f"生命周期时延均值为{l['lifetime_delay_mean_ms']:.3f}/{b['lifetime_delay_mean_ms']:.3f}/{e['lifetime_delay_mean_ms']:.3f} ms，"
            f"单位能耗为{l['energy_J_per_input_Gbit']:.3f}/{b['energy_J_per_input_Gbit']:.3f}/{e['energy_J_per_input_Gbit']:.3f} J/Gbit。"
            f"链路峰值为{l['peak_link_util_pct']:.3f}%/{b['peak_link_util_pct']:.3f}%/{e['peak_link_util_pct']:.3f}%，"
            f"链路CV为{l['mean_link_cv']:.3f}/{b['mean_link_cv']:.3f}/{e['mean_link_cv']:.3f}。"
            f"均衡/能耗相对时延的非配对单位能耗差为{contrasts[sc,'balance']['normalized_energy_difference_pct']:+.2f}%/"
            f"{contrasts[sc,'energy']['normalized_energy_difference_pct']:+.2f}%；因输入不同，这不是同任务节能率。")
        active=[runs[sc,st]['stats']['ALL']['time_avg_active'] for st in STRATEGIES]
        offered=[workload[sc,st,'ALL']['offered_input_Gbit'] for st in STRATEGIES]
        ll=[runs[sc,st]['stats']['LL'] for st in STRATEGIES]
        paragraph(doc,f"本场景三组请求数据量范围为{min(offered):.1f}–{max(offered):.1f} Gbit，实际平均并发范围为{min(active):.2f}–{max(active):.2f}，"
            '并非严格等负载。LL按时延/均衡/能耗顺序的生命周期最坏时延均值为'
            +'/'.join(f"{s['lifetime_delay_mean_ms']:.3f}" for s in ll)+' ms，最大值为'
            +'/'.join(f"{s['lifetime_delay_max_ms']:.3f}" for s in ll)+' ms；已服务LL段均受20 ms预算约束。')
    reject=sum(s['rejected_no_feasible'] for _,_,s in cases);drop=sum(s['dropped'] for _,_,s in cases)
    paragraph(doc,f"九组共记录不可行初始拒绝{int(reject)}次、已接纳后掉线{int(drop)}次，期末仍在服均为0。"
        '容量拒绝对应短峰超额请求，不等同于算法找不到路径；各业务在峰值期的到达先后影响哪类任务被拒绝，因此生成比例不等于完成比例。'
        '本轮没有可配对的共同任务；J/Gbit只能部分消除数据量尺度，不能消除端点距离及业务样本差异，不能以总能耗更低直接宣称算法更省电。')
    paragraph(doc,'负载均衡以预测平方利用率增量为目标，倾向避免原始高码率长路径，并在源端先处理压缩，因而也能省能。'
        '它可能接受更长的传播路径以分散热点，生命周期时延不必最低；节点算力本就宽裕时，链路项更能影响当前目标。'
        '目标同时包含链路和算力，不保证每一类资源的每个CV、峰值都最优。')
    b2=runs[2,'balance']['stats']['ALL'];e2=runs[2,'energy']['stats']['ALL']
    paragraph(doc,f"值得关注的非一致排序：场景2中均衡/能耗策略的最忙链路峰值为{b2['peak_link_util_pct']:.3f}%/{e2['peak_link_util_pct']:.3f}%，"
        f"但链路CV为{b2['mean_link_cv']:.3f}/{e2['mean_link_cv']:.3f}，本批能耗策略的CV反而更低。"
        f"两者单位能耗为{b2['energy_J_per_input_Gbit']:.5f}/{e2['energy_J_per_input_Gbit']:.5f} J/Gbit，仅相差"
        f"{100*abs(b2['energy_J_per_input_Gbit']/e2['energy_J_per_input_Gbit']-1):.3f}%。"
        '这说明压低最大热点不等于全网CV最低，且这两组实际业务输入不同，不能把如此接近的能耗差当成可靠的算法优势。'
        '本轮未为保持某种排序而修改目标、参数或重新挑选随机任务。')
    paragraph(doc,'能耗优化倾向源端完成压缩，再传较小的数据量，减小β×bit的逐跳累积；是否使用地面由综合路径代价决定，不是地面算力越强越应卸载。'
        '低迁移表明计算位置稳定，不表示不逐片重算；当前模型未计迁移状态传输，不能把少迁移额外折算成已测得的节能。'
        '功耗策略优化的是动态能量，平均W随总J同步，但任意输入下的瞬时峰值W不保证最小。')
    hc=runs[2,'latency']['stats']['HC'];all_s2=runs[2,'latency']['stats']['ALL']
    paragraph(doc,f"场景2的时延策略中，HC占全网业务增量能量的{hc['energy_total_J']/all_s2['energy_total_J']:.2%}，"
        f"全体路由段地面处理比例为{all_s2['ground_compute_segment_pct']:.2f}%，最忙链路达{all_s2['peak_link_util_pct']:.3f}%。"
        'HC的输入速率远高于其他业务，原始数据在处理前搬运往往主导带宽与传输能耗；地面卸载带来的计算优势可能不足以抵消原始数据搬运成本。'
        '当前模型没有显式链路排队，即使某链路接近满载，也不能据此给真实系统作同等时延保证。')
    image(doc,batch,'03_load.png',f'图{figures+4}：链路峰值与CV。各策略任务数据独立，跨场景资源总量和拓扑也不同，按非配对结果解读。')
    image(doc,batch,'04_power.png',f'图{figures+5}：同一时刻计算＋传输动态功率的平均与峰值，不包含平台静态功耗。')
    image(doc,batch,'05_placement.png',f'图{figures+6}：计算放置和迁移。源端处理段比例按路由服务段计数，不是按任务数或时长加权。')
    image(doc,batch,'07_solver.png',f'图{figures+7}：实测求解P99，受本机并行进程影响；不是业务时延组成部分。')
    doc.add_heading(f'{chapter}.7 A/B两批不同输入的比较与稳健性',level=2)
    table(doc,['场景/策略','A/B平均并发','A/B时延均值ms','A/B J/Gbit','单位能耗变化','A/B链路CV'],[
        [f'{sc}/{LABEL[st]}',f"{a['time_avg_active']:.2f}/{s['time_avg_active']:.2f}",
         f"{a['lifetime_delay_mean_ms']:.3f}/{s['lifetime_delay_mean_ms']:.3f}",
         f"{a['energy_J_per_input_Gbit']:.3f}/{s['energy_J_per_input_Gbit']:.3f}",
         f"{100*(s['energy_J_per_input_Gbit']/a['energy_J_per_input_Gbit']-1):+.2f}%",f"{a['mean_link_cv']:.3f}/{s['mean_link_cv']:.3f}"]
        for sc,st,s in cases for a in [old[sc,st]]])
    image(doc,batch,'08_cross_batch.png',f'图{figures+8}：两批业务输入不同，淡色A、深色B。跨批次只观察归一化指标波动和策略排序，不是逐任务配对。')
    for sc in (1,2,3):
        st_order=sorted(STRATEGIES,key=lambda st:runs[sc,st]['stats']['ALL']['energy_J_per_input_Gbit'])
        old_order=sorted(STRATEGIES,key=lambda st:old[sc,st]['energy_J_per_input_Gbit'])
        changes=[100*(runs[sc,st]['stats']['ALL']['energy_J_per_input_Gbit']/old[sc,st]['energy_J_per_input_Gbit']-1) for st in STRATEGIES]
        paragraph(doc,f"场景{sc}：B相对A的单位能耗变化按时延/均衡/能耗策略为"
            +'/'.join(f'{v:+.2f}%' for v in changes)+'；单位能耗从低到高，A为'
            +'→'.join(LABEL[st] for st in old_order)+'，B为'+'→'.join(LABEL[st] for st in st_order)+'。'
            '两批的速率、端点、到达与时长同时变化；差值不能称为算法升级收益。')
    paragraph(doc,'HC通常主导总体能耗与链路占用，随机端点决定传播距离，随机速率/时长决定实际输入量；即使任务总数相同，整体指标也可能变化。'
        'J/Gbit能消除一部分数据量尺度差异，但仍受端点距离、业务混合与服务状态影响，不是完全无偏的跨样本比较。'
        '总体时延又混合四种业务，业务子集与各业务分位数应一同查看。')
    paragraph(doc,'本次新增一批只能提供第二个样本，不能计算有说服力的稳定置信区间或宣称统计显著。'
        '若要形成正式算法评估，可按本设计继续多批独立抽样，比较各场景/业务的样本分布、均值和不确定性；'
        '如需隔离算法因果收益，可另设三策略共用输入的配对控制试验，但本章未擅自改回配对。'
        '本轮按用户要求仅新增一组，不以更多重复测试替代这一实际结果。')
    doc.add_heading(f'{chapter}.8 数据完整性、接口与复现',level=2)
    paragraph(doc,'15项回归测试通过，新增覆盖九组27个随机流互异、同种子重放、旧默认输入不变，以及从原始CSV重建补充表与原生命周期逐字段一致。'
        '补充表生成中断不等于仿真失败：完整原始导出可以恢复汇总；恢复过程只重放生成器核验输入，不重新计算路由或修改成功/拒绝结果。')
    paragraph(doc,'自动检查：新旧输入哈希不同、B九组任务哈希均不同、27个派生种子互异、算法和物理配置不变；'
        '所有任务区间位于[0,3600]、无ACTIVE_AT_END、完整完成任务服务秒数等于配置时长、实际峰值达到场景上限；'
        '进一步逐资源检查上下界、逐服务段检查时延/丢包、连续性与能耗/功率积分。完整审计见comparison_manifest.json、batch_input_audit.csv与各组审计文件。')
    paragraph(doc,'结果根目录：'+str(batch.relative_to(ROOT))+'。strategy_summary.csv有9行总体，mixed_all_metrics.csv有45行总体及分业务，'
        'mixed_all_slot_metrics.csv有1080片；unpaired_strategy_comparison.csv有6行非配对策略对照，不生成伪同任务服务区间对照；'
        'cross_batch_comparison.csv有45行跨批非配对汇总，two_batch_metrics.csv有90行A/B原指标，workload_summary.csv有45行请求负载摘要。'
        '每组保留generated_tasks、routing_allocations（路径＋计算节点）、resource_allocations（带宽/算力/功率事件）、lifecycle_tasks、event_timeline、slot_stats等CSV。')
    paragraph(doc,'新批次：python experiments/run_randomized_batch.py --prepare（可选--seed重放）；终端打印新目录。'
        '接着对三场景三策略执行python experiments/run_randomized_batch.py --batch <新目录> --scenario <1/2/3> --strategy <latency/balance/energy>。'
        '完成后用同一入口--batch <新目录> --summarize，再运行python experiments/append_randomized_batch_report.py --batch <新目录>。'
        '只有--prepare会产生新业务批次；同一--batch内也按场景和策略分别抽取独立输入。重复指定同一组配置是复现而非新样本；'
        '普通config.yaml入口仍遵守固定种子，历史复现实验不自动变成新业务。')
    return doc,before,figures


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--batch',type=Path,required=True)
    args=ap.parse_args();batch=args.batch.resolve()
    charts.OUT=batch;charts.FIG=batch/'figures'
    runs=charts.load_data();charts.plots(runs,independent=True)
    old={(int(float(r['scenario'])),r['strategy']):number(r) for r in read(PREVIOUS/'strategy_summary.csv')}
    cross_plot(batch,runs,old)
    backup=batch/'追加随机批次前完整报告备份.docx'
    if not backup.exists():shutil.copy2(MAIN,backup)
    doc,before,old_figures=append_document(batch,runs,old,backup)
    output=batch/NAME;doc.save(output)
    original=Document(backup);saved=Document(output)
    assert [p.text for p in original.paragraphs]==[p.text for p in saved.paragraphs[:before]]
    for a,b in zip(original.tables,saved.tables):
        assert [[c.text for c in r.cells] for r in a.rows]==[[c.text for c in r.cells] for r in b.rows]
    assert len(saved.inline_shapes)==old_figures+8
    assert sum(p.text==f'{next_chapter(original)}. {TITLE_SUFFIX}' for p in saved.paragraphs)==1
    mirrors=[MAIN,ROOT/'results'/'mixed_3600_latency_sparse_peak_20261016'/NAME,PREVIOUS/NAME]
    replacements=[]
    for target in mirrors:
        record=dict(path=str(target),replaced=False)
        target_backup=target.with_name(target.stem+'_追加随机批次前备份.docx')
        try:
            if target.exists() and not target_backup.exists():shutil.copy2(target,target_backup)
            shutil.copy2(output,target);record.update(replaced=True,backup=str(target_backup))
        except PermissionError:record['reason']='Word占用；请保存关闭后重试，完整新版已写入新批次目录'
        replacements.append(record)
    manifest=dict(report=str(output),previous_chapters_preserved='PASS',old_paragraphs=before,
        paragraphs=len(saved.paragraphs),old_tables=len(original.tables),tables=len(saved.tables),
        original_figures=old_figures,figures=len(saved.inline_shapes),replacement=replacements,
        baseline_sha256=hashlib.sha256(backup.read_bytes()).hexdigest(),
        builder_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    dump(batch/'append_manifest.json',manifest)
    path=batch/'comparison_manifest.json';meta=json.loads(path.read_text(encoding='utf-8'))
    meta.update(word_report_modified=True,report=str(output));dump(path,meta)
    print(json.dumps(manifest,ensure_ascii=True))


if __name__=='__main__':main()
