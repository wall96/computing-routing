"""Build a field-complete Word guide from actual CSV headers, without running simulations."""
from __future__ import annotations

from collections import defaultdict
import csv
from datetime import datetime, timezone, timedelta
import hashlib
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from experiments.csv_guide_content import INFO, MIXED_NAMES, definition
from experiments.inspect_csv_outputs import inspect, read_rows

NAME = 'CSV输出文件与字段使用说明'
OUT = ROOT / 'results' / 'csv_output_guide_20261016'
BOUNDED = ROOT / 'results' / 'mixed_3600_bounded_all_strategies_20261016'
FRESH = ROOT / 'results' / 'mixed_3600_random_batch_20261016_111126_73e445bc6f54'
GENERATED_NAMES = {'csv_file_inventory.csv', 'csv_schema_catalog.csv', 'csv_field_dictionary.csv'}


def write_csv(path, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def inventory():
    """Read headers only: the existing CSV corpus is approximately 4.9 GB."""
    schemas = defaultdict(list)
    files = []
    for path in sorted((ROOT / 'results').rglob('*.csv')):
        if OUT in path.parents:
            continue
        with path.open(encoding='utf-8-sig', newline='') as stream:
            header = next(csv.reader(stream), [])
        if not header or len(header) != len(set(header)):
            raise ValueError(f'Empty or duplicate CSV header: {path}')
        relative = path.relative_to(ROOT).as_posix()
        stat = path.stat()
        fingerprint = hashlib.sha256(json.dumps(header, ensure_ascii=False).encode()).hexdigest()
        key = (path.name, tuple(header))
        schemas[key].append(relative)
        files.append(dict(path=relative, filename=path.name, bytes=stat.st_size,
            mtime_ns=stat.st_mtime_ns, header_sha256=fingerprint))
    groups = []
    mapping = {}
    for i, ((name, header), paths) in enumerate(sorted(schemas.items()), 1):
        sid = f'C{i:02}'
        for path in paths:
            mapping[path] = sid
        groups.append(dict(schema_id=sid, filename=name, columns=list(header), paths=paths,
            file_count=len(paths), header_sha256=hashlib.sha256(json.dumps(list(header), ensure_ascii=False).encode()).hexdigest()))
    for row in files:
        row['schema_id'] = mapping[row['path']]
    return groups, files


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.7)
    for name, size, font in [('Normal', 10, '宋体'), ('Title', 25, '黑体'),
        ('Subtitle', 13, '黑体'), ('Heading 1', 16, '黑体'), ('Heading 2', 12, '黑体'), ('Heading 3', 10.5, '黑体')]:
        style = doc.styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.15
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    section.header.paragraphs[0].text = '智能路由仿真 | CSV输出文件与字段使用说明'
    for run in section.header.paragraphs[0].runs:
        run.font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')


def table(doc, headers, rows, widths=None, size=8.3):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = str(value)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for data in rows:
        for cell, value in zip(t.add_row().cells, data):
            cell.text = str(value)
    for row_index, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if widths:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(size)
                    if row_index:
                        run.font.color.rgb = RGBColor(24, 32, 40)
                    else:
                        run.bold = True
    doc.add_paragraph('')
    return t


def code(doc, source):
    for line in source.strip('\n').splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line or ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
    doc.add_paragraph('')


RAW_NOTES = {
    'generated_tasks.csv': [
        '本表是需求全集，统计生成比例/初始拒绝时以它为分母。速率与时间均保留原精度，不要用tasks.csv的三位小数速率重新计算精确业务量。',
        '计算关系：D=r×10^6×duration_s；ω=基础op/bit+目的数×附加op/bit；总操作量=Dω；算力占用=r×10^6ω。算力与能量基于原始输入，不再乘压缩比例ρ。',
        '当前典型ω：LL=50、HR=200、HC=1000、MC三目的=85 op/bit。TFLOPS容量在本模型用10^12等效op/s换算，并非严格硬件实测吞吐。',
        '旧22列版本缺少处理逻辑/附加每目的开销等字段。缺列表示旧版未导出，不能把缺列补0后声称业务没有处理需求。'],
    'tasks.csv': [
        '终态任务一行的“速览表”，没有完整路径历史。跨片任务在此保留最后一次compute_node/delay/loss，但事件版能量列是累计能量，两种时间范围不能混为一谈。',
        'REJECT:scenario_concurrency_cap表示初始并发上限拒绝；REJECT:no_feasible表示初始无可行候选。拒绝行的路由/能量字段为空，不是0。',
        '如需LL平均时延，单片试验可对成功行delay_ms求均值；混合跨片应筛lifecycle中initially_accepted=1，再平均lifetime_max_delay_ms。'],
    'routing_allocations.csv': [
        '每个allocation_id表示一段可执行的“网络路径+计算节点+资源配置”，而非一项新业务。以task_id连接生成输入后，关系为一任务对多段。',
        '前段承载原始速率r；每条目的后段承载ρr。同一无向链路被多段/多目的使用时，bw_per_link_json已合并，不要再重复叠加。MC按多条目的路径分别复制，不是共享树去重。',
        '具有service_t_end/service_duration_s的事件版本：每段已按真实服务区间重算能量，可对段求和核对lifecycle。没有这两列的传统跨片版本：能量可能重复代表完整任务估算，不能直接按历史行求和。',
        'max_delay_ms=最大目的(前段传播+处理时延+后段传播)；处理项=fixed+Lref×ω/C×1000/(1−u)，u为分配前利用率并按代码裁剪。模型不含逐链路序列化、显式队列、重传和迁移状态复制时延。'],
    'resource_allocations.csv': [
        '这是事件长表，不是每一行一个完整路由。一次任务占用会产生多条BANDWIDTH以及COMPUTE、可能的POWER行；直接连接路由后求和能量会被行数放大。',
        '利用率=(总量−剩余)/总量。带宽用Mbps，算力用等效op/s，功率用W；只能在同类型单位内相加。POWER仅约束计算增量功率，不是电池能量或传输RF功率。',
        '按event_seq顺序应用remain_after；相同t不能任意重排。换拓扑时资源重置，边界同刻含旧片释放和新片占用，需以快照边界划分，不能把整个3600 s流水当一个永不重置的资源池。',
        '仅本表被触及资源的均值不等于全网均值：闲置节点/链路未必有事件行；全网CV/Jain必须纳入完整快照的空闲资源。若只需已分配路径和带宽，用routing表更简单。'],
    'slot_stats.csv': [
        '传统版本8列或11列：一片一批调度并记录快照；事件版本23列：片内到达/结束立即更新、边界重新计算。是否存在active_mean等列可帮助识别，但应以effective_config确认。',
        '事件版t是片首，active是片末状态，active_mean是本片真实时间均值，active_max是片内峰值。任务恰在3600 s结束时，末片active可以为0，而平均并发仍显著大于0。',
        '全窗口平均并发=Σ(active_mean×片长)/T；峰值=max(active_max)。只要各片相等长，active_mean算术均值也等价；不能对片末active平均来替代。',
        'accepted/rejected只计本片新请求；rerouted计存量业务重算成功。换片并非重新产生任务，accepted+rerouted不能当作新业务数。'],
    'lifecycle_tasks.csv': [
        '优先用于混合跨片终态和任务级汇总。COMPLETED为真实完成；DROPPED为先接纳后掉线；REJECTED为初始拒绝；ACTIVE_AT_END是窗口右截尾，不是失败也不是完成。',
        '有限窗口正式A/B批次将任务在生成时截断至3600 s，期末在服应为0。requested_t_end是截断后实际预定结束时刻，原始抽样时间看duration_adjustments。',
        '拒绝任务lifetime_max_delay_ms/loss等导出0占位。求时延均值前须筛initially_accepted=1，否则拒绝多会造成平均时延假性降低。'],
    'event_timeline.csv': [
        '同一t可以有多行BEFORE/AFTER观察状态；区间积分用上一行状态×(下一行t−上一行t)，同刻间隔0。不要先按t去重后任意选一行。',
        '例如10 s从2个任务变成3个，20 s降为1个，则10～20 s贡献30任务秒。最后3600 s状态虽归0，不改变之前区间服务量。',
        '计算能量=Σ上一行compute_power_W×dt，可与路由段energy_compute_J和lifecycle核对；时间线未含传输功率，不能用它直接得到总能量。'],
    'duration_adjustments.csv': [
        '实际duration_s=min(sampled_duration_s,hard_end_sec−arrival_t)。截断发生在总bit和总操作量计算之前；不是只修改终态文字。',
        '核验实际t_end≤3600、duration_s>0及t_end=arrival_t+duration_s；如要分析原时长分布，用sampled_duration_s，不能把截断后的尾部短任务说成抽样错误。'],
    'arrival_schedule.csv': [
        'background_X是背景计划，arrivals_X是全部计划；额外突发量=arrivals_X−background_X。跨四业务/120片求和可核对generated_tasks的业务数量。',
        '这是到达计数计划，不是精确到达时间列表，也不是实际在服并发。四业务3:2:3:1保证全窗口总数，不要求每片比例都相同。'],
}


class Guide:
    def __init__(self, schemas, files, dictionary):
        self.doc = Document()
        configure(self.doc)
        self.schemas = schemas
        self.files = files
        self.dictionary = dictionary
        self.covered = set()
        self.covered_fields = set()

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text):
        self.doc.add_paragraph(text)

    def file(self, names, title):
        if isinstance(names, str):
            names = [names]
        self.h(title, 2)
        combined = []
        for name in names:
            info = INFO[name]
            versions = [s for s in self.schemas if s['filename'] == name]
            self.covered.add(name)
            self.p(f'{name}：{info[0]}。行粒度：{info[1]} 关联键：{info[2]}。用途：{info[3]}')
            self.p(f'核对入口：{info[4]}。现有{sum(s["file_count"] for s in versions)}份、{len(versions)}种表头；版本索引：'+
                   '、'.join(s['schema_id'] for s in versions)+'。')
            for schema in versions:
                for field in schema['columns']:
                    self.covered_fields.add((name, field))
                    if field not in combined:
                        combined.append(field)
        rows = []
        for field in combined:
            source_name = next(n for n in names if any(s['filename'] == n and field in s['columns'] for s in self.schemas))
            unit, meaning = definition(source_name, field, self.dictionary)
            rows.append([field, unit, meaning])
        self.p('逐列说明如下；这是本组历史表头的字段并集，某个CSV不存在的列表示该版本未导出。实际列顺序以文件首行为准，不能按此表位置硬编码。')
        table(self.doc, ['列名', '存储单位/格式', '含义、口径及使用注意'], rows, [5.5, 2.1, 10.0])
        for note in RAW_NOTES.get(names[0], []):
            self.p(note)


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    schemas, files = inventory()
    dictionary = {r['field']: (r['unit'], r['definition']) for r in read_rows(
        ROOT / 'results' / 'technical_scheme_experiments_20261016' / 'metric_dictionary.csv')}
    fields = []
    for s in schemas:
        if s['filename'] not in INFO:
            raise KeyError(f'Undocumented file: {s["filename"]}')
        for position, field in enumerate(s['columns'], 1):
            unit, meaning = definition(s['filename'], field, dictionary)
            fields.append(dict(schema_id=s['schema_id'], filename=s['filename'], column_index=position,
                field=field, unit=unit, definition=meaning, source=INFO[s['filename']][4]))
    write_csv(OUT / 'csv_file_inventory.csv', files)
    write_csv(OUT / 'csv_schema_catalog.csv', [dict(schema_id=s['schema_id'], filename=s['filename'],
        file_count=s['file_count'], column_count=len(s['columns']), header_sha256=s['header_sha256'],
        columns_json=json.dumps(s['columns'], ensure_ascii=False), paths_json=json.dumps(s['paths'], ensure_ascii=False)) for s in schemas])
    write_csv(OUT / 'csv_field_dictionary.csv', fields)
    g = Guide(schemas, files, dictionary)
    g.doc.core_properties.title = NAME
    g.doc.core_properties.subject = 'CSV内容、逐列含义、关联键、版本差异与程序对接'
    g.doc.add_heading(NAME, 0)
    g.doc.add_paragraph('智能路由仿真 · 数据接口与结果阅读手册', style='Subtitle')
    stamp = datetime.now(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M %Z')
    g.p(f'基于现有代码与本地实验结果核对；生成于{stamp}。')
    g.p(f'覆盖results目录已有{len(files):,}份CSV、{len(INFO)}种文件名和{len(schemas)}种有序表头。相同结构的重复实验目录合并说明，所有实际列均在本手册和机器可读字典中逐项登记。')
    g.p('本次仅新建说明文档、字段字典和只读示例，不重新运行仿真、不改变算法、不覆盖任何既有实验CSV和原Word报告。')
    g.p('阅读建议：先看第1～2章定位文件；普通结果阅读重点第3～4章；接口开发重点第3、8～9章。可用Word导航窗格按文件名定位章节。')
    g.doc.add_page_break()
    g.h('目录')
    p = g.doc.add_paragraph()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
    p._p.append(field)
    g.doc.add_page_break()

    g.h('1. 从哪里开始看：目录、版本和统计边界')
    g.p('项目根目录是本手册所称code目录。下列路径均相对此目录；实际工作区外层还存在同名code文件夹。不要把__MACOSX中的._*.csv当实验结果，它们不是本项目导出的CSV。')
    table(g.doc, ['需求', '推荐入口', '使用方式'], [
        ['最新独立随机业务三策略', FRESH.relative_to(ROOT).as_posix() + '/strategy_summary.csv', '总体9行；四业务完整指标看同目录mixed_all_metrics.csv。非配对。'],
        ['相同混合输入的三策略公平对比', BOUNDED.relative_to(ROOT).as_posix() + '/strategy_summary.csv', '配对有限窗口实验；进一步看common_service_comparison.csv。'],
        ['单业务×三场景×三策略', 'results/all_strategy_single_biz_20261016/all_metrics.csv', '36行，按scenario/biz/strategy筛选。'],
        ['检查具体任务、路由、分配', '上述结果目录下latency、balance或energy/scN/', '先generated，再lifecycle，再routing；资源审计看resource。单业务另有业务子目录。'],
        ['复现综合技术报告图表', 'results/technical_scheme_experiments_20261016/report_data.csv', '按phase、scenario、strategy、biz筛选；225是指标行数，不是试验次数。'],
    ], [3.4, 8.2, 6.0], 8)
    g.p('特别注意：IDE中旧入口results/mixed_3600_latency_20261016/mixed_all_metrics.csv仍是当时的历史数据。后来该目录中的Word更新，并不意味着旧CSV被重新仿真覆盖；判断数据版本必须看目录、配置、manifest和输入指纹，不能只看Word修改时间。')
    g.p('现行场景：1纯500 LEO并发上限200；2为500 LEO+18 EDGE+10 DC并发250；3为双层各500 LEO+6 GEO并发500。历史较早结果可能采用旧节点数或旧容量，不能只凭scenario=1就视为同条件。')
    g.p('正式混合观察窗口为0～3600 s，拓扑片长30 s。单业务压力试验主要为300 s共10片，每任务30 s；较早低负载单业务还有3600 s。以各运行effective_config.json为准，缺失配置的早期目录应保留复现限制。')
    g.h('1.1 统一读取约定', 2)
    table(g.doc, ['项目', '约定和注意'], [
        ['编码/CSV解析', '主要为UTF-8 BOM，Python使用utf-8-sig和csv.DictReader；不能line.split(",")，JSON字段内本身有逗号。'],
        ['空值', '空白表示不适用、未保存或未计算，不自动等于0。全网资源指标在分业务行常为空。'],
        ['数字/比例', 'rate的Mbps为10^6 bit/s，Gbit为10^9 bit；_pct已乘100，_rate与loss多为0～1。CV可大于1。'],
        ['标识/布尔', '节点ID、task_id、allocation_id按字符串读；布尔可能True/False、0/1或PASS检查标记，不用bool("False")转换。'],
        ['精度', 'tasks.csv是显示格式，速率、时延、能量可能已四舍五入；精确计算优先生成表、路由表和生命周期表。'],
        ['版本', '使用实际表头而非固定列序。旧版缺字段不是0；_partial是不完整过程预览，优先无后缀正式文件。'],
        ['输入与输出指纹', 'input_sha256只校验剔除终态列后的业务输入；source_file_index.sha256校验文件字节；二者不能互换。'],
    ], [3.5, 14.1])

    g.h('2. 文件之间如何关联')
    table(g.doc, ['上游 → 下游', '关系和键', '典型用途'], [
        ['generated_tasks → lifecycle_tasks / tasks', '同一run_id内task_id一对一', '把原始需求与终态、累计实服能量对应。'],
        ['generated_tasks → routing_allocations', '同一run_id内task_id一对多', '查看跨片历次计算节点与路径。'],
        ['routing_allocations → resource_allocations', '同一run_id内allocation_id一对多', '核查某段每条链路/节点的占用与释放。'],
        ['generated_tasks → duration_adjustments', '同一run_id内task_id一对一', '确认实际时长如何从原抽样时长截断。'],
        ['arrival_schedule → generated_tasks', '计划slot对应generated_slot，一对多', '核对每片每业务到达计数；不是单任务关联。'],
        ['event_timeline → slot_stats', '按时间片区间积分/最大值聚合', '核对片平均并发、峰值和计算功率。'],
        ['mixed_metrics → mixed_all_metrics → strategy_summary', '单运行5行 → 多运行合并 → biz=ALL子集', '同一批次的不同视图，不是互相独立的实验。'],
    ], [5.2, 5.6, 6.8])
    g.p('建议自建run_id=批次目录/场景/策略/单业务子目录；原始单次运行CSV通常没有batch、scenario和strategy全套列，可从父目录、kpi.json.meta与effective_config补齐。跨运行主键至少为(run_id,task_id)，路由主键为(run_id,allocation_id)。')
    g.p('非配对新批次中，不同策略即使都出现T-LL-000001也不是同一任务。不得用同名task_id跨策略硬连接。配对比较必须先校验输入指纹一致，再匹配任务和重叠服务区间。')
    g.p('防止重复聚合：连接resource后，一段能量会复制到多条资源行；应先按allocation_id聚合或单独在routing/lifecycle层算能量。ALL行已包含四业务，不与四业务行再次相加。')

    g.h('3. 单次运行的九类核心CSV：逐列说明')
    raw = ['generated_tasks.csv', 'tasks.csv', 'routing_allocations.csv', 'resource_allocations.csv',
           'slot_stats.csv', 'lifecycle_tasks.csv', 'event_timeline.csv', 'duration_adjustments.csv', 'arrival_schedule.csv']
    for i, name in enumerate(raw, 1):
        g.file(name, f'3.{i} {name}')

    g.h('4. 汇总指标与单业务对照文件')
    for i, name in enumerate(['compare.csv', 'single_biz_satellite_dc_summary_20261016.csv',
        'single_biz_redefined_summary_20261016.csv', 'concurrency_stress_summary_20261016.csv',
        'all_metrics.csv', 'all_slot_metrics.csv', 'algorithm_before_after.csv'], 1):
        g.file(name, f'4.{i} {name}')
        if name == 'compare.csv':
            g.p('识别方法：有tag和qos_overall是results/compare.csv旧矩阵版；有biz_only是单业务版；再有stress_mode、avg_generated_concurrency等是压力扩展。新旧accepted/final_accepted/ever_allocated不是同一列，不能用列位置替换。')
        if name == 'algorithm_before_after.csv':
            g.p('重要命名例外：old_energy_J与new_energy_J来自旧avg_energy_per_task_J和新energy_mean_J，因此单位应为J/任务。若误当总J，会把算法前后能量结论算错一个任务数倍数。')
    g.file(list(MIXED_NAMES), '4.8 混合业务完整指标及其多种视图')
    g.p('文件关系：mixed_metrics是一组5行；mixed_all_metrics合并多组；strategy_summary只取ALL；new_strategies_metrics取balance/energy；two_batch_metrics多一列batch。它们有重复数据视图，不应拼成“更多样本”。旧fixed30/variable_random两方案汇总通常30行；完整三场景三策略为45行；两批合并90行。')
    g.p('三种时延均值不同：任务生命周期最坏时延均值先对每任务取最大再等权平均；service_weighted_delay按路由服务秒加权；tasks.delay_ms只最后一次分配。若要跨策略画同一种指标，必须保持同一口径和同一样本范围。')
    g.p('CV=σ/μ；Jain=(Σu)²/(nΣu²)，在单个非空闲状态可写为1/(1+CV²)，但时间平均后不能直接用平均CV代入。指标纳入空闲资源，全空闲时间Jain按实现约定0；低利用率的大规模网络CV很大不代表容量被占满。')
    g.p('混合终态计数守恒：generated=completed+active_at_end+dropped+rejected_cap+rejected_no_feasible；initially_accepted=completed+active_at_end+dropped。有限窗口批次active_at_end=0，完成率和final_survival_rate才一致。')
    g.file(['mixed_all_slot_metrics.csv', 'mixed_all_slot_metrics_partial.csv'], '4.9 混合业务跨运行时间片汇总')
    g.p('除scenario/strategy或variant标签外，各列与事件版slot_stats同义。绘图横轴t可换为分钟t/60；按场景和策略分组后绘active_mean/active_max。早期variant表没有strategy列，来自只测latency的历史方案，不应推断为未知策略。')

    g.h('5. 到达计划与失败诊断')
    for i, name in enumerate(['arrival_profile_summary.csv', 'rejection_diagnostics.csv',
        'rejection_audit.csv', 'mixed_failure_diagnostics.csv'], 1):
        g.file(name, f'5.{i} {name}')
    g.p('诊断限制：no_feasible只说明当前求解器候选未找到可行分配。传播下界>预算可以证明该时刻时延不可能满足；下界≤预算不能证明可行，还要满足计算位置、带宽、功率、丢包等约束。')
    g.p('rejection_diagnostics通过同片满资源快照反事实求解：若成功则标RESOURCE_CONTENTION；单独放宽时延或丢包预算后成功则对应DELAY_BUDGET/LOSS_BUDGET；都放宽后才成功为DELAY_AND_LOSS；仍不成功为TOPOLOGY_OR_ALGORITHM_NO_PATH。源/目的无相邻可见节点另有NO_VISIBLE_SATELLITE类。该分类不替代严格全局可行性证明，且重新诊断必须使用匹配的历史配置。')

    g.h('6. 配对、非配对与跨批次比较')
    g.file(['common_service_comparison.csv', 'common_service_comparison_partial.csv'], '6.1 同输入、共同服务区间比较')
    g.p('使用顺序：验证同场景输入一致→匹配同一任务→取双方真实服务区间交集→以各段能量/段秒重算交集能量。正的common_energy_reduction_pct才表示该共同服务量下节能；不能把未服务/被拒绝任务带来的能量减少当算法节能。')
    g.file('cross_batch_comparison.csv', '6.2 A/B跨批次宽表：所有前缀列')
    g.p('本表对每个指标展开batch_A_、batch_B_、change_pct_三列。change_pct=100×(B/A−1)，例如接纳率0.8→0.9为+12.5%相对变化，而不是+10个百分点；需百分点差时另算100×(B−A)。A为0时留空避免除零。各真实列已逐列列于上表。')
    g.file('batch_input_audit.csv', '6.3 输入变化和物理配置审计')
    g.file('unpaired_strategy_comparison.csv', '6.4 各策略独立业务的非配对比较')
    g.p('此处normalized_energy_difference_pct=100×(策略J/Gbit÷基线J/Gbit−1)，负值才表示较低；与6.1节的节能率正负方向相反。即使J/Gbit归一化，源目的距离、业务混合、速率及持续时间仍不同，因此只能描述观察到的差异，不能宣称纯算法因果收益或统计显著。')
    g.file('workload_summary.csv', '6.5 需求工作负载分布')
    g.p('offered为所有生成请求的需求侧负载，observed为实际已服务负载。比较非配对策略时先查看本表，确认生成数、输入Gbit、持续时长和端点分布差异，再分析策略指标。')

    g.h('7. 技术方案Word的配套数据与追溯文件')
    for i, name in enumerate(['report_data.csv', 'metric_dictionary.csv', 'experiment_catalog.csv',
        'figure_index.csv', 'formula_index.csv', 'source_file_index.csv', 'node_cv_recalculation.csv'], 1):
        g.file(name, f'7.{i} {name}')
    g.p('report_data是报告专用统一别名底表，不是模拟器新增原始输出。phase条件和slot/event口径可能不同；早期completed/end_active等未独立记录会留空，不允许补0再做完成率比较。节点CV/Jain的独立重建与原导出在本批浮点误差内一致，复核文件不代表原始CSV存在错误。')

    g.h('8. 如何使用：Excel、Python和后续接口')
    g.h('8.1 Excel/Power Query查看', 2)
    for text in [
        '小文件可直接用Excel打开；中文或分列异常时使用“数据→从文本/CSV”，选择UTF-8、逗号分隔、双引号文本限定。JSON列保持为文本，不再按逗号分列。',
        '把task_id、allocation_id、节点ID和SHA列设为文本；数值列按需转换。status/reject_reason用筛选；不要把空值一律填0。',
        '总体策略比较：strategy_summary按scenario筛选三策略；逐业务对比mixed_all_metrics再筛biz=LL/HR/HC/MC。ALL已含四业务，不参与分业务再求和。',
        '资源流水可能超过Excel单表1,048,576行限制；用Power Query、数据库或Python流式读。不要为“看得全”而另存覆盖原CSV。',
        '画时序：用mixed_all_slot_metrics，按场景+策略分组，以t/60为横轴，active_mean/active_max或link_util_max为纵轴；比例转%时只乘100一次。',
    ]:
        g.p(text)
    g.h('8.2 Python标准库安全读取及路径展开', 2)
    code(g.doc, '''import csv, json
from pathlib import Path
run = Path("results/mixed_3600_bounded_all_strategies_20261016/latency/sc1")
with (run / "routing_allocations.csv").open(
        encoding="utf-8-sig", newline="") as f:
    for row in csv.DictReader(f):
        pre = json.loads(row["path_pre_json"])
        post = json.loads(row["path_post_json"])
        links = json.loads(row["bw_per_link_json"])
        compute = row["compute_node"]
        for link in links:
            u, v, bw = link["u"], link["v"], float(link["bw_mbps"])
        break  # 示例只看第一段；完整处理时移除此行''')
    g.p('历史node_dist是Python字典文本而非JSON，仅对这个已知历史字段使用ast.literal_eval；绝不使用eval。所有*_json字段按标准json.loads解析。')
    g.h('8.3 只读核对脚本：可直接运行', 2)
    g.p('本次新增experiments/inspect_csv_outputs.py，只读输入并打印JSON，不改变仿真或结果；需要事件分段版和lifecycle。--horizon是观察长度，--hard-end是可选绝对截止时刻，不对旧未截断数据强加3600截止检查。')
    code(g.doc, '''$runPath = 'results/mixed_3600_bounded_all_strategies_20261016/latency/sc1'
python experiments/inspect_csv_outputs.py --run $runPath --hard-end 3600''')
    g.p('脚本完成：任务ID唯一性→生成/生命周期全集一致→每任务服务段时长与能量核对→active时间积分与服务秒核对→compute_power积分与计算能量核对→可选实际结束边界检查。输出同时示范完整JSON路径/带宽解析。')
    result = inspect(BOUNDED / 'latency' / 'sc1', hard_end=3600)
    (OUT / 'example_readonly_audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    table(g.doc, ['实际样例：有限窗口场景1时延策略', '只读重算值', '说明'], [
        ['生成/曾接纳/路由段', f'{result["generated"]}/{result["initially_accepted"]}/{result["segments"]}', '任务数与路由分段数不同。'],
        ['终态计数', json.dumps(result['states'], ensure_ascii=False), '所有任务实际结束≤3600，期末在服0。'],
        ['平均并发', f'{result["time_avg_active"]:.6f}', 'Σ实服任务秒/3600，而非片末active均值。'],
        ['生命周期最坏时延任务均值ms', f'{result["lifetime_delay_mean_ms"]:.6f}', '排除初始拒绝任务的0占位。'],
        ['实际输入Gbit/总动态J', f'{result["observed_input_Gbit"]:.6f} / {result["energy_total_J"]:.6f}', '仅实际服务区间，不含拒绝需求。'],
        ['J/Gbit / 平均动态W', f'{result["energy_J_per_input_Gbit"]:.6f} / {result["mean_total_dynamic_power_W"]:.6f}', '分别按输入量、观察窗口归一化，分母不同。'],
    ], [6, 5.1, 6.5])
    g.h('8.4 一项真实LL任务的完整关联示例', 2)
    generated = next(read_rows(BOUNDED / 'latency' / 'sc1' / 'generated_tasks.csv'))
    route = next(read_rows(BOUNDED / 'latency' / 'sc1' / 'routing_allocations.csv'))
    life = next(read_rows(BOUNDED / 'latency' / 'sc1' / 'lifecycle_tasks.csv'))
    g.p(f'本样例task_id={generated["task_id"]}，到达{float(generated["arrival_t"]):.6f} s，持续{float(generated["duration_s"]):.6f} s，原速率{float(generated["rate_mbps"]):.9f} Mbps，ω={generated["cycles_per_bit_effective"]} op/bit，ρ={generated["compression_rho"]}。')
    g.p(f'第一段allocation_id={route["allocation_id"]}，区间[{float(route["t_alloc"]):.6f}, {float(route["service_t_end"]):.6f})，计算点{route["compute_node"]}。前段{route["path_pre_json"]}，后段{route["path_post_json"]}。链路需求{route["bw_per_link_json"]}。前段只有源节点表示源端处理，并非缺失路径。')
    g.p(f'该段实服{float(route["service_duration_s"]):.6f} s，输出{float(route["delivered_data_bits"]):.3f} bit，段总能量{float(route["energy_total_J"]):.9f} J；整任务实际分配{life["allocation_segments"]}段、重算{life["reroutes"]}次、迁移{life["migrations"]}次，终态{life["status"]}，累计能量{float(life["energy_total_J"]):.9f} J。段能量与全任务累计不能互相替代。')
    g.h('8.5 推荐接口模型与检查清单', 2)
    table(g.doc, ['接口实体', '主键与来源', '实施建议'], [
        ['runs运行元信息', 'run_id；配置/kpi.meta/manifest', '保存batch、scenario、strategy、solver、slot_sec、窗口起止、输入指纹、源码/配置版本。'],
        ['tasks任务需求', '(run_id,task_id)；generated', '保留精确速率、到达/结束、QoS、处理op/bit；目的可拆到task_destinations子表。'],
        ['task_results生命周期', '(run_id,task_id)；lifecycle', '原始终态枚举不改写；累计能量和实服bit分开记录。'],
        ['allocations路由段', '(run_id,allocation_id)；routing', 'task_id外键；保存[start,end)和计算节点，路径可JSON存储或按hop展开。'],
        ['resource_events流水', '(run_id,event_seq)；resource', '保留资源类型、单位、前后剩余和拓扑快照标识；只追加，不把历史负值删除。'],
        ['slot_metrics/summary', '运行+片序/业务标签', '派生视图可重算，避免与原始观测混作同一种实体。'],
    ], [4, 5.4, 8.2])
    g.p('快照标识注意：事件资源表没有显式slot列。t恰好等于边界时，旧片释放与新片重算占用共用时刻；结合event_seq及RELEASE_SLOT_END/REROUTE_ALLOCATE等动作划定边界，不能仅floor(t/30)后忽略事件类型。完整链路拓扑/空闲资源需由保存配置按network.py构建，本表不是全量拓扑快照。')
    g.p('最小检查：主键无重复；JSON可解析且后段数=目的数；实际duration>0；有限窗口t_end≤3600；service_t_end−t_alloc=段秒；段秒≤任务秒且连续；生命周期能量=分配段能量之和；总能量=计算+传输；resource剩余在[0,total]容差内；任务状态数量守恒。小于阈值的浮点残差允许容差，不应据此改写原始数据。')
    g.p('统计建议：固定窗口平均总功率=总J/T；总峰值需在同一时间对计算+传输求和再取最大。业务间总J会因HC高数据量主导，比较算法同时报告接纳/完成、J/Gbit、QoS、并发、资源公平性及求解耗时，不能只看总J较小。')

    g.h('9. 实际表头版本索引与覆盖范围')
    g.p('每种有序表头分配C编号；同名列顺序不同也分别登记。各文件在本手册第3～7章按字段并集说明，以下给出版本差异；完整顺序、全部1099个路径和每列字典见随附CSV。')
    for name in sorted(INFO):
        versions = [s for s in schemas if s['filename'] == name]
        union = {f for s in versions for f in s['columns']}
        g.h(name, 2)
        for s in versions:
            missing = sorted(union - set(s['columns']))
            g.p(f'{s["schema_id"]}：{s["file_count"]}份，{len(s["columns"])}列；样例：{s["paths"][0]}。'+
                ('相对该文件字段并集未导出：'+', '.join(missing)+'。' if missing else '包含该文件所有已发现字段。'))
    g.h('附录A. 本手册随附的新文件')
    table(g.doc, ['文件', '每列及使用方式'], [
        ['csv_file_inventory.csv', '每原CSV一行：path项目相对路径；filename文件名；bytes字节大小；mtime_ns修改时间纳秒；header_sha256表头JSON有序列表指纹；schema_id关联版本表。用于定位已有结果及检查原文件未被本次修改。'],
        ['csv_schema_catalog.csv', '每版本一行：schema_id；filename；file_count同表头文件数；column_count列数；header_sha256；columns_json实际有序列名列表；paths_json所有同版原文件路径列表。'],
        ['csv_field_dictionary.csv', '每版本每列一行：schema_id；filename；column_index从1开始的列位置；field列名；unit存储单位/格式；definition定义和注意；source核对的导出代码入口。可直接供接口元数据或自动映射使用。'],
        ['example_readonly_audit.json', '第8章实际样例的机器可读核对结果，非新仿真。'],
        ['guide_manifest.json / guide_validation.json', '文档覆盖范围/原CSV文件状态快照及结构、渲染校验结果；不计入原43类CSV。'],
    ], [5.5, 12.1])
    g.p('复现手册：在code目录执行python experiments/build_csv_guide.py；再执行python experiments/validate_csv_guide.py完成Word目录更新、PDF导出、字段覆盖与原文件未修改检查。只在本手册新目录写入说明产物，不修改任何实验结果。')
    expected = {(s['filename'], f) for s in schemas for f in s['columns']}
    assert g.covered == set(INFO) == {s['filename'] for s in schemas}
    assert g.covered_fields == expected
    g.doc.save(OUT / (NAME + '.docx'))
    manifest = dict(generated_at=stamp, name=NAME, original_csv_files=len(files),
        original_csv_bytes=sum(f['bytes'] for f in files), filename_types=len(INFO), schemas=len(schemas),
        schema_column_definitions=len(fields), unique_filename_fields=len(expected),
        covered_filename_fields=len(g.covered_fields), covered_filenames=sorted(g.covered),
        field_coverage='PASS', tables=len(g.doc.tables), original_files=files,
        generated_sidecar_csv_excluded_from_original_inventory=sorted(GENERATED_NAMES),
        scripts=['experiments/build_csv_guide.py', 'experiments/csv_guide_content.py', 'experiments/inspect_csv_outputs.py'])
    (OUT / 'guide_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in manifest.items() if k not in ('original_files', 'covered_filenames')}, ensure_ascii=True))


if __name__ == '__main__':
    build()
