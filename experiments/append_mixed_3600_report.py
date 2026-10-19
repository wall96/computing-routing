"""Audit six mixed runs and append their figures/results to the existing DOCX."""
from __future__ import annotations
from collections import Counter,defaultdict
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import shutil
import statistics
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from experiments.build_all_strategy_report import table,paragraph

OUT=ROOT/'results'/'mixed_3600_latency_20261016'
LEGACY_OUT=OUT
SPARSE_OUT=ROOT/'results'/'mixed_3600_latency_sparse_peak_20261016'
REVISED=False
DOC_DIR=ROOT/'results'/'all_strategy_single_biz_20261016'
DOC=DOC_DIR/'全场景全策略单业务测试.docx'
BACKUP=DOC_DIR/'全场景全策略单业务测试_追加混合实验前备份.docx'
BIZ=('LL','HR','HC','MC')
VARIANTS=('fixed30','variable_random')
LABEL={'fixed30':'固定30秒','variable_random':'原时长＋随机到达'}
COLOR={'fixed30':'#2864A0','variable_random':'#D88722'}


def read(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))


def numeric(row):
    result={}
    for key,value in row.items():
        if key in ('variant','biz','status','task_id','src','dst_json','reject_reason'):result[key]=value
        else:
            try:result[key]=float(value) if value!='' else None
            except ValueError:result[key]=value
    return result


def save_csv(path,rows):
    with Path(path).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)


def audit_data(scenarios=(1,2,3),save=True):
    runs={};all_rows=[];all_slots=[];pairing={};audits=[];failure_details=[]
    for sc in scenarios:
        canonical=[]
        for variant in VARIANTS:
            path=(LEGACY_OUT if REVISED and variant=='fixed30' else OUT)/variant/f'sc{sc}'
            cfg=json.loads((path/'effective_config.json').read_text(encoding='utf-8'))
            generated=read(path/'generated_tasks.csv')
            lifecycle={r['task_id']:numeric(r) for r in read(path/'lifecycle_tasks.csv')}
            slots=[numeric(r) for r in read(path/'slot_stats.csv')]
            stats={r['biz']:numeric(r) for r in read(path/'mixed_metrics.csv')}
            assert len(slots)==120
            assert cfg['simulation']['t_end_sec']==3600
            assert cfg['simulation']['strategy']=='latency' and cfg['simulation']['solver']=='S1'
            cap={1:200,2:250,3:500}[sc]
            assert stats['ALL']['peak_concurrency']==cap
            if not REVISED or variant=='fixed30':assert stats['ALL']['time_avg_active']>=cap/2
            demand={r['task_id']:r for r in generated}
            counts=Counter(r['biz'] for r in generated)
            assert [counts[b]/w for b,w in zip(BIZ,(3,2,3,1))]==[counts['MC']]*4
            per_slot=Counter((int(r['generated_slot']),r['biz']) for r in generated)
            if not REVISED or variant=='fixed30':
                for i in range(120):
                    assert per_slot[i,'MC']>0
                    assert all(per_slot[i,b]==w*per_slot[i,'MC'] for b,w in zip(BIZ,(3,2,3,1)))
            paired=[{k:r[k] for k in ('task_id','biz','src','dst_json','rate_mbps','delay_budget_ms',
                         'loss_budget','cycles_per_bit_effective','compression_rho','generated_slot')}
                    for r in sorted(generated,key=lambda x:x['task_id'])]
            canonical.append(paired)
            physical_keys=('constellation','ground','compute_capacity_tflops',
                'energy_compute_eta_J_per_cycle','compute_power_budget_W',
                'energy_link_beta_J_per_bit','link_params','channel')
            if variant=='fixed30':
                paired_physics={key:cfg[key] for key in physical_keys}
            else:
                assert paired_physics=={key:cfg[key] for key in physical_keys}
            offsets=defaultdict(list)
            for task in generated:
                arrival=float(task['arrival_t']);duration=float(task['duration_s'])
                offset=arrival-int(task['generated_slot'])*30
                assert 0<=offset<30,task
                offsets[task['biz']].append(offset)
                if variant=='fixed30':assert duration==30 and offset==0
                else:
                    bounds=cfg['traffic']['qos'][task['biz']]['duration_s']
                    assert bounds[0]<=duration<=bounds[1]
                if task['biz']=='LL':assert 'GEO' not in task['src']+task['dst_json']
            energy=0.0;service=0.0;last_end={};segment_count=0;max_delay_by_task=defaultdict(float)
            compute_integrals=defaultdict(float)
            for r in read(path/'routing_allocations.csv'):
                task=demand[r['task_id']]
                start,end=float(r['t_alloc']),float(r['service_t_end']);dt=float(r['service_duration_s'])
                assert abs(end-start-dt)<1e-7
                assert start>=float(task['arrival_t'])-1e-8
                assert end<=min(3600,float(task['t_end']))+1e-8
                assert end<=30*(int((start+1e-8)//30)+1)+1e-7
                if r['task_id'] in last_end:assert abs(start-last_end[r['task_id']])<1e-7
                else:assert abs(start-float(task['arrival_t']))<1e-7
                last_end[r['task_id']]=end
                assert abs(float(r['energy_compute_J'])-float(r['power_used_W'])*dt)<1e-5
                assert float(r['max_delay_ms'])<=float(task['delay_budget_ms'])+1e-6
                assert float(r['max_loss'])<=float(task['loss_budget'])+1e-12
                energy+=float(r['energy_total_J']);service+=dt;segment_count+=1
                compute_integrals[r['biz']]+=float(r['compute_latency_ms'])*dt
                max_delay_by_task[r['task_id']]=max(max_delay_by_task[r['task_id']],float(r['max_delay_ms']))
            assert abs(energy-stats['ALL']['energy_total_J'])<max(.01,energy*1e-10)
            assert abs(service-stats['ALL']['observed_service_seconds'])<1e-3
            assert abs(service-sum(s['active_mean']*30 for s in slots))<1e-3
            fair_integral=defaultdict(float);peak_business=defaultdict(float);previous=None
            cap_seconds=near_cap_seconds=pre_pulse_integral=post_pulse_integral=0.0
            first_cap_t=None;timeline=[]
            with (path/'event_timeline.csv').open(encoding='utf-8-sig',newline='') as f:
                for event in csv.DictReader(f):
                    now=float(event['t'])
                    if previous is not None:
                        dt=now-float(previous['t']);assert dt>=-1e-7
                        for prefix in ('link','node'):
                            cv=float(previous[f'{prefix}_util_cv']);mean=float(previous[f'{prefix}_util_mean'])
                            fair_integral[prefix]+=max(0,dt)/(1+cv*cv) if mean>1e-15 else 0
                        active=float(previous['active'])
                        if active==cap:cap_seconds+=max(0,dt)
                        if active>=.9*cap:near_cap_seconds+=max(0,dt)
                        begin=float(previous['t'])
                        pre_pulse_integral+=active*max(0,min(now,1800)-begin)
                        post_pulse_integral+=active*max(0,now-max(begin,1830))
                    for b in BIZ:peak_business[b]=max(peak_business[b],float(event[f'active_{b}']))
                    if float(event['active'])==cap and first_cap_t is None:first_cap_t=now
                    timeline.append((now,float(event['active'])))
                    previous=event
            for biz,stat in stats.items():
                accepted_lives=[r for r in lifecycle.values() if r['initially_accepted'] and (biz=='ALL' or r['biz']==biz)]
                stat['lifetime_loss_max']=max((r['lifetime_max_loss'] for r in accepted_lives),default=0.0)
                stat['compute_delay_service_weighted_ms']=(sum(compute_integrals.values()) if biz=='ALL' else compute_integrals[biz])/max(1e-12,stat['observed_service_seconds'])
                stat['mean_node_cv']=statistics.fmean(s['node_util_cv_mean'] for s in slots) if biz=='ALL' else None
                stat['mean_link_jain']=fair_integral['link']/3600 if biz=='ALL' else None
                stat['mean_node_jain']=fair_integral['node']/3600 if biz=='ALL' else None
                if biz!='ALL':stat['peak_concurrency']=peak_business[biz]
                stat['time_at_cap_s']=cap_seconds if biz=='ALL' else None
                stat['time_at_cap_pct']=cap_seconds/36 if biz=='ALL' else None
                stat['time_above_90pct_cap_s']=near_cap_seconds if biz=='ALL' else None
                stat['first_cap_t']=first_cap_t if biz=='ALL' else None
                stat['pre_pulse_time_avg_active']=pre_pulse_integral/1800 if biz=='ALL' else None
                stat['post_pulse_time_avg_active']=post_pulse_integral/1770 if biz=='ALL' else None
            if REVISED and variant=='variable_random':
                from experiments.run_mixed_3600 import planned_timing
                timing_cfg=cfg['traffic'].copy()
                timing_cfg['fixed_arrival_overrides']={int(k):v for k,v in cfg['traffic']['fixed_arrival_overrides'].items()}
                predicted_peak,predicted_mean=planned_timing(timing_cfg,120)
                assert predicted_peak==cfg['traffic']['arrival_profile']['offered_peak']
                actual_offered=[]
                for task in generated:
                    actual_offered.extend(((float(task['arrival_t']),1),(min(3600,float(task['t_end'])),-1)))
                active=offered_peak=0
                for when,delta in sorted(actual_offered):active+=delta;offered_peak=max(offered_peak,active)
                assert offered_peak==predicted_peak,('timing forecast',offered_peak,predicted_peak)
                assert cap_seconds>0 and cap_seconds<360,('peak should be brief',cap_seconds)
                assert 1800<=first_cap_t<1830
                assert sum(counts.values())<cap*3
            for tid,r in lifecycle.items():
                if r['status']=='COMPLETED':assert abs(r['service_seconds']-r['duration_s'])<1e-5
                if r['status']=='ACTIVE_AT_END':assert r['requested_t_end']>3600
                if r['initially_accepted']:assert abs(max_delay_by_task[tid]-r['lifetime_max_delay_ms'])<1e-6
            candidates=[(tid,r) for tid,r in lifecycle.items() if r['status']=='DROPPED' or r['reject_reason']=='no_feasible']
            if candidates:
                from simulator import Simulator
                from network import build_snapshot
                from models import Task,BizType
                from solvers.base import single_source
                physical=json.loads(json.dumps(cfg))
                for k in ('scenario_shells','scenario_geo','scenario_ground','scenario_max_concurrency'):
                    physical['constellation'][k]={int(k2):v for k2,v in physical['constellation'][k].items()}
                sim=Simulator(physical,scenario=sc,no_scale=True)
                cache={}
                for tid,life in candidates:
                    raw=demand[tid]
                    when=float(life['last_service_end']) if life['status']=='DROPPED' else float(raw['arrival_t'])
                    epoch=int((when+1e-7)//30)
                    if epoch not in cache:cache[epoch]=build_snapshot(epoch*30,sim.all_nodes,physical,sc)
                    task=Task(tid,raw['src'],json.loads(raw['dst_json']),BizType(raw['biz']),
                        int(raw['priority']),raw['mode'],float(raw['rate_mbps']),
                        float(raw['delay_budget_ms']),float(raw['loss_budget']),
                        float(raw['cycles_per_bit']),float(raw['Lref_bits']),float(raw['arrival_t']),float(raw['duration_s']))
                    dist,_=single_source(cache[epoch],task.src,task,'latency',0,0,4)
                    lower=max(dist[d] for d in task.dst_set)+float(raw['fixed_processing_ms_total'])
                    failure_details.append(dict(scenario=sc,variant=variant,task_id=tid,biz=raw['biz'],
                        kind=life['status'],t_sec=when,src=raw['src'],dst_json=raw['dst_json'],
                        delay_budget_ms=task.delay_budget_ms,delay_lower_bound_ms=lower,
                        cause='传播时延下界超过预算' if lower>task.delay_budget_ms else '资源或候选限制；未证明全局不可行'))
            # Stream resource checks; large CSVs are not retained in memory.
            resources_checked=0
            with (path/'resource_allocations.csv').open(encoding='utf-8-sig',newline='') as f:
                for r in csv.DictReader(f):
                    remain,total=float(r['remain_after']),float(r['total_capacity'])
                    tolerance=max(1e-6,total*1e-12)
                    assert -tolerance<=remain<=total+tolerance
                    assert float(r['t'])<=3600
                    resources_checked+=1
            current=dict(scenario=sc,variant=variant,groups_ratio=dict(counts),
                ratio_scope='whole_horizon' if REVISED and variant=='variable_random' else 'every_slot',
                exact_3_2_3_1='PASS',actual_peak_hits_cap='PASS',time_at_cap_s=cap_seconds,
                segments=segment_count,resource_events=resources_checked,
                energy_segment_sum='PASS',service_time_integral='PASS',
                true_arrival_and_end_times='PASS',cross_slot_continuity='PASS',
                QoS_and_resources='PASS',LL_GEO_endpoints='PASS',
                arrival_offset_mean_s={b:statistics.fmean(offsets[b]) for b in BIZ},
                arrival_offset_std_s={b:statistics.pstdev(offsets[b]) for b in BIZ})
            audits.append(current)
            for r in stats.values():all_rows.append(r)
            for r in slots:all_slots.append({'scenario':sc,'variant':variant,**r})
            runs[sc,variant]=dict(stats=stats,slots=slots,cfg=cfg,audit=current,lifecycle=lifecycle,timeline=timeline)
        if not REVISED:assert canonical[0]==canonical[1],('pairing',sc)
        pairing[f'sc{sc}']={variant:hashlib.sha256(json.dumps(rows,sort_keys=True).encode()).hexdigest()
                            for variant,rows in zip(VARIANTS,canonical)}
    if save:
        save_csv(OUT/'mixed_all_metrics.csv',all_rows)
        save_csv(OUT/'mixed_all_slot_metrics.csv',all_slots)
    audit={'run_groups':len(scenarios)*2,'time_slots':len(scenarios)*240,'metric_rows':len(all_rows),
           'paired_endpoints_rates_profiles':'NOT_APPLICABLE: arrival counts intentionally reduced' if REVISED else 'PASS',
           'paired_physical_configuration':'PASS',
           'average_half_cap_required_for_variable':not REVISED,
           'paired_sha256':pairing,'runs':audits,
           'source_sha256':{name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for name in (
               'simulator.py','event_simulation.py','metrics.py','tasks.py','solvers/s1_dijkstra.py',
               'experiments/run_mixed_3600.py','experiments/append_mixed_3600_report.py')}}
    if save:
        (OUT/'mixed_audit_manifest.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
        if failure_details:save_csv(OUT/'mixed_failure_diagnostics.csv',failure_details)
        if REVISED:
            save_csv(OUT/'arrival_profile_summary.csv',[
                dict(scenario=sc,**{k:v for k,v in runs[sc,'variable_random']['cfg']['traffic']['arrival_profile'].items()
                                   if k!='background_counts_by_slot'},
                     **{k:runs[sc,'variable_random']['stats']['ALL'][k] for k in ('generated','initially_accepted',
                        'initial_accept_rate','time_avg_active','peak_concurrency','first_cap_t','time_at_cap_s',
                        'time_at_cap_pct','pre_pulse_time_avg_active','post_pulse_time_avg_active')}) for sc in scenarios])
    return runs


def plots(runs):
    folder=OUT/'figures';folder.mkdir(exist_ok=True)
    plt.rcParams.update({'font.sans-serif':['Microsoft YaHei','SimHei','DejaVu Sans'],
        'axes.unicode_minus':False,'font.size':10.5,'axes.spines.top':False,'axes.spines.right':False})
    def finish(fig,axes,name,title):
        fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',ncol=2,fontsize=9)
        fig.suptitle(title,fontsize=12);fig.tight_layout(rect=(0,.08,1,.95))
        fig.savefig(folder/name,dpi=220);plt.close(fig)
    fig,axes=plt.subplots(1,3,figsize=(8.5,3.6))
    for sc,ax in enumerate(axes,1):
        for var in VARIANTS:
            ss=runs[sc,var]['slots']
            if REVISED and var=='variable_random':
                timeline=runs[sc,var]['timeline']
                ax.step([t/60 for t,a in timeline],[a for t,a in timeline],where='post',
                        label=LABEL[var],color=COLOR[var],lw=1.1)
            else:
                ax.plot([s['t']/60 for s in ss],[s['active_mean'] for s in ss],
                        label=LABEL[var],color=COLOR[var],lw=1.2)
        ax.axhline({1:200,2:250,3:500}[sc],color='#888888',ls=':',lw=1)
        ax.set_title(f'场景{sc}');ax.set_xlabel('仿真时间（min）');ax.set_ylabel('并发任务数');ax.grid(alpha=.2)
    finish(fig,axes,'mixed_01_concurrency.png','3600秒混合业务：实际并发随时间变化')
    for name,key,title,ylabel in [
        ('mixed_02_survival.png','final_survival_rate','完成或期末在服任务 / 到达任务','比例'),
        ('mixed_03_delay.png','lifetime_delay_mean_ms','已接纳任务全生命周期最大时延的平均值','时延（ms）'),
        ('mixed_04_energy.png','energy_J_per_input_Gbit','按实际服务原始输入数据量归一化的能耗','能耗（J/Gbit输入）')]:
        fig,axes=plt.subplots(1,3,figsize=(8.5,3.6));x=np.arange(4)
        for sc,ax in enumerate(axes,1):
            for j,var in enumerate(VARIANTS):
                vals=[runs[sc,var]['stats'][b][key] for b in BIZ]
                ax.bar(x+(j-.5)*.34,vals,.34,color=COLOR[var],label=LABEL[var])
            ax.set_xticks(x,BIZ);ax.set_title(f'场景{sc}');ax.set_ylabel(ylabel);ax.grid(axis='y',alpha=.2)
        finish(fig,axes,name,title)
    fig,ax=plt.subplots(figsize=(8.5,3.7));bottom=np.zeros(6)
    keys=[('completed','已完成','#27856B'),('active_at_end','期末仍在服','#56B4E9'),
          ('dropped','跨片掉线','#C04C56'),('rejected_cap','并发上限拒绝','#D88722'),
          ('rejected_no_feasible','初始不可行','#777777')]
    cases=[(s,v) for s in (1,2,3) for v in VARIANTS]
    for key,label,color in keys:
        values=np.array([runs[s,v]['stats']['ALL'][key]/runs[s,v]['stats']['ALL']['generated']*100 for s,v in cases])
        ax.bar(range(6),values,bottom=bottom,color=color,label=label);bottom+=values
    ax.set_xticks(range(6),[f'场景{s}\n{LABEL[v]}' for s,v in cases],fontsize=9)
    ax.set_ylabel('全部生成任务占比（%）');ax.set_ylim(0,100);ax.set_title('观测结束时的任务状态构成')
    ax.legend(loc='upper center',bbox_to_anchor=(.5,-.20),ncol=3,fontsize=9)
    fig.tight_layout();fig.savefig(folder/'mixed_05_outcomes.png',dpi=220);plt.close(fig)
    fig,axes=plt.subplots(1,3,figsize=(8.5,3.6))
    for sc,ax in enumerate(axes,1):
        ss=runs[sc,'variable_random']['slots']
        ax.stackplot([s['t']/60 for s in ss],*[np.array([s[f'active_mean_{b}'] for s in ss]) for b in BIZ],
            labels=BIZ,colors=['#2864A0','#56B4E9','#D88722','#27856B'])
        ax.set_title(f'场景{sc}');ax.set_xlabel('仿真时间（min）');ax.set_ylabel('实际活跃任务数')
    fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',ncol=4,fontsize=9)
    fig.suptitle('原业务时长：不同业务的活跃占用构成',fontsize=12)
    fig.tight_layout(rect=(0,.08,1,.95));fig.savefig(folder/'mixed_06_active_mix.png',dpi=220);plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(8.5,3.6));x=np.arange(3)
    for ax,key,title in zip(axes,('peak_link_util_pct','peak_node_util_pct'),('最忙链路','最忙计算节点')):
        for j,var in enumerate(VARIANTS):
            ax.bar(x+(j-.5)*.34,[runs[sc,var]['stats']['ALL'][key] for sc in (1,2,3)],.34,
                color=COLOR[var],label=LABEL[var])
        ax.set_xticks(x,['场景1','场景2','场景3']);ax.set_title(title);ax.set_ylabel('峰值利用率（%）');ax.grid(axis='y',alpha=.2)
    finish(fig,axes,'mixed_07_resources.png','混合业务的资源瓶颈对比')


def picture(document,name,caption):
    from docx.shared import Cm,Pt
    document.add_picture(str(OUT/'figures'/name),width=Cm(17))
    document.paragraphs[-1].paragraph_format.keep_with_next=True
    p=document.add_paragraph(caption)
    for r in p.runs:r.italic=True;r.font.size=Pt(9)


def append_document(runs):
    if not BACKUP.exists():shutil.copy2(DOC,BACKUP)
    document=Document(BACKUP)
    original_paragraphs=len(document.paragraphs)
    document.add_page_break()
    document.add_heading('11. 追加：3600秒混合业务最小时延实验',level=1)
    paragraph(document,'本章在前述三场景、算力/链路配置、S1最小时延策略的基础上，新增两种时间模型的混合业务测试。'
        'LL∶HR∶HC∶MC的生成任务数严格为3∶2∶3∶1，每种时间模型在三个场景各运行3600 s、120个30 s时间片，共6组。'
        '每个场景先完成固定30 s实验，再恢复原时长并随机到达。前文单业务36组原始数据及正文保持保留。')
    document.add_heading('11.1 负载与测试条件',level=2)
    table(document,['条件','第一组：固定30秒','第二组：原时长、随机到达'],[
        ['业务构成','每片严格3∶2∶3∶1','相同数量、相同端点与速率'],
        ['任务时长','全部30 s','LL 60–300；HR 120–600；HC/MC 600–1800 s'],
        ['到达时刻','片首','所有业务均匀随机分布于片内[0,30) s'],
        ['拓扑与路由','片内静态；片首统一调度','片内静态；到达时分配；片首重算全部存续任务'],
        ['释放资源','任务结束时','按真实结束时刻；迁移时释放旧分配'],
        ['能耗统计','实际30 s服务分段积分','实际服务各分段积分，截断于3600 s'],
        ['片内抢占','不启用','不启用；同时刻按优先级调度，未来任务不提前预留']])
    table(document,['场景','并发硬上限','常规片任务数','峰值片任务数','总生成数'],[
        [sc,{1:200,2:250,3:500}[sc],sum(runs[sc,'fixed30']['cfg']['traffic']['fixed_arrivals_per_slot'].values()),
         sum(runs[sc,'fixed30']['cfg']['traffic']['fixed_arrival_overrides']['5'].values()),
         int(runs[sc,'fixed30']['stats']['ALL']['generated'])] for sc in (1,2,3)])
    paragraph(document,'每10片的第5号片为峰值片（150、450、…、3450 s）。为了每片严格保持3∶2∶3∶1，'
        '总请求数取9的倍数；常规片按约60%上限取最近值，峰值至少达到上限。'
        '场景1/2因此也会有少量主动过载请求，按并发上限拒绝，不能解释为物理网络不可行。')
    paragraph(document,'两组使用相同主随机数流，时长和到达偏移使用独立时间随机数流（种子1042），'
        '逐任务校验源、目的、速率、QoS、处理参数和生成片号完全一致。配置仅在各实验effective_config.json中覆盖；'
        'config.yaml中原有四类持续时间未被替换为30 s。每种业务的端点限制与前文一致，LL始终禁止GEO端点。')
    table(document,['原时长随机偏移均值s','LL','HR','HC','MC'],[
        [f'场景{sc}',*[f"{runs[sc,'variable_random']['audit']['arrival_offset_mean_s'][b]:.3f}" for b in BIZ]]
        for sc in (1,2,3)])
    paragraph(document,'均匀[0,30)分布的理论均值为15 s、标准差约8.66 s。偏移均值用于检查四类业务均有片内随机到达；'
        '到达精确时刻与初始分配时刻在逐任务审计中相等，求解程序墙钟耗时未额外加进模拟网络时延。')
    paragraph(document,'为缩短长时仿真，最小时延策略使用同一目标的延迟候选路径重建：先按Dijkstra距离分组，'
        '只为有机会入选的计算节点重建完整路径，近等时延候选按原公式精算并保持原顺序。'
        '28个全规模样本与原实现的计算节点、路径、时延和能耗逐项一致；不改变容量或QoS。')
    paragraph(document,'软件验证包含6项定向检查：原4项策略目标/可行性检查，加上真实到达、完成释放、跨片续服、'
        '期末截断、能耗积分及近等时延候选一致性。另将历史单业务首片120个任务的路由与能耗逐项回归，结果一致。'
        '数据层再独立检查任务比例、两组配对、资源范围、路由段连续性、QoS和服务时间积分。')
    document.add_heading('11.2 时间与能耗口径修正',level=2)
    paragraph(document,'旧片首批处理不能直接代表随机片内到达，且“仅取最后一次Allocation的整任务能耗”不能代表跨片迁移的能量积分。'
        '本轮采用可选事件循环：拓扑仍每30 s更新，片内仅处理真实到达和完成事件；每到新片，对当前仍存在的任务重路由。'
        '例如任务在14 s开始、34 s结束，则生成[14,30)和[30,34)两段，绝不为两次路由各重复计算20 s能耗。')
    paragraph(document,'每段能耗Ecomp=η×计算吞吐×段长；Etx=该段路径单位秒能耗×段长。'
        '任务能耗为所有已服务分段之和，包含后续掉线任务已消耗的能量。资源占用按真实持续时间积分，'
        '时间平均并发=所有任务实际服务秒数之和/3600，不能用片末活跃数的算术平均替代。'
        '固定30 s组在片末已全部释放，片末活跃数可能为0，但整片平均并发仍为该片接纳数。')
    paragraph(document,'任务持续时间代表流会话的存活时间，不直接加到毫秒级的单块传播＋处理QoS中。'
        '当前仍未计入协议排队、重传、路由计算墙钟和状态迁移通信开销；迁移次数只记录计算节点变化，'
        '不意味着每次迁移都是物理上必要的。纯时延排序下，同等代价候选变化也会产生迁移。')
    paragraph(document,'任务状态分为：初始拒绝、已完成、跨片掉线、期末仍在服。3600 s时仍在服是右删失观测，'
        '只统计截至3600 s已发生的服务和能耗，不把其未来部分计入，也不称为完整任务成功。'
        '“完成或期末在服比例”用于与原最终接纳口径衔接；完整完成数和掉线数另外列出。')
    document.add_heading('11.3 六组总体结果',level=2)
    table(document,['场景/组别','接纳/生成','已完成','期末在服','掉线','平均/峰值并发'],[
        [f'{sc}/{LABEL[v]}',f"{int(r['initially_accepted'])}/{int(r['generated'])}",
         int(r['completed']),int(r['active_at_end']),int(r['dropped']),
         f"{r['time_avg_active']:.2f}/{int(r['peak_concurrency'])}"]
        for sc in (1,2,3) for v in VARIANTS for r in [runs[sc,v]['stats']['ALL']]])
    table(document,['场景/组别','上限拒绝','初始不可行','平均时延ms','J/Gbit输入','链路峰值%'],[
        [f'{sc}/{LABEL[v]}',int(r['rejected_cap']),int(r['rejected_no_feasible']),
         f"{r['lifetime_delay_mean_ms']:.3f}",f"{r['energy_J_per_input_Gbit']:.3f}",
         f"{r['peak_link_util_pct']:.3f}"]
        for sc in (1,2,3) for v in VARIANTS for r in [runs[sc,v]['stats']['ALL']]])
    picture(document,'mixed_01_concurrency.png','图9：片内时间平均并发。固定30 s组有周期性峰值，原时长组可能因长期积压而持续贴近硬上限。')
    picture(document,'mixed_05_outcomes.png','图10：任务最终状态构成。期末在服不等于完整成功；拒绝按上限和不可行分开统计。')
    document.add_heading('11.4 分业务完整指标',level=2)
    for sc in (1,2,3):
        document.add_heading(f'场景{sc}',level=3)
        for v in VARIANTS:
            paragraph(document,LABEL[v])
            table(document,['业务','接纳率','完成/在服/掉线','均值/P95 ms','能耗J/Gbit','平均/峰值活跃'],[
                [b,f"{r['initial_accept_rate']:.2%}",
                 f"{int(r['completed'])}/{int(r['active_at_end'])}/{int(r['dropped'])}",
                 f"{r['lifetime_delay_mean_ms']:.2f}/{r['lifetime_delay_p95_ms']:.2f}",
                 f"{r['energy_J_per_input_Gbit']:.3f}",f"{r['time_avg_active']:.2f}/{int(r['peak_concurrency'])}"]
                for b in BIZ for r in [runs[sc,v]['stats'][b]]])
            table(document,['业务','最大时延ms','最大丢包率','处理ms','服务Gbit','能耗kJ'],[
                [b,f"{r['lifetime_delay_max_ms']:.3f}",f"{r['lifetime_loss_max']:.2e}",
                 f"{r['compute_delay_service_weighted_ms']:.4f}",f"{r['observed_input_Gbit']:.3f}",
                 f"{r['energy_total_J']/1000:.3f}"] for b in BIZ for r in [runs[sc,v]['stats'][b]]])
    picture(document,'mixed_02_survival.png','图11：分业务完成或期末在服比例，以全部生成请求为分母。较低比例需结合上限拒绝和跨片掉线解释。')
    picture(document,'mixed_03_delay.png','图12：已接纳任务在全生命周期经历的最坏时延，再求任务间均值；包含后来掉线任务已服务分段。')
    picture(document,'mixed_04_energy.png','图13：每Gbit实际服务原始输入的能耗。可减少时长不同、接纳数量不同对平均每任务能耗的干扰。')
    document.add_heading('11.5 资源、路由与算力分配',level=2)
    table(document,['场景/组别','路由段数','重路由','迁移','算力峰值%','功率总和峰值W','求解P99 ms'],[
        [f'{sc}/{LABEL[v]}',int(r['allocation_segments']),int(r['reroutes']),int(r['migrations']),
         f"{r['peak_node_util_pct']:.4f}",f"{r['peak_compute_power_W']:.4f}",f"{r['solve_p99_ms']:.3f}"]
        for sc in (1,2,3) for v in VARIANTS for r in [runs[sc,v]['stats']['ALL']]])
    table(document,['场景/组别','链路CV','源端计算%','地面计算%','服务加权时延ms','总能耗MJ'],[
        [f'{sc}/{LABEL[v]}',f"{r['mean_link_cv']:.3f}",f"{r['source_compute_segment_pct']:.2f}",
         f"{r['ground_compute_segment_pct']:.2f}",f"{r['service_weighted_delay_ms']:.3f}",
         f"{r['energy_total_J']/1e6:.3f}"]
        for sc in (1,2,3) for v in VARIANTS for r in [runs[sc,v]['stats']['ALL']]])
    paragraph(document,'链路CV先按片内事件区间积分再对整个观测窗口平均；源端/地面计算比例按路由服务分段计数，'
        '不等于唯一任务的比例。服务加权时延与计算时延以实际服务秒数为权重，可体现长时业务占用。'
        '所有已执行路由分段都通过时延与丢包约束；因新片无可行路径而掉线的后续未服务区间不计为合格服务。')
    paragraph(document,'功率总和峰值是全网同时运行任务的动态计算功率之和；算力峰值是最忙单节点相对自身容量的比例。'
        '两者含义不同，均不包含卫星平台静态功耗或射频发射机功率。混合实验的链路与节点资源峰值为四类业务共享资源的总体值。')
    table(document,['场景/组别','链路CV','链路Jain','节点CV','节点Jain'],[
        [f'{sc}/{LABEL[v]}',f"{r['mean_link_cv']:.3f}",f"{r['mean_link_jain']:.4f}",
         f"{r['mean_node_cv']:.3f}",f"{r['mean_node_jain']:.4f}"]
        for sc in (1,2,3) for v in VARIANTS for r in [runs[sc,v]['stats']['ALL']]])
    paragraph(document,'Jain指数按每个实际事件区间的利用率分布计算后再做时间平均，包含空闲资源；越接近1越均匀。'
        '各业务的峰值可能发生在不同时刻，不能将分业务峰值直接相加当作全网峰值。')
    picture(document,'mixed_07_resources.png','图14：资源峰值。链路和算力分别用各自容量归一化；计算功率利用率与算力利用率按当前η标定相同。')
    picture(document,'mixed_06_active_mix.png','图15：原时长组的业务活跃构成。生成比例严格3∶2∶3∶1，不代表活跃任务占比也应为该比例。')
    document.add_heading('11.6 结果分析',level=2)
    for sc in (1,2,3):
        fixed=runs[sc,'fixed30']['stats']['ALL'];variable=runs[sc,'variable_random']['stats']['ALL']
        paragraph(document,f"场景{sc}：固定30 s组初始接纳率{fixed['initial_accept_rate']:.2%}，"
            f"平均并发{fixed['time_avg_active']:.2f}；恢复原时长后初始接纳率{variable['initial_accept_rate']:.2%}，"
            f"平均并发{variable['time_avg_active']:.2f}，上限拒绝{int(variable['rejected_cap'])}个、"
            f"跨片掉线{int(variable['dropped'])}个。链路峰值由{fixed['peak_link_util_pct']:.3f}%变为"
            f"{variable['peak_link_util_pct']:.3f}%。已接纳后未发生掉线的比例为"
            f"{variable['final_survived']/max(1,variable['initially_accepted']):.2%}（包含期末仍在服的右删失任务）。")
    paragraph(document,'固定30 s组的任务在同一片首进入，按LL、HR、HC、MC优先级排序。'
        '因此峰值的纯并发上限拒绝往往集中在排队靠后的MC，不代表MC本身带宽或计算需求更难满足。'
        '原时长组按真实到达先后处理，只有同时刻事件才按优先级排序，不为尚未到达的高优先级任务预留资源。')
    paragraph(document,'两组保持相同每片请求数，而混合任务的平均持续时间从30 s变为约673.33 s：'
        '(3×180+2×360+3×1200+1×1200)/9。稳定态潜在并发需求约放大22.44倍，'
        '所以原时长组的大量上限拒绝主要是负载积压，不应归因于路径求解退化。'
        '若以后要在相同“并发负载”下比较时间模型，应另做降低到达强度约22倍的控制实验；本轮没有偷偷调整请求数。')
    paragraph(document,'在无限容量、相同到达比例下，活跃时间占比约为LL 8.91%、HR 11.88%、HC 59.41%、MC 19.80%。'
        '长时业务会占据更多任务秒数；实际还受有限并发、优先级、随机入网顺序和跨片不可行影响。'
        '因此每片生成3∶2∶3∶1与观测到不同的在服比例并不矛盾。')
    for sc in (1,2,3):
        rs=runs[sc,'variable_random']['stats']
        worst=max(BIZ,key=lambda b:rs[b]['dropped']/max(1,rs[b]['initially_accepted']))
        paragraph(document,f"场景{sc}中，相对初始接纳数的掉线比例最高为{worst}："
            f"{int(rs[worst]['dropped'])}/{int(rs[worst]['initially_accepted'])}。"
            '业务的源和目的在生成时确定，跨片仅改变路径和计算节点，不重新挑选目的端；'
            '初始近邻关系不保证整个60–1800 s服务期间始终满足时延。')
    paragraph(document,'场景2的星地卸载可能降低个别任务传播与处理代价，但1 Gbps标称星地链路很容易形成热点，'
        '而地面计算能力仍有余量。场景3的GEO端点继续带来长传播时延，长会话跨片检查更容易暴露尾时延问题。'
        '原时长组的总能耗同时受到在服时长、被拒任务、掉线和路径迁移影响，不能仅比较总J或每任务J；'
        '应结合实际服务数据量归一化能耗和服务时间加权时延。')
    paragraph(document,'接近100%利用率的链路在当前模型中仍只受带宽硬约束，尚无显式的链路队列。'
        '因此报告中的QoS达标不能视为接近饱和时真实分组网络也能保持相同时延，工程上仍应预留容量裕量。')
    failures=read(OUT/'mixed_failure_diagnostics.csv') if (OUT/'mixed_failure_diagnostics.csv').exists() else []
    if failures:
        classified=Counter((r['scenario'],r['variant'],r['kind'],r['cause']) for r in failures)
        table(document,['场景/组别','失败类型','个数','满资源反事实诊断'],[
            [f'{sc}/{LABEL[v]}',kind,count,cause] for (sc,v,kind,cause),count in sorted(classified.items())])
        paragraph(document,'诊断在失败时刻对应的满资源拓扑上计算源到目的的传播下界并加最小固定处理项。'
            '下界已经超过预算时，无论怎样选择计算节点或绕行都无法达标；否则只标为资源/候选限制，'
            '不将启发式未找到路径等同于全局无解。逐任务记录在mixed_failure_diagnostics.csv。')
    document.add_heading('11.7 与前述单业务基准的对照',level=2)
    baseline=read(DOC_DIR/'all_metrics.csv')
    table(document,['场景/业务','单业务均值ms','混合30 s均值ms','单业务J/任务','混合30 s J/任务'],[
        [f"{r['scenario']}/{r['biz']}",f"{float(r['delay_mean_ms']):.3f}",
         f"{runs[int(r['scenario']),'fixed30']['stats'][r['biz']]['lifetime_delay_mean_ms']:.3f}",
         f"{float(r['energy_mean_J']):.6f}",
         f"{runs[int(r['scenario']),'fixed30']['stats'][r['biz']]['energy_per_initially_accepted_J']:.6f}"]
        for r in baseline if r['strategy']=='latency'])
    paragraph(document,'该表是观察性参考：单业务基准运行300 s且一种业务独享总并发预算；'
        '本轮运行3600 s，四类业务共享预算，且流量生成比例固定。网络物理参数相同，但观察窗口和竞争组成改变，'
        '差异不能全部解释为算法本身的提升或下降。')
    document.add_heading('11.8 后续方案与复现',level=2)
    paragraph(document,'建议分别处理三类问题：并发过载可在算法层引入会话准入、按业务配额或等待队列；'
        '跨片掉线可尝试时延裕量、有限滚动窗口可行性预测和对旧会话的稳定性惩罚；'
        '星地热点可使用前文验证有效的预测负载代价或联合能耗代价。'
        '这些是后续算法实验建议，本轮始终采用最小时延策略，没有改容量和QoS来掩盖拒绝。')
    paragraph(document,'输出位于results/mixed_3600_latency_20261016/{fixed30,variable_random}/sc{1,2,3}/。'
        '每组保留原六类输出，另有lifecycle_tasks.csv（完整生命周期）、event_timeline.csv（真实时刻并发与资源）、'
        'mixed_metrics.csv、effective_config.json和run_audit.json。统一指标mixed_all_metrics.csv共30行，'
        '时片指标mixed_all_slot_metrics.csv共720行，审计见mixed_audit_manifest.json。')
    paragraph(document,'复现命令：python experiments/run_mixed_3600.py --scenario 1 --variant fixed30；'
        '再执行 --variant variable_random，场景替换为2、3。追加文档：python experiments/append_mixed_3600_report.py。'
        '首次追加前会保留同目录“追加混合实验前备份.docx”，重复生成从该备份重新追加，避免重复章节。')
    report_path=DOC
    try:
        document.save(report_path)
    except PermissionError:
        # A user may have the baseline report open in Word. Keep it intact and
        # deliver the complete appended report under the same name elsewhere.
        report_path=OUT/DOC.name
        document.save(report_path)
    return original_paragraphs,len(document.paragraphs),report_path


def main():
    runs=audit_data();plots(runs)
    before,after,report_path=append_document(runs)
    old=Document(BACKUP);new=Document(report_path)
    assert [p.text for p in old.paragraphs]==[p.text for p in new.paragraphs[:len(old.paragraphs)]]
    for a,b in zip(old.tables,new.tables):
        assert [[c.text for c in r.cells] for r in a.rows]==[[c.text for c in r.cells] for r in b.rows]
    manifest=dict(report=str(report_path),baseline_file_locked=report_path!=DOC,
        backup=str(BACKUP),new_figures=7,original_content_preserved='PASS',
        paragraph_count_before=before,paragraph_count_after=after,
        original_backup_sha256=hashlib.sha256(BACKUP.read_bytes()).hexdigest())
    (OUT/'append_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=True))


if __name__=='__main__':main()
