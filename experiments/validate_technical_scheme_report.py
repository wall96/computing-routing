"""Validate the new report, update its TOC in an owned hidden Word instance, export PDF."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import re
import shutil
import zipfile
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
DEFAULT=ROOT/'results'/'technical_scheme_experiments_20261016'/'智能路由算法技术方案及实验.docx'


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--docx',type=Path,default=DEFAULT);args=ap.parse_args()
    path=args.docx.resolve();out=path.parent;manifest=json.loads((out/'source_manifest.json').read_text(encoding='utf-8'))
    doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes)==manifest['figures']==37
    assert len(doc.tables)==manifest['tables']
    for name in ('S0','S1','S2','S3','S4','S5','非配对','每bit','期末在服','流守恒','模型','代码框架'):
        assert name in text,name
    math_count=len(doc.element.xpath('.//m:oMath'))
    assert math_count>=manifest['equations']==43
    for source in manifest['source_documents']:
        assert hashlib.sha256((ROOT/source['path']).read_bytes()).hexdigest()==source['sha256'],'Existing report changed since source snapshot'
    with zipfile.ZipFile(path) as z:assert z.testzip() is None
    import win32com.client
    word=opened=None
    try:
        word=win32com.client.DispatchEx('Word.Application');word.Visible=False;word.DisplayAlerts=0;word.AutomationSecurity=3
        opened=word.Documents.Open(str(path),ReadOnly=False,AddToRecentFiles=False,Visible=False)
        opened.Fields.Update()
        for toc in opened.TablesOfContents:toc.Update()
        opened.Repaginate()
        for toc in opened.TablesOfContents:toc.UpdatePageNumbers()
        opened.Save()
        pdf=out/(path.stem+'.pdf');opened.ExportAsFixedFormat(str(pdf),17)
    finally:
        if opened is not None:opened.Close(False)
        if word is not None:word.Quit()
    import pymupdf
    pdf_doc=pymupdf.open(pdf);empty=[];outside=[];pages=[];previews=out/'page_previews';previews.mkdir(exist_ok=True)
    selected={0,1,2,3,len(pdf_doc)-1,len(pdf_doc)-2}
    selected.update(range(5,len(pdf_doc),10))
    for i,page in enumerate(pdf_doc):
        content=page.get_text();pages.append(content)
        if len(content.strip())<25 and not page.get_images():empty.append(i+1)
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                for span in line.get('spans',[]):
                    x0,y0,x1,y1=span['bbox']
                    if x0<8 or x1>page.rect.width-8:
                        outside.append(dict(page=i+1,text=span['text'],bbox=list(span['bbox'])))
        if any(key in content for key in ('5.6 S5','5.6  S5','9. E1','15. E7','16. E8','附录A.')):selected.add(i)
    for i in sorted(selected):
        pix=pdf_doc[i].get_pixmap(matrix=pymupdf.Matrix(1,1));pix.save(str(previews/f'page_{i+1:03}.png'))
    (out/'rendered_text.txt').write_text('\n\n'.join(pages),encoding='utf-8')
    labels=set(int(n) for n in re.findall(r'图\s*(\d+)\s', '\n'.join(pages)))
    assert labels>=set(range(1,manifest['figures']+1)),('Missing figure captions',labels)
    assert not empty,('Empty pages',empty)
    result=dict(docx_zip='PASS',figures=manifest['figures'],equation_definitions=manifest['equations'],native_math_objects=math_count,
        tables=manifest['tables'],pages=len(pdf_doc),empty_pages=empty,horizontal_overflow=outside,
        word_export='PASS',toc_updated='PASS',all_figure_captions_found='PASS',source_reports_unchanged='PASS')
    (out/'report_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    mirror=ROOT/path.name
    if out==DEFAULT.parent:
        try:
            shutil.copy2(path,mirror);shutil.copy2(pdf,ROOT/pdf.name)
            result['project_root_copy']='PASS'
        except PermissionError:result['project_root_copy']='LOCKED; use the complete report in its results directory'
        (out/'report_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=True),flush=True)


if __name__=='__main__':main()
