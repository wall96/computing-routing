"""Replace the mixed chapter with sparse-background / short-peak experiments.

Keep the original single-business report and the existing fixed-30s reference.
The old high-arrival raw runs are archived, not overwritten or relabelled.
"""
from __future__ import annotations
import hashlib
import json
from pathlib import Path
import shutil
import sys
from collections import Counter

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from docx import Document
from experiments import append_mixed_3600_report as report
from experiments.build_all_strategy_report import table,paragraph

BIZ=report.BIZ
OUT=report.SPARSE_OUT


def add_revised_chapter(runs):
    document=Document(report.BACKUP)
    before=len(document.paragraphs)
    document.add_page_break()
    document.add_heading('11. 追加：3600秒混合业务最小时延实验（低背景＋短峰值修订版）',level=1)
    paragraph(document,'本章替换上版混合实验补充，不重复追加旧章节。前文三场景、三策略、四类单业务的36组测试及其正文不变。'
        '固定30 s混合实验沿用上次已验证结果作为参考；本次只重跑原时长混合实验，三个场景各3600 s。'
        '不再沿用“持续高并发、平均至少半满”的单业务压测机制，仅要求实际承载并发在一个短时高峰达到场景上限，其他时段允许低负载。')
    document.add_heading('11.1 场景、业务与不变条件',level=2)
    table(document,['场景','参与网络','并发上限','本轮策略'],[
        ['1','500 LEO；无EDGE、DC、GEO',200,'S1 / latency'],
        ['2','500 LEO＋18 EDGE＋10 DC',250,'S1 / latency'],
        ['3','双层LEO各500＋6 GEO；无地面',500,'S1 / latency']])
    paragraph(document,'业务源、目的均为卫星；场景2的地面节点仅用于中转或计算，LL始终禁止GEO端点。'
        '算力仍为LEO 400、GEO 1350、EDGE 6000、DC 32000 TFLOPS，计算功率上限分别为150、500、1400、10200 W。'
        '拓扑、链路、算力、功率、业务速率、压缩和处理参数、QoS以及路由目标均未更改，本次只调整到达过程。')
    table(document,['业务','原时长s（均匀分布）','时延/丢包预算','生成总数权重','每bit等效操作数'],[
        ['LL','60–300','20 ms / 1%',3,50],['HR','120–600','200 ms / 0.02%',2,200],
        ['HC','600–1800','200 ms / 1%',3,1000],['MC','600–1800','200 ms / 1%',1,'40＋15×3目的＝85']])
    paragraph(document,'3∶2∶3∶1约束的是整个3600 s内的生成任务总数，不强求每个时间片有四类业务或固定比例。'
        '片内到达、持续时间均随机；已接纳任务的比例和活跃时间占比允许不同，不能通过事后删任务来凑接纳比例。'
        '原config.yaml的时长等业务画像保持原样，实际覆盖配置写入每组effective_config.json。')
    document.add_heading('11.2 新随机到达机制及复现方法',level=2)
    paragraph(document,'第一步：计算加权平均时长E[S]=(3×180＋2×360＋3×1200＋1200)/9＝673.333 s。'
        '按Little定律用20%并发上限作为背景负载规划值：背景任务数约为0.20×上限×3600/E[S]，取最接近的9倍数。'
        '这是背景负载的规划值，不是整段观测平均并发的硬性约束；仿真从空系统启动。')
    paragraph(document,'第二步：按全局比例生成有限数量的背景任务，用种子2042独立随机选择120个时间片；'
        '各业务在各片的数量自然波动，可为0。片内到达偏移均匀分布于[0,30) s，独立时间流种子1042，源目的及速率种子42。')
    paragraph(document,'第三步：仅在1800–1830 s增加一次30 s到达脉冲，脉冲本身也按3∶2∶3∶1添加任务。'
        '脉冲数量从0开始按9递增，利用已知的到达和结束时间预测“无限容量下请求重叠数”，'
        '取第一个预测峰值不低于并发上限105%的数量。此处不运行路由求解、不参考接纳结果，也不调容量或QoS。'
        '少量超额请求提供到峰裕量；真实接纳峰值仍须由事件日志独立验证，不能将计划峰值当作实际峰值。')
    table(document,['场景','背景任务','脉冲任务','总生成','请求重叠峰值','实际峰值'],[
        [sc,p['background_tasks'],p['pulse_tasks'],int(s['generated']),p['offered_peak'],int(s['peak_concurrency'])]
        for sc in (1,2,3) for p in [runs[sc,'variable_random']['cfg']['traffic']['arrival_profile']]
        for s in [runs[sc,'variable_random']['stats']['ALL']]])
    table(document,['场景','LL生成','HR生成','HC生成','MC生成','片内偏移均值s LL/HR/HC/MC'],[
        [sc,*[int(runs[sc,'variable_random']['stats'][b]['generated']) for b in BIZ],
         '/'.join(f"{runs[sc,'variable_random']['audit']['arrival_offset_mean_s'][b]:.2f}" for b in BIZ)] for sc in (1,2,3)])
    paragraph(document,'固定30 s参考组仍采用上次每片固定业务比例、常规约60%并发上限、每10片一次峰值的流量，'
        '并未重跑或改写其原始数据。因此两种时长模型不再逐任务配对，请求数和端点样本均可能不同；'
        '两组仅保持物理配置与路由策略一致，图表用于工程现象对照，不用于宣称同等输入下的算法增益。')
    document.add_heading('11.3 事件执行与指标定义',level=2)
    paragraph(document,'拓扑每30 s更新、片内静态。任务实际到达时才分配资源，实际结束时即释放；'
        '每个新片重新计算当前尚未结束的业务，不为未来业务预留资源。片内不抢占，同时刻先完成释放，再按优先级处理到达。'
        '跨片任务按连续服务段记账，任务在14 s到达、34 s结束时，服务区间为[14,30)与[30,34)。')
    table(document,['指标','定义与注意事项'],[
        ['初始接纳率','有初始路由的任务数 / 全部生成任务数；拒绝分并发上限与no_feasible'],
        ['完成/在服/掉线','COMPLETED / ACTIVE_AT_END / DROPPED；期末在服是右删失，不称为完整完成'],
        ['完成或期末在服比例','(已完成＋3600 s仍在服)/生成；另列已接纳后未掉线比例'],
        ['平均/峰值并发','实际服务秒数之和/3600；峰值来自真实事件时刻，不是生成量或片末均值'],
        ['满并发秒数','对active＝场景上限的事件区间长度求和，同时列≥90%上限的持续时间'],
        ['生命周期时延均值/P95/最大','先对每个已接纳任务的全部已服务段取最大端到端时延，再对任务统计；含后来掉线任务'],
        ['服务加权时延/处理时延','按各路由段真实服务秒数加权；与按唯一任务求均值的口径不同'],
        ['端到端时延组成','压缩前路径传播＋节点处理＋压缩后路径传播；处理＝固定项＋参考块bit数×每bit操作数/节点吞吐×负载放大，多播取最差目的'],
        ['能耗与单位数据能耗','按每段实际服务秒数积分计算和传输能耗，包含掉线前消耗，截断于3600 s；J/Gbit按实际服务原始输入计'],
        ['资源与功率','链路/节点峰值按自身容量归一化；全网动态计算功率峰值为同时在服任务功率之和，不含平台静态功耗'],
        ['CV/Jain','事件区间内利用率分布（含空闲资源）计算后按时长平均；CV越小、Jain越大表示越均匀'],
        ['路由段/重算/迁移','一条实际分配是一段；跨片重算不一定换节点；迁移特指计算节点改变'],
        ['求解P99','单次路由求解墙钟耗时99分位；并行执行负载可影响该值，不额外计入模拟网络时延']])
    paragraph(document,'未计入链路串行化发送、显式链路排队、抖动、重传、路由计算等待和迁移状态传输；计算处理项包含1/(1−节点利用率)的简化负载放大（利用率截断于0.95），不是显式排队仿真。任务持续60–1800 s是流会话时长，'
        '不直接累加进毫秒级分组/块时延；近满带宽时的时延结果不能直接等同于真实排队网络表现。')
    document.add_heading('11.4 新原时长实验总体结果及峰值验证',level=2)
    table(document,['场景','接纳/生成','接纳率','完成/在服/掉线','上限拒绝/不可行','平均/峰值并发'],[
        [sc,f"{int(s['initially_accepted'])}/{int(s['generated'])}",f"{s['initial_accept_rate']:.2%}",
         f"{int(s['completed'])}/{int(s['active_at_end'])}/{int(s['dropped'])}",
         f"{int(s['rejected_cap'])}/{int(s['rejected_no_feasible'])}",f"{s['time_avg_active']:.2f}/{int(s['peak_concurrency'])}"]
        for sc in (1,2,3) for s in [runs[sc,'variable_random']['stats']['ALL']]])
    table(document,['场景','首次满并发s','满并发累计s（占比）','≥90%上限s','高峰前均值','高峰后均值'],[
        [sc,f"{s['first_cap_t']:.3f}",f"{s['time_at_cap_s']:.3f}（{s['time_at_cap_pct']:.3f}%）",
         f"{s['time_above_90pct_cap_s']:.3f}",f"{s['pre_pulse_time_avg_active']:.2f}",f"{s['post_pulse_time_avg_active']:.2f}"]
        for sc in (1,2,3) for s in [runs[sc,'variable_random']['stats']['ALL']]])
    paragraph(document,'高峰前均值取[0,1800) s，高峰后取[1830,3600) s；脉冲只集中到达30 s，'
        '不代表这些长时会话也只服务30 s。后半段仍有HC/MC会话逐渐退出，所以负载呈长尾，但不再持续贴近上限。')
    report.picture(document,'mixed_01_concurrency.png','图9：新原时长组画真实事件并发阶梯线；固定30 s参考画每片时间平均并发。虚线为硬上限，可见一次短峰和后续自然回落。')
    report.picture(document,'mixed_05_outcomes.png','图10：两种负载条件下任务状态构成；固定30 s是旧参考，不是相同请求数的配对试验。')
    document.add_heading('11.5 新原时长实验分业务完整指标',level=2)
    for sc in (1,2,3):
        document.add_heading(f'场景{sc}',level=3)
        table(document,['业务','接纳/生成','接纳率','完成/在服/掉线','上限拒绝/不可行','平均/峰值活跃'],[
            [b,f"{int(s['initially_accepted'])}/{int(s['generated'])}",f"{s['initial_accept_rate']:.2%}",
             f"{int(s['completed'])}/{int(s['active_at_end'])}/{int(s['dropped'])}",
             f"{int(s['rejected_cap'])}/{int(s['rejected_no_feasible'])}",f"{s['time_avg_active']:.2f}/{int(s['peak_concurrency'])}"]
            for b in BIZ for s in [runs[sc,'variable_random']['stats'][b]]])
        table(document,['业务','时延均值/P95/最大ms','服务加权时延ms','处理ms','最大丢包率','能耗J/Gbit'],[
            [b,f"{s['lifetime_delay_mean_ms']:.3f}/{s['lifetime_delay_p95_ms']:.3f}/{s['lifetime_delay_max_ms']:.3f}",
             f"{s['service_weighted_delay_ms']:.3f}",f"{s['compute_delay_service_weighted_ms']:.4f}",
             f"{s['lifetime_loss_max']:.2e}",f"{s['energy_J_per_input_Gbit']:.3f}"]
            for b in BIZ for s in [runs[sc,'variable_random']['stats'][b]]])
        table(document,['业务','实际服务秒','服务Gbit','计算/传输能耗kJ','路由段/重算/迁移','源端/地面计算段%'],[
            [b,f"{s['observed_service_seconds']:.2f}",f"{s['observed_input_Gbit']:.3f}",
             f"{s['energy_compute_J']/1000:.3f}/{s['energy_tx_J']/1000:.3f}",
             f"{int(s['allocation_segments'])}/{int(s['reroutes'])}/{int(s['migrations'])}",
             f"{s['source_compute_segment_pct']:.2f}/{s['ground_compute_segment_pct']:.2f}"]
            for b in BIZ for s in [runs[sc,'variable_random']['stats'][b]]])
    report.picture(document,'mixed_02_survival.png','图11：以全部请求为分母的完成或期末在服比例。业务间差异须结合到达顺序、上限拒绝和跨片掉线判断。')
    report.picture(document,'mixed_03_delay.png','图12：已接纳任务生命周期最坏时延的任务均值；新原时长组与固定30 s参考的负载和样本不同。')
    report.picture(document,'mixed_04_energy.png','图13：按实际服务输入数据量归一化的能耗，避免直接比较不同会话时长下每任务J造成误读。')
    document.add_heading('11.6 资源、功率与固定30秒参考',level=2)
    table(document,['场景/组别','接纳/生成','均值/P95时延ms','平均/峰值并发','总能耗MJ','J/Gbit'],[
        [f'{sc}/{report.LABEL[v]}',f"{int(s['initially_accepted'])}/{int(s['generated'])}",
         f"{s['lifetime_delay_mean_ms']:.3f}/{s['lifetime_delay_p95_ms']:.3f}",
         f"{s['time_avg_active']:.2f}/{int(s['peak_concurrency'])}",f"{s['energy_total_J']/1e6:.3f}",
         f"{s['energy_J_per_input_Gbit']:.3f}"]
        for sc in (1,2,3) for v in report.VARIANTS for s in [runs[sc,v]['stats']['ALL']]])
    table(document,['场景/组别','链路峰值%','节点峰值%','计算功率峰值W','链路CV/Jain','节点CV/Jain','求解P99 ms'],[
        [f'{sc}/{report.LABEL[v]}',f"{s['peak_link_util_pct']:.3f}",f"{s['peak_node_util_pct']:.4f}",
         f"{s['peak_compute_power_W']:.4f}",f"{s['mean_link_cv']:.3f}/{s['mean_link_jain']:.4f}",
         f"{s['mean_node_cv']:.3f}/{s['mean_node_jain']:.4f}",f"{s['solve_p99_ms']:.3f}"]
        for sc in (1,2,3) for v in report.VARIANTS for s in [runs[sc,v]['stats']['ALL']]])
    report.picture(document,'mixed_07_resources.png','图14：全网最忙链路与最忙节点的峰值利用率，二者按不同的自身容量归一化，不是同一资源。')
    report.picture(document,'mixed_06_active_mix.png','图15：新原时长实验的业务活跃构成，HC/MC因会话更长，占用更多任务秒。生成比例不等于在服比例。')
    document.add_heading('11.7 与上版高到达实验对照及具体分析',level=2)
    old_rows={int(float(r['scenario'])):report.numeric(r) for r in report.read(report.LEGACY_OUT/'mixed_all_metrics.csv')
              if r['variant']=='variable_random' and r['biz']=='ALL'}
    table(document,['场景','旧→新请求数','减少比例','旧→新接纳率','旧→新平均并发','旧→新上限拒绝'],[
        [sc,f"{int(o['generated'])}→{int(s['generated'])}",f"{1-s['generated']/o['generated']:.2%}",
         f"{o['initial_accept_rate']:.2%}→{s['initial_accept_rate']:.2%}",
         f"{o['time_avg_active']:.2f}→{s['time_avg_active']:.2f}",f"{int(o['rejected_cap'])}→{int(s['rejected_cap'])}"]
        for sc in (1,2,3) for o in [old_rows[sc]] for s in [runs[sc,'variable_random']['stats']['ALL']]])
    for sc in (1,2,3):
        ss=runs[sc,'variable_random']['stats'];s=ss['ALL'];ll=ss['LL']
        paragraph(document,f"场景{sc}：实际峰值{int(s['peak_concurrency'])}，平均{s['time_avg_active']:.2f}，"
            f"平均仅占上限{s['time_avg_active']/s['peak_concurrency']:.2%}；满并发累计{s['time_at_cap_s']:.3f} s，"
            f"占窗口{s['time_at_cap_pct']:.3f}%，满足短峰达标、平时低负载的目标。"
            f"初始接纳率{s['initial_accept_rate']:.2%}；{int(s['rejected_cap'])}个上限拒绝、"
            f"{int(s['rejected_no_feasible'])}个初始不可行、{int(s['dropped'])}个跨片掉线。"
            f"LL生命周期时延均值{ll['lifetime_delay_mean_ms']:.3f} ms、P95 {ll['lifetime_delay_p95_ms']:.3f} ms、"
            f"最大{ll['lifetime_delay_max_ms']:.3f} ms，实际在服后的未掉线比例"
            f"{s['final_survived']/max(1,s['initially_accepted']):.2%}（含期末右删失任务）。")
    paragraph(document,'接纳率改善主要来自请求到达负载恢复合理，而不是最小时延算法变强。'
        '旧版把30 s会话的高到达强度直接用于平均673.33 s的会话，潜在稳定态并发约放大22.44倍，'
        '持续撞击200/250/500硬上限。本轮降低请求数并集中一次高峰后，能观察到多业务竞争、跨片续服和自然退出，而不是主要测量准入拒绝。')
    paragraph(document,'到达脉冲特意给出约5%的请求重叠裕量，出现少量并发上限拒绝属于预期。'
        '不能把达到上限后的正常拒绝归类成带宽不足，也不需要放大算力、放宽QoS或提高并发上限来消除它。'
        '任务按实际到达先后接纳，只有同一时刻按优先级，因此新实验拒绝不必集中在低优先级MC。')
    for sc in (1,2,3):
        ss=runs[sc,'variable_random']['stats'];total=ss['ALL']
        active_share='/'.join(f"{ss[b]['observed_service_seconds']/total['observed_service_seconds']:.1%}" for b in BIZ)
        paragraph(document,f'场景{sc}的LL/HR/HC/MC实际服务时间占比为{active_share}。'
            '无限容量稳定态的理论时间占比约为8.91%/11.88%/59.41%/19.80%；有限窗口、空系统启动、'
            '一次性脉冲、少量拒绝和随机时长使实测值偏离该近似，不应要求活跃业务也严格满足3∶2∶3∶1。')
    s2=runs[2,'variable_random']['stats']['ALL']
    paragraph(document,f"场景2链路峰值仍达{s2['peak_link_util_pct']:.3f}%，即使全程平均并发已降低，局部星地热点仍可能存在；"
        f"地面计算约占{s2['ground_compute_segment_pct']:.2f}%的路由段。节点算力利用率与动态计算功率远未饱和，"
        '优先改善路径分流和星地链路热点比继续扩大计算容量更有针对性。片内拓扑静止只排除了拓扑变化，不会消除共享带宽和并发名额的竞争。')
    paragraph(document,'能耗以传输项为主；当前η和β模型下计算功率很低，与400 TFLOPS级节点的较大吞吐余量一致，'
        '但这不代表真实卫星平台只有数瓦总功耗。本报告只计算业务增量计算与传输能量，不包括平台静态功耗。'
        '期末仍在服任务尚未完成，不宜用“完整完成数/生成数”单独评价路由成功率，也不应把未来未服务能量计入本窗口。')
    failures=report.read(OUT/'mixed_failure_diagnostics.csv') if (OUT/'mixed_failure_diagnostics.csv').exists() else []
    if failures:
        counts=Counter((r['scenario'],r['variant'],r['kind'],r['cause']) for r in failures)
        table(document,['场景/组别','失败类型','个数','满资源传播下界诊断'],[
            [f'{sc}/{report.LABEL[v]}',kind,count,cause] for (sc,v,kind,cause),count in sorted(counts.items())])
        paragraph(document,'诊断在失败时刻的满资源拓扑上取最短传播下界，加最小固定处理项。下界超过预算才证明无法通过绕路或换计算节点解决；'
            '其他情况仅归为资源/候选限制，不将启发式未找到方案等同于全局无解。表中固定30 s失败为上次参考数据，本轮不重新解释为新失败。')
    document.add_heading('11.8 结果文件、验证与复现',level=2)
    paragraph(document,'本轮原时长结果：results/mixed_3600_latency_sparse_peak_20261016/variable_random/sc1、sc2、sc3。'
        '每组包含generated_tasks、tasks、routing_allocations、resource_allocations、slot_stats等CSV以及kpi.json；'
        'lifecycle_tasks.csv提供累计服务与状态，event_timeline.csv提供真实并发，arrival_schedule.csv提供各片背景与脉冲输入。'
        '路由路径、计算节点、带宽及算力/功率变化可用task_id和allocation_id关联。')
    paragraph(document,'新根目录中的mixed_all_metrics.csv共30行（固定30 s参考＋新原时长，各场景总体及四业务），'
        'mixed_all_slot_metrics.csv共720行；arrival_profile_summary.csv单独列本次3场景到峰时刻、持续秒数与到达参数。'
        '上版高到达原时长原始CSV留在旧目录，仅用于历史对照，不冒充本次输入。审计文件记录数据来源及SHA-256。')
    paragraph(document,'复现：python experiments/run_mixed_3600.py --scenario 1 --variant variable_random --arrival-profile sparse_peak；'
        '场景2、3替换--scenario即可。更新本章：python experiments/revise_mixed_3600_report.py。'
        '重跑旧高到达方案需显式传--arrival-profile legacy_dense，不会默认覆盖历史数据。')
    paragraph(document,'7项自动测试通过；独立结果审计校验全局业务比例、时间预测与实际生成一致、实际峰值和满并发持续时间、'
        'QoS、资源非负且不超容量、真实到达才分配、跨片服务连续、状态分区、路由能量和服务时间积分。'
        '只要求短峰达上限，不再要求平均并发达到上限的一半；本轮单种子、单次脉冲结果是工程样本，非多种子统计显著性结论。')
    return document,before


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    report.OUT=OUT;report.REVISED=True
    report.LABEL['variable_random']='原时长·低背景短峰'
    runs=report.audit_data();report.plots(runs)
    document,before=add_revised_chapter(runs)
    previous=Path(json.loads((report.LEGACY_OUT/'append_manifest.json').read_text(encoding='utf-8'))['report'])
    archive=previous.with_name(previous.stem+'_高到达补充章备份.docx')
    if previous.exists() and not archive.exists():shutil.copy2(previous,archive)
    output=OUT/report.DOC.name
    document.save(output)
    replaced=False
    try:
        document.save(previous);replaced=True
    except PermissionError:
        pass
    original=Document(report.BACKUP);updated=Document(output)
    assert [p.text for p in original.paragraphs]==[p.text for p in updated.paragraphs[:before]]
    for a,b in zip(original.tables,updated.tables):
        assert [[c.text for c in r.cells] for r in a.rows]==[[c.text for c in r.cells] for r in b.rows]
    text='\n'.join(p.text for p in updated.paragraphs)
    assert text.count('11. 追加：3600秒混合业务')==1
    manifest=dict(report=str(output),previous_report=str(previous),previous_supplement_replaced=replaced,
        previous_full_report_backup=str(archive),original_single_business_backup=str(report.BACKUP),
        original_content_preserved='PASS',mixed_chapter_count=1,new_figures=7,
        paragraph_count_before=before,paragraph_count_after=len(updated.paragraphs),
        fixed30_data_root=str(report.LEGACY_OUT/'fixed30'),variable_data_root=str(OUT/'variable_random'),
        source_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    (OUT/'append_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    summary=['# 3600 s原时长混合业务：低背景＋短峰值修订结果','',
             'LL:HR:HC:MC=3:2:3:1（全局生成比例）；只改变到达过程，原始时长、容量和QoS不变。','',
             '|场景|生成|接纳率|平均/峰值并发|满并发秒|已完成/期末在服/掉线|','|---|---:|---:|---:|---:|---|']
    for sc in (1,2,3):
        s=runs[sc,'variable_random']['stats']['ALL']
        summary.append(f"|{sc}|{int(s['generated'])}|{s['initial_accept_rate']:.2%}|{s['time_avg_active']:.2f}/{int(s['peak_concurrency'])}|"
            f"{s['time_at_cap_s']:.3f}|{int(s['completed'])}/{int(s['active_at_end'])}/{int(s['dropped'])}|")
    summary.extend(['','接纳率变化主要来自降低过载请求，不是算法性能提升；期末在服不等于完整完成。',
                    '完整分析、指标说明与7幅混合实验图见同目录Word文档第11章；原单业务正文保留。'])
    (OUT/'混合业务修订分析.md').write_text('\n'.join(summary)+'\n',encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=True))


if __name__=='__main__':main()
