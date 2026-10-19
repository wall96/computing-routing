"""Check DOCX contents, render in hidden Word, and create page previews.

Word/PyMuPDF are optional validation dependencies, not simulation dependencies.
"""
from pathlib import Path
import argparse
import json
import zipfile

from docx import Document

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results'/'all_strategy_single_biz_20261016'


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--docx',type=Path,default=OUT/'全场景全策略单业务测试.docx')
    parser.add_argument('--expected-figures',type=int,help='Explicit figure count for reports with additional experiment chapters')
    args=parser.parse_args()
    path=args.docx.resolve()
    output=path.parent
    doc=Document(path)
    text='\n'.join(p.text for p in doc.paragraphs)
    assert '主要发现' in text and '具体结论与适用策略' in text
    has_mixed='追加：3600秒混合业务' in text
    random_chapters=sum(p.style.name=='Heading 1' and '新随机业务批次' in p.text for p in doc.paragraphs)
    expected_figures=args.expected_figures or ((15 if has_mixed else 8)+8*random_chapters)
    assert len(doc.inline_shapes)==expected_figures
    assert len(doc.tables)>=35
    for required in ('传播时延','负载均衡','能耗优化','可行性补充搜索','300 s','SHA-256'):
        assert required in text, required
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        assert len([n for n in z.namelist() if n.startswith('word/media/')])==expected_figures
    result={'docx_zip':'PASS','embedded_figures':len(doc.inline_shapes),'tables':len(doc.tables)}
    import win32com.client
    word=None
    opened=None
    try:
        word=win32com.client.DispatchEx('Word.Application')
        word.Visible=False
        word.DisplayAlerts=0
        word.AutomationSecurity=3
        opened=word.Documents.Open(str(path),ReadOnly=True,AddToRecentFiles=False,Visible=False)
        opened.Repaginate()
        pdf=output/(path.stem+'.pdf')
        opened.ExportAsFixedFormat(str(pdf),17)
        result['word_pdf_export']='PASS'
    finally:
        if opened is not None:opened.Close(False)
        if word is not None:word.Quit()
    import pymupdf
    document=pymupdf.open(pdf)
    result['pages']=len(document)
    result['empty_pages']=[]
    previews=output/'page_previews';previews.mkdir(exist_ok=True)
    page_text=[]
    for i,page in enumerate(document):
        extracted=page.get_text().strip()
        page_text.append(extracted)
        if len(extracted)<25 and not page.get_images():result['empty_pages'].append(i+1)
    preview_numbers={0,1,2,4,7,len(document)//2,len(document)-2,len(document)-1}
    if has_mixed:preview_numbers.update({17,20,25})
    if '新随机业务批次' in text:preview_numbers.update({29,30,32,35})
    for number in sorted(preview_numbers):
        if number<0 or number>=len(document):continue
        pix=document[number].get_pixmap(matrix=pymupdf.Matrix(1.1,1.1))
        pix.save(str(previews/f'page_{number+1:02}.png'))
    (output/'rendered_text.txt').write_text('\n\n'.join(page_text),encoding='utf-8')
    assert not result['empty_pages'], result
    (output/'document_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':main()
